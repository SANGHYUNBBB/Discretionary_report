from __future__ import annotations

import bisect
import os
import re
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Tuple

try:
    import xlrd
except ImportError as exc:
    raise ImportError(
        "메리츠 원본이 .xls 형식이므로 xlrd가 필요합니다. "
        "명령 프롬프트에서 'pip install xlrd'를 실행한 뒤 다시 실행해 주세요."
    ) from exc

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


# ============================================================
# 환경설정
# ============================================================

BASE_DIR = Path.cwd()

INPUT_PATTERNS = [
    "meritz*.xls",
    "meritz*.xlsx",
    "meritz*.xlsm",
    "MERITZ*.xls",
    "MERITZ*.xlsx",
    "MERITZ*.xlsm",
    "메리츠증권*.xls",
    "메리츠증권*.xlsx",
    "메리츠증권*.xlsm",
]

EXCHANGE_DIR_NAMES = [
    "exchange_rate",
    "Exchange_Rate",
    "EXCHANGE_RATE",
]

OUTPUT_SUFFIX = "_정리.xlsx"

# 증권사 원본에 계좌번호와 이름이 없는 경우 직접 입력할 수 있다.
# 예: ACCOUNT_NUMBER_OVERRIDE = "1234-567890-12"
# 예: HOLDER_NAME_OVERRIDE = "홍길동"
ACCOUNT_NUMBER_OVERRIDE = ""
HOLDER_NAME_OVERRIDE = ""

OUTPUT_COLUMNS = [
    "계좌번호",
    "계약자명",
    "PF명",
    "구분",
    "종목명",
    "매매일자",
    "수량",
    "매매단가",
    "매매금액",
    "위탁매매수수료",
    "각종세금",
]

FOREIGN_CODES = {
    "USD",
    "JPY",
    "HKD",
    "CNY",
    "EUR",
    "GBP",
    "AUD",
    "CAD",
    "CHF",
    "SGD",
    "TWD",
}

# 원본 거래적요가 아래 문자열과 정확히 일치할 때만 제외한다.
# 주의: "해외주식 매수", "해외주식 매도"처럼 중간에 공백이 있는 거래는
# 정상 거래로 반영해야 하므로 여기에서 제외하지 않는다.
SKIP_TYPES = {
    "해외주식매수",
    "해외주식매도",
}

HEADER_ALIASES = {
    "거래종류": {
        "거래적요",
        "거래종류",
        "적요명",
        "거래구분",
        "거래내용",
    },
    "거래일자": {
        "거래일자",
        "거래일",
        "일자",
    },
    "종목명": {
        "종목명",
        "종목",
        "종목명(거래상대명)",
        "거래상대명",
    },
    "통화구분": {
        "통화구분",
        "통화",
        "통화코드",
        "외화구분",
    },
    "거래수량": {
        "수량",
        "거래수량",
        "주문수량",
    },
    "거래단가": {
        "단가",
        "거래단가",
        "거래단가(외화)",
        "주문단가",
    },
    "거래금액": {
        "거래금액",
        "금액",
        "매매금액(외화)",
        "거래금액(외화)",
        "외화금액",
    },
    "수수료": {
        "수수료",
        "수수료금액",
        "수수료(외화)",
        "외화수수료",
    },
    "제세금": {
        "제세금",
        "세금",
        "거래세",
        "세금합계",
        "제비용(외화)",
        "제비용",
    },
}

REQUIRED_HEADERS = [
    "거래종류",
    "거래일자",
    "종목명",
    "통화구분",
    "거래수량",
    "거래단가",
    "거래금액",
    "수수료",
    "제세금",
]

ZERO = Decimal("0")
HUNDRED = Decimal("100")


# ============================================================
# 환율 데이터 구조
# ============================================================

@dataclass
class ExchangeTable:
    dates_ord: List[int]
    rates: List[Decimal]

    def lookup(self, target_date: date) -> Decimal:
        """
        거래일과 동일한 날짜의 환율을 우선 사용한다.
        해당 날짜가 없으면 가장 가까운 이전 날짜의 환율을 사용한다.
        """
        target_ordinal = target_date.toordinal()
        position = bisect.bisect_right(
            self.dates_ord,
            target_ordinal,
        ) - 1

        if position < 0:
            raise ValueError(
                f"{target_date} 이전에 사용할 수 있는 환율이 없습니다."
            )

        return self.rates[position]


# ============================================================
# 공통 함수
# ============================================================

def clean_text(value) -> str:
    if value is None:
        return ""

    return str(value).strip()


def normalize_header(value) -> str:
    """
    헤더의 줄바꿈과 모든 공백을 제거한다.

    예:
    '거래\n번호' -> '거래번호'
    """
    return re.sub(
        r"\s+",
        "",
        clean_text(value),
    )


def normalize_transaction(value) -> str:
    """
    거래적요 비교를 위해 모든 공백을 제거한다.

    예:
    '해외주식 매수' -> '해외주식매수'
    """
    return re.sub(
        r"\s+",
        "",
        clean_text(value),
    )


def safe_excel_text(value) -> str:
    return ILLEGAL_CHARACTERS_RE.sub(
        "",
        clean_text(value),
    )


def set_plain_text(cell, value) -> None:
    cell.value = safe_excel_text(value)
    cell.data_type = "s"


def normalize_numeric_text(value) -> str:
    """
    숫자 문자열에 들어간 쉼표와 모든 종류의 공백을 제거한다.

    예:
    '6 420'   -> '6420'
    '96\u00a0620' -> '96620'
    '1,234.5' -> '1234.5'
    """
    text = unicodedata.normalize(
        "NFKC",
        str(value),
    )

    text = text.replace(",", "")

    text = "".join(
        character
        for character in text
        if not character.isspace()
        and unicodedata.category(character)
        not in {
            "Zs",
            "Zl",
            "Zp",
            "Cf",
        }
    )

    # 괄호 음수도 처리한다. 예: (1,000) -> -1000
    if (
        len(text) >= 2
        and text.startswith("(")
        and text.endswith(")")
    ):
        text = f"-{text[1:-1]}"

    return text


def to_decimal(value) -> Decimal:
    if value is None or value == "":
        return ZERO

    if isinstance(value, Decimal):
        result = value

    elif isinstance(value, bool):
        result = Decimal(int(value))

    else:
        try:
            normalized_value = normalize_numeric_text(
                value
            )

            if normalized_value == "":
                return ZERO

            result = Decimal(normalized_value)

        except (
            InvalidOperation,
            AttributeError,
            ValueError,
        ):
            return ZERO

    if not result.is_finite():
        return ZERO

    return result


def parse_date_safe(value) -> Optional[date]:
    if value is None or value == "":
        return None

    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    if isinstance(value, (int, float)):
        try:
            return date(1899, 12, 30) + timedelta(
                days=int(value)
            )
        except Exception:
            return None

    text = normalize_numeric_text(value)

    for date_format in (
        "%Y%m%d",
        "%Y/%m/%d",
        "%Y-%m-%d",
        "%Y.%m.%d",
        "%Y-%m-%d%H:%M:%S",
    ):
        try:
            return datetime.strptime(
                text,
                date_format,
            ).date()

        except ValueError:
            continue

    return None


# ============================================================
# 파일 검색
# ============================================================

def find_exchange_dir(
    base_dir: Path,
) -> Optional[Path]:
    for directory_name in EXCHANGE_DIR_NAMES:
        directory_path = base_dir / directory_name

        if (
            directory_path.exists()
            and directory_path.is_dir()
        ):
            return directory_path

    return None


def find_input_files(
    base_dir: Path,
) -> List[Path]:
    files: List[Path] = []

    for pattern in INPUT_PATTERNS:
        files.extend(
            base_dir.glob(pattern)
        )

    result = {
        file_path.resolve()
        for file_path in files
        if not file_path.name.startswith("~$")
        and not file_path.stem.endswith("_정리")
        and "_작성중" not in file_path.stem
    }

    return sorted(result)


# ============================================================
# 환율 처리
# ============================================================

def adjust_fx_rate(
    currency_code: str,
    rate: Decimal,
) -> Decimal:
    """
    JPY 환율은 100엔당 환율이므로 100으로 나눈다.
    다른 통화는 그대로 사용한다.
    """
    if clean_text(currency_code).upper() == "JPY":
        return rate / HUNDRED

    return rate


def extract_currency_code_from_filename(
    file_path: Path,
) -> str:
    stem_upper = file_path.stem.strip().upper()

    if stem_upper in FOREIGN_CODES:
        return stem_upper

    match = re.search(
        r"(?<![A-Z])([A-Z]{3})(?![A-Z])",
        stem_upper,
    )

    if match:
        return match.group(1)

    return stem_upper


def load_exchange_rates(
    exchange_dir: Optional[Path],
) -> Dict[str, ExchangeTable]:
    """
    exchange_rate 폴더의 xlsx/xlsm 파일을 읽는다.

    - 파일명은 통화구분과 동일한 이름을 우선 사용한다.
    - 첫 번째 시트의 A열 날짜, C열 환율을 읽는다.
    - 날짜가 있는 모든 행을 읽으므로 시작 행이 바뀌어도 처리된다.
    """
    rate_map: Dict[str, ExchangeTable] = {}

    if exchange_dir is None:
        return rate_map

    candidate_files: List[Path] = []

    for extension in (
        "*.xlsx",
        "*.xlsm",
    ):
        candidate_files.extend(
            exchange_dir.glob(extension)
        )

        candidate_files.extend(
            exchange_dir.glob(
                f"**/{extension}"
            )
        )

    seen_paths = set()

    for file_path in candidate_files:
        resolved_path = str(
            file_path.resolve()
        )

        if resolved_path in seen_paths:
            continue

        if file_path.name.startswith("~$"):
            continue

        seen_paths.add(resolved_path)

        currency_code = (
            extract_currency_code_from_filename(
                file_path
            )
        )

        if currency_code not in FOREIGN_CODES:
            continue

        workbook = load_workbook(
            file_path,
            data_only=True,
            read_only=True,
        )

        try:
            worksheet = workbook[
                workbook.sheetnames[0]
            ]

            date_rate_pairs: List[
                Tuple[int, Decimal]
            ] = []

            for row_number in range(
                1,
                worksheet.max_row + 1,
            ):
                exchange_date = parse_date_safe(
                    worksheet.cell(
                        row_number,
                        1,
                    ).value
                )

                exchange_rate = to_decimal(
                    worksheet.cell(
                        row_number,
                        3,
                    ).value
                )

                if (
                    exchange_date is None
                    or exchange_rate == ZERO
                ):
                    continue

                date_rate_pairs.append(
                    (
                        exchange_date.toordinal(),
                        exchange_rate,
                    )
                )

        finally:
            workbook.close()

        if not date_rate_pairs:
            raise ValueError(
                "환율 파일에서 날짜와 환율 데이터를 "
                f"찾지 못했습니다: {file_path}"
            )

        date_rate_pairs.sort(
            key=lambda item: item[0]
        )

        rate_map[
            currency_code
        ] = ExchangeTable(
            dates_ord=[
                item[0]
                for item in date_rate_pairs
            ],
            rates=[
                item[1]
                for item in date_rate_pairs
            ],
        )

    return rate_map


def get_fx_rate(
    row: dict,
    fx_tables: Dict[str, ExchangeTable],
) -> Decimal:
    currency_code = clean_text(
        row.get("통화구분")
    ).upper()

    if currency_code in (
        "",
        "KRW",
    ):
        return Decimal("1")

    trade_date = parse_date_safe(
        row.get("거래일자")
    )

    if trade_date is None:
        raise ValueError(
            "환율 조회에 필요한 거래일자를 "
            "해석할 수 없습니다."
        )

    exchange_table = fx_tables.get(
        currency_code
    )

    if exchange_table is None:
        available_codes = (
            ", ".join(
                sorted(fx_tables.keys())
            )
            if fx_tables
            else "없음"
        )

        raise KeyError(
            f"통화구분 {currency_code}와 같은 "
            "환율 파일을 찾지 못했습니다. "
            f"사용 가능한 통화: {available_codes}"
        )

    raw_rate = exchange_table.lookup(
        trade_date
    )

    return adjust_fx_rate(
        currency_code,
        raw_rate,
    )


# ============================================================
# 헤더 및 계좌정보 처리
# ============================================================

def canonical_header_map(
    raw_headers: Iterable,
) -> Dict[str, int]:
    normalized_headers = [
        normalize_header(value)
        for value in raw_headers
    ]

    header_map: Dict[str, int] = {}

    for canonical_name, aliases in HEADER_ALIASES.items():
        normalized_aliases = {
            normalize_header(alias)
            for alias in aliases
        }

        for column_index, header_value in enumerate(
            normalized_headers
        ):
            if header_value in normalized_aliases:
                header_map[
                    canonical_name
                ] = column_index
                break

    return header_map


def parse_account_info(
    text: str,
) -> Tuple[str, str]:
    text = clean_text(text)

    if text == "":
        return "", ""

    account_match = re.search(
        r"(\d+(?:-\d+)+|\d{8,})",
        text,
    )

    if not account_match:
        return "", ""

    account_number = account_match.group(1)

    holder = text[
        account_match.end():
    ].strip(" []()_-/")

    return account_number, holder


def resolve_account_info(
    top_values: Iterable,
    input_file: Path,
) -> Tuple[str, str]:
    if (
        ACCOUNT_NUMBER_OVERRIDE
        or HOLDER_NAME_OVERRIDE
    ):
        return (
            clean_text(
                ACCOUNT_NUMBER_OVERRIDE
            ),
            clean_text(
                HOLDER_NAME_OVERRIDE
            ),
        )

    for value in top_values:
        account_number, holder = (
            parse_account_info(
                clean_text(value)
            )
        )

        if account_number:
            return account_number, holder

    # 파일명에 계좌번호와 이름이 있는 경우도 지원한다.
    return parse_account_info(
        input_file.stem
    )


# ============================================================
# 거래종류별 계산
# ============================================================

def calculate_row(
    row: dict,
    fx_tables: Dict[str, ExchangeTable],
) -> Tuple[
    Optional[List[Decimal]],
    Optional[str],
]:
    transaction_original = clean_text(
        row.get("거래종류")
    )

    # 중간 공백이 전혀 없는 원본 거래명만 제외한다.
    # "해외주식 매수"와 "해외주식 매도"는 아래 계산 규칙으로 반영한다.
    if transaction_original in SKIP_TYPES:
        return None, None

    transaction = normalize_transaction(
        transaction_original
    )

    quantity = to_decimal(
        row.get("거래수량")
    )

    unit_price = to_decimal(
        row.get("거래단가")
    )

    trade_amount = to_decimal(
        row.get("거래금액")
    )

    fee = to_decimal(
        row.get("수수료")
    )

    tax = to_decimal(
        row.get("제세금")
    )

    currency_code = clean_text(
        row.get("통화구분")
    ).upper()

    # ========================================================
    # 해외주식매도대금 / 해외주식매수대금
    # ========================================================
    if transaction in {
        "해외주식매도대금",
        "해외주식매수대금",
    }:
        exchange_rate = get_fx_rate(
            row,
            fx_tables,
        )

        output_unit_price = (
            unit_price
            * exchange_rate
        )

        return [
            quantity,
            output_unit_price,
            quantity * output_unit_price,
            fee * exchange_rate,
            tax * exchange_rate,
        ], ""

    # ========================================================
    # 해외주식 매수 / 해외주식 매도
    # ========================================================
    # 원본 거래명에 중간 공백이 있는 경우만 이 규칙에 도달한다.
    # 중간 공백이 없는 "해외주식매수", "해외주식매도"는 위에서 제외된다.
    if transaction in {
        "해외주식매수",
        "해외주식매도",
    }:
        exchange_rate = get_fx_rate(
            row,
            fx_tables,
        )

        output_unit_price = (
            unit_price
            * exchange_rate
        )

        return [
            quantity,
            output_unit_price,
            quantity * output_unit_price,
            fee * exchange_rate,
            tax * exchange_rate,
        ], ""

    # ========================================================
    # 환전외화매도(자체) / 환전외화매수(자체)
    # ========================================================
    if transaction in {
        "환전외화매도(자체)",
        "환전외화매수(자체)",
    }:
        exchange_rate = get_fx_rate(
            row,
            fx_tables,
        )

        output_quantity = trade_amount

        return [
            output_quantity,
            exchange_rate,
            output_quantity * exchange_rate,
            fee * exchange_rate,
            tax * exchange_rate,
        ], ""

    # ========================================================
    # 예탁금이용료
    # ========================================================
    if transaction == "예탁금이용료":
        return [
            ZERO,
            ZERO,
            trade_amount,
            fee,
            tax,
        ], ""

    # ========================================================
    # 외화예탁금이용료
    # ========================================================
    if transaction == "외화예탁금이용료":
        exchange_rate = get_fx_rate(
            row,
            fx_tables,
        )

        output_quantity = trade_amount

        return [
            output_quantity,
            exchange_rate,
            output_quantity * exchange_rate,
            fee * exchange_rate,
            tax * exchange_rate,
        ], ""

    # ========================================================
    # 배당금 - 통화구분이 존재하는 외화 배당금
    # ========================================================
    if (
        transaction == "배당금"
        and currency_code not in (
            "",
            "KRW",
        )
    ):
        exchange_rate = get_fx_rate(
            row,
            fx_tables,
        )

        output_quantity = trade_amount

        return [
            output_quantity,
            exchange_rate,
            output_quantity * exchange_rate,
            fee * exchange_rate,
            tax * exchange_rate,
        ], ""

    # ========================================================
    # 배당세금출금(원화)
    # ========================================================
    if transaction == "배당세금출금(원화)":
        return [
            ZERO,
            ZERO,
            ZERO,
            fee,
            tax,
        ], ""

    # ========================================================
    # 무상주입고
    # ========================================================
    if transaction == "무상주입고":
        return [
            quantity,
            ZERO,
            ZERO,
            fee,
            tax,
        ], ""

    # ========================================================
    # 무상세금출금
    # ========================================================
    if transaction == "무상세금출금":
        return [
            quantity,
            ZERO,
            ZERO,
            fee,
            tax,
        ], ""

    return [
        quantity,
        unit_price,
        trade_amount,
        fee,
        tax,
    ], (
        f"규칙 미지정 거래적요: "
        f"{transaction_original}"
    )


# ============================================================
# XLS 읽기
# ============================================================

def find_xls_header(
    sheet,
) -> Tuple[int, Dict[str, int]]:
    search_row_count = min(
        sheet.nrows,
        10,
    )

    for row_index in range(
        search_row_count
    ):
        raw_headers = [
            sheet.cell_value(
                row_index,
                column_index,
            )
            for column_index in range(
                sheet.ncols
            )
        ]

        header_map = canonical_header_map(
            raw_headers
        )

        if all(
            header_name in header_map
            for header_name in REQUIRED_HEADERS
        ):
            return row_index, header_map

    raise KeyError(
        f"시트 '{sheet.name}'에서 변경된 메리츠 헤더를 "
        "찾지 못했습니다."
    )


def read_xls_file(
    input_file: Path,
    fx_tables: Dict[str, ExchangeTable],
) -> Tuple[List[List], List[str]]:
    workbook = xlrd.open_workbook(
        str(input_file)
    )

    output_rows: List[List] = []
    warnings: List[str] = []

    for sheet in workbook.sheets():
        header_row_index, header_map = (
            find_xls_header(sheet)
        )

        top_values = []

        for row_index in range(
            header_row_index
        ):
            for column_index in range(
                sheet.ncols
            ):
                top_values.append(
                    sheet.cell_value(
                        row_index,
                        column_index,
                    )
                )

        account_number, holder = (
            resolve_account_info(
                top_values,
                input_file,
            )
        )

        if not account_number and not holder:
            warnings.append(
                f"[{sheet.name}] 증권사 원본에는 계좌번호와 이름이 없습니다. "
                "코드 상단 ACCOUNT_NUMBER_OVERRIDE와 "
                "HOLDER_NAME_OVERRIDE에 직접 입력해 주세요."
            )

        for row_index in range(
            header_row_index + 1,
            sheet.nrows,
        ):
            row_data = {
                header_name: sheet.cell_value(
                    row_index,
                    column_index,
                )
                for header_name, column_index
                in header_map.items()
            }

            transaction_name = clean_text(
                row_data.get("거래종류")
            )

            trade_date_value = row_data.get(
                "거래일자"
            )

            if (
                transaction_name == ""
                and trade_date_value in (
                    None,
                    "",
                )
            ):
                continue

            try:
                calculation, warning_message = (
                    calculate_row(
                        row_data,
                        fx_tables,
                    )
                )

            except Exception as error:
                calculation = [ZERO] * 5
                warning_message = (
                    f"계산 실패({transaction_name}): "
                    f"{error}"
                )

            if calculation is None:
                continue

            if warning_message:
                warnings.append(
                    f"[{sheet.name} R{row_index + 1}] "
                    f"{warning_message}"
                )

            (
                output_quantity,
                output_unit_price,
                output_trade_amount,
                output_fee,
                output_tax,
            ) = calculation

            output_rows.append(
                [
                    account_number,
                    holder,
                    "",
                    transaction_name,
                    clean_text(
                        row_data.get("종목명")
                    ),
                    parse_date_safe(
                        row_data.get("거래일자")
                    ),
                    float(output_quantity),
                    float(output_unit_price),
                    float(output_trade_amount),
                    float(output_fee),
                    float(output_tax),
                ]
            )

    return output_rows, warnings


# ============================================================
# XLSX/XLSM 읽기
# ============================================================

def find_xlsx_header(
    worksheet,
) -> Tuple[int, Dict[str, int]]:
    search_row_count = min(
        worksheet.max_row,
        10,
    )

    for row_number in range(
        1,
        search_row_count + 1,
    ):
        raw_headers = [
            worksheet.cell(
                row_number,
                column_number,
            ).value
            for column_number in range(
                1,
                worksheet.max_column + 1,
            )
        ]

        zero_based_map = canonical_header_map(
            raw_headers
        )

        if all(
            header_name in zero_based_map
            for header_name in REQUIRED_HEADERS
        ):
            one_based_map = {
                header_name: column_index + 1
                for header_name, column_index
                in zero_based_map.items()
            }

            return row_number, one_based_map

    raise KeyError(
        f"시트 '{worksheet.title}'에서 변경된 메리츠 헤더를 "
        "찾지 못했습니다."
    )


def read_xlsx_file(
    input_file: Path,
    fx_tables: Dict[str, ExchangeTable],
) -> Tuple[List[List], List[str]]:
    workbook = load_workbook(
        input_file,
        data_only=True,
        read_only=True,
    )

    output_rows: List[List] = []
    warnings: List[str] = []

    try:
        for worksheet in workbook.worksheets:
            header_row_number, header_map = (
                find_xlsx_header(
                    worksheet
                )
            )

            top_values = []

            for row_number in range(
                1,
                header_row_number,
            ):
                for column_number in range(
                    1,
                    worksheet.max_column + 1,
                ):
                    top_values.append(
                        worksheet.cell(
                            row_number,
                            column_number,
                        ).value
                    )

            account_number, holder = (
                resolve_account_info(
                    top_values,
                    input_file,
                )
            )

            if not account_number and not holder:
                warnings.append(
                    f"[{worksheet.title}] 증권사 원본에는 계좌번호와 이름이 없습니다. "
                    "코드 상단 ACCOUNT_NUMBER_OVERRIDE와 "
                    "HOLDER_NAME_OVERRIDE에 직접 입력해 주세요."
                )

            for row_number in range(
                header_row_number + 1,
                worksheet.max_row + 1,
            ):
                row_data = {
                    header_name: worksheet.cell(
                        row_number,
                        column_number,
                    ).value
                    for header_name, column_number
                    in header_map.items()
                }

                transaction_name = clean_text(
                    row_data.get("거래종류")
                )

                trade_date_value = row_data.get(
                    "거래일자"
                )

                if (
                    transaction_name == ""
                    and trade_date_value in (
                        None,
                        "",
                    )
                ):
                    continue

                try:
                    calculation, warning_message = (
                        calculate_row(
                            row_data,
                            fx_tables,
                        )
                    )

                except Exception as error:
                    calculation = [ZERO] * 5
                    warning_message = (
                        f"계산 실패({transaction_name}): "
                        f"{error}"
                    )

                if calculation is None:
                    continue

                if warning_message:
                    warnings.append(
                        f"[{worksheet.title} R{row_number}] "
                        f"{warning_message}"
                    )

                (
                    output_quantity,
                    output_unit_price,
                    output_trade_amount,
                    output_fee,
                    output_tax,
                ) = calculation

                output_rows.append(
                    [
                        account_number,
                        holder,
                        "",
                        transaction_name,
                        clean_text(
                            row_data.get("종목명")
                        ),
                        parse_date_safe(
                            row_data.get("거래일자")
                        ),
                        float(output_quantity),
                        float(output_unit_price),
                        float(output_trade_amount),
                        float(output_fee),
                        float(output_tax),
                    ]
                )

    finally:
        workbook.close()

    return output_rows, warnings


def read_input_file(
    input_file: Path,
    fx_tables: Dict[str, ExchangeTable],
) -> Tuple[List[List], List[str]]:
    suffix = input_file.suffix.lower()

    if suffix == ".xls":
        return read_xls_file(
            input_file,
            fx_tables,
        )

    if suffix in {
        ".xlsx",
        ".xlsm",
    }:
        return read_xlsx_file(
            input_file,
            fx_tables,
        )

    raise ValueError(
        f"지원하지 않는 입력 형식입니다: {input_file}"
    )


# ============================================================
# 결과 파일 저장
# ============================================================

def autosize_columns(
    worksheet,
) -> None:
    for column_number, column_cells in enumerate(
        worksheet.iter_cols(
            1,
            worksheet.max_column,
        ),
        start=1,
    ):
        maximum_length = 0

        for cell in column_cells:
            if cell.value is None:
                continue

            maximum_length = max(
                maximum_length,
                len(str(cell.value)),
            )

        worksheet.column_dimensions[
            get_column_letter(
                column_number
            )
        ].width = min(
            max(
                maximum_length + 2,
                10,
            ),
            28,
        )


def save_output(
    output_path: Path,
    rows: List[List],
    warnings: List[str],
) -> None:
    temporary_path = output_path.with_name(
        f".{output_path.stem}_작성중.xlsx"
    )

    if temporary_path.exists():
        temporary_path.unlink()

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "정리"

    for column_number, header_name in enumerate(
        OUTPUT_COLUMNS,
        start=1,
    ):
        cell = worksheet.cell(
            1,
            column_number,
        )

        set_plain_text(
            cell,
            header_name,
        )

        cell.font = Font(
            bold=True
        )

    for row_number, row_values in enumerate(
        rows,
        start=2,
    ):
        for column_number, value in enumerate(
            row_values,
            start=1,
        ):
            cell = worksheet.cell(
                row_number,
                column_number,
            )

            if column_number <= 5:
                set_plain_text(
                    cell,
                    value,
                )
            else:
                cell.value = value

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = (
        worksheet.dimensions
    )

    if worksheet.max_row >= 2:
        for row_cells in worksheet.iter_rows(
            min_row=2,
            max_row=worksheet.max_row,
        ):
            row_cells[5].number_format = "yyyy-mm-dd"

            for column_index in (
                6,
                7,
                8,
                9,
                10,
            ):
                row_cells[
                    column_index
                ].number_format = "#,##0.00"

    autosize_columns(
        worksheet
    )

    if warnings:
        warning_sheet = workbook.create_sheet(
            "검토필요"
        )

        warning_sheet.cell(
            1,
            1,
            "메시지",
        ).font = Font(
            bold=True
        )

        for row_number, warning_message in enumerate(
            warnings,
            start=2,
        ):
            set_plain_text(
                warning_sheet.cell(
                    row_number,
                    1,
                ),
                warning_message,
            )

        warning_sheet.column_dimensions[
            "A"
        ].width = 120

    try:
        workbook.save(
            temporary_path
        )

        workbook.close()

        verification_workbook = load_workbook(
            temporary_path,
            data_only=False,
            read_only=True,
        )

        verification_workbook.close()

        try:
            os.replace(
                temporary_path,
                output_path,
            )

        except PermissionError as error:
            raise PermissionError(
                f"'{output_path.name}' 파일이 Excel에서 열려 있습니다. "
                "파일을 완전히 닫고 다시 실행해 주세요."
            ) from error

    except Exception:
        workbook.close()

        if temporary_path.exists():
            try:
                temporary_path.unlink()
            except OSError:
                pass

        raise


# ============================================================
# 메인 실행
# ============================================================

def main() -> None:
    input_files = find_input_files(
        BASE_DIR
    )

    if not input_files:
        raise FileNotFoundError(
            f"작업폴더({BASE_DIR})에서 메리츠 원본을 찾지 못했습니다. "
            f"예상 패턴: {', '.join(INPUT_PATTERNS)}"
        )

    exchange_directory = find_exchange_dir(
        BASE_DIR
    )

    if exchange_directory is None:
        raise FileNotFoundError(
            f"작업폴더({BASE_DIR}) 안에서 exchange_rate 폴더를 "
            "찾지 못했습니다."
        )

    fx_tables = load_exchange_rates(
        exchange_directory
    )

    if not fx_tables:
        raise ValueError(
            "exchange_rate 폴더에서 사용할 수 있는 "
            "환율 파일을 찾지 못했습니다."
        )

    for input_file in input_files:
        rows, warnings = read_input_file(
            input_file,
            fx_tables,
        )

        output_path = input_file.with_name(
            f"{input_file.stem}{OUTPUT_SUFFIX}"
        )

        save_output(
            output_path,
            rows,
            warnings,
        )

        print(
            f"완료: {output_path}"
        )

        if warnings:
            print(
                "  - 검토필요 건수: "
                f"{len(warnings)}"
            )


if __name__ == "__main__":
    main()
