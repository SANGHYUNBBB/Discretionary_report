from __future__ import annotations

import bisect
import re
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

BASE_DIR = Path.cwd()
INPUT_PATTERNS = [
    "meritz26q1.xlsx",
    "MERITZ26q1.xlsx",
    "메리츠증권*.xlsx",
    "메리츠증권*.xlsm",
]
EXCHANGE_DIR_NAMES = ["exchange_rate", "Exchange_Rate", "EXCHANGE_RATE"]
OUTPUT_SUFFIX = "_정리.xlsx"

OUTPUT_COLUMNS = [
    "계좌번호", "계약자명", "PF명", "구분", "종목명", "매매일자",
    "수량", "매매단가", "매매금액", "위탁매매수수료", "각종세금",
]

SKIP_TYPES = {"해외주식매수", "해외주식매도"}

HEADER_ALIASES = {
    "거래종류": ["거래종류", "적요명", "구분", "거래구분", "거래내용"],

    # 🔥 주문일자만 우선으로 따로 처리
    "주문일자": ["주문일자", "주문일"],

    # 거래일자는 따로 분리
    "거래일자": ["거래일자", "거래일", "일자"],

    "종목명": ["종목명", "종목", "종목명(거래상대명)", "거래상대명"],
    "통화구분": ["통화구분", "통화", "통화코드", "외화구분"],
    "거래수량": ["거래수량", "수량", "주문수량"],
    "거래단가(외화)": ["거래단가(외화)", "거래단가", "단가", "주문단가"],
    "매매금액(외화)": ["매매금액(외화)", "거래금액(외화)", "금액(외화)", "외화금액"],
    "수수료(외화)": ["수수료(외화)", "외화수수료", "수수료외화"],
    "제비용(외화)": ["제비용(외화)", "외화제비용", "제비용외화"],
}


@dataclass
class ExchangeTable:
    dates_ord: List[int]
    rates: List[Decimal]

    def lookup(self, target_date: date) -> Decimal:
        ordinal = target_date.toordinal()
        pos = bisect.bisect_right(self.dates_ord, ordinal) - 1
        if pos < 0:
            raise ValueError(f"{target_date} 이전 환율이 없습니다.")
        return self.rates[pos]


def to_decimal(value) -> Decimal:
    if value is None or value == "":
        return Decimal("0")
    if isinstance(value, Decimal):
        return value
    if isinstance(value, bool):
        return Decimal(int(value))
    try:
        return Decimal(str(value).replace(",", "").strip())
    except (InvalidOperation, AttributeError):
        return Decimal("0")


def clean_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def parse_date_safe(val) -> Optional[date]:
    if val is None:
        return None
    if isinstance(val, datetime):
        return val.date()
    if isinstance(val, date):
        return val
    if isinstance(val, (int, float)):
        try:
            return date(1899, 12, 30) + timedelta(days=int(val))
        except Exception:
            return None
    text = clean_text(val)
    for fmt in ("%Y/%m/%d", "%Y-%m-%d", "%Y.%m.%d", "%Y%m%d", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def find_exchange_dir(base_dir: Path) -> Optional[Path]:
    for name in EXCHANGE_DIR_NAMES:
        p = base_dir / name
        if p.exists() and p.is_dir():
            return p
    return None


def find_input_files(base_dir: Path) -> List[Path]:
    files: List[Path] = []
    for pattern in INPUT_PATTERNS:
        files.extend(base_dir.glob(pattern))
    return sorted({p.resolve() for p in files})


def adjust_fx_rate(currency_code: str, rate: Decimal) -> Decimal:
    if clean_text(currency_code).upper() == "JPY":
        return rate / Decimal("100")
    return rate


def load_exchange_rates(exchange_dir: Optional[Path]) -> Dict[str, ExchangeTable]:
    rate_map: Dict[str, ExchangeTable] = {}
    if exchange_dir is None:
        return rate_map

    candidates: List[Path] = []
    for ext in ("*.xlsx", "*.xlsm"):
        candidates.extend(exchange_dir.glob(ext))
        candidates.extend(exchange_dir.glob(f"**/{ext}"))

    seen = set()
    unique_candidates: List[Path] = []
    for p in candidates:
        rp = str(p.resolve())
        if rp not in seen:
            seen.add(rp)
            unique_candidates.append(p)

    for file_path in unique_candidates:
        code = file_path.stem.strip().upper()
        wb = load_workbook(file_path, data_only=True, read_only=True)
        ws = wb[wb.sheetnames[0]]

        dates_ord: List[int] = []
        rates: List[Decimal] = []
        for r in range(10, ws.max_row + 1):
            d = parse_date_safe(ws[f"A{r}"].value)
            fx = ws[f"C{r}"].value
            if d is None or fx in (None, ""):
                continue
            dates_ord.append(d.toordinal())
            rates.append(to_decimal(fx))
        wb.close()

        if not dates_ord:
            raise ValueError(f"환율 파일에서 데이터를 찾지 못했습니다: {file_path}")

        paired = sorted(zip(dates_ord, rates), key=lambda x: x[0])
        rate_map[code] = ExchangeTable(
            dates_ord=[x[0] for x in paired],
            rates=[x[1] for x in paired],
        )
    return rate_map


def extract_account_info_from_top(ws) -> Tuple[str, str]:
    text = clean_text(ws["A1"].value)

    # 계좌번호: 하이픈 포함 형식 우선
    m_no = re.search(r"(\d+(?:-\d+)+)", text)
    if m_no:
        account_no = m_no.group(1)
        holder = text[m_no.end():].strip()
        return account_no, holder

    # fallback: 숫자만 있는 계좌번호
    m_digits = re.search(r"(\d{8,})", text)
    if m_digits:
        account_no = m_digits.group(1)
        holder = text[m_digits.end():].strip()
        return account_no, holder

    return "", ""
def build_header_map(ws) -> Tuple[int, Dict[str, int]]:
    for r in range(1, min(ws.max_row, 20) + 1):
        row_values = [clean_text(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
        found = {}
        for canonical, aliases in HEADER_ALIASES.items():
            for idx, val in enumerate(row_values, start=1):
                if val in aliases:
                    found[canonical] = idx
                    break
        if "거래종류" in found and "주문일자" in found:
            return r, found
    raise KeyError("헤더 행을 찾지 못했습니다. 헤더명을 확인해주세요.")


def get_fx_rate(row: dict, fx_tables: Dict[str, ExchangeTable]) -> Decimal:
    currency = clean_text(row.get("통화구분")).upper()
    if currency in ("", "KRW"):
        return Decimal("1")

    order_date = parse_date_safe(row.get("주문일자"))

    # 주문일자가 없으면 거래일자로 fallback
    if order_date is None:
        order_date = parse_date_safe(row.get("거래일자"))

    if order_date is None:
        raise ValueError("주문일자/거래일자 모두 해석 불가")

    table = fx_tables.get(currency)
    if table is None:
        available = ", ".join(sorted(fx_tables.keys())) if fx_tables else "없음"
        raise KeyError(f"통화코드 {currency} 에 해당하는 환율 파일을 찾지 못했습니다. 사용가능 코드: {available}")

    fx = table.lookup(order_date)
    return adjust_fx_rate(currency, fx)


def calculate_row(row: dict, fx_tables: Dict[str, ExchangeTable]) -> Tuple[Optional[List], Optional[str]]:
    tx = clean_text(row.get("거래종류"))
    if tx in SKIP_TYPES:
        return None, None

    qty = to_decimal(row.get("거래수량"))
    unit = to_decimal(row.get("거래단가(외화)"))
    amount_fx = to_decimal(row.get("매매금액(외화)"))
    fee_fx = to_decimal(row.get("수수료(외화)"))
    cost_fx = to_decimal(row.get("제비용(외화)"))
    fx = get_fx_rate(row, fx_tables)

    if tx in {"해외주식매도대금", "해외주식매수대금"}:
        out_qty = qty
        out_unit = unit * fx
        out_amount = out_qty * out_unit
        out_fee = fee_fx * fx
        out_tax = cost_fx * fx
        return [out_qty, out_unit, out_amount, out_fee, out_tax], ""

    if tx in {"환전외화매도(자체)", "환전외화매수(자체)"}:
        out_qty = amount_fx
        out_unit = fx
        out_amount = out_qty * out_unit
        out_fee = fee_fx * fx
        out_tax = cost_fx * fx
        return [out_qty, out_unit, out_amount, out_fee, out_tax], ""

    if tx == "외화예탁금이용료":
        out_qty = amount_fx
        out_unit = fx
        out_amount = out_qty * out_unit
        out_fee = fee_fx * fx
        out_tax = cost_fx * fx
        return [out_qty, out_unit, out_amount, out_fee, out_tax], ""

    note = f"규칙 미지정 거래종류: {tx}"
    return [qty, unit, Decimal("0"), Decimal("0"), Decimal("0")], note


def read_sheet_rows(ws, fx_tables: Dict[str, ExchangeTable]) -> Tuple[List[List], List[str]]:
    account_no, holder = extract_account_info_from_top(ws)
    header_row_idx, header_map = build_header_map(ws)

    required = ["거래종류", "주문일자", "종목명", "통화구분", "거래수량", "거래단가(외화)", "매매금액(외화)", "수수료(외화)", "제비용(외화)"]
    for key in required:
        if key not in header_map:
            raise KeyError(f"시트 '{ws.title}' 에서 필수 헤더 '{key}' 를 찾지 못했습니다.")

    output_rows: List[List] = []
    warnings: List[str] = []

    for r in range(header_row_idx + 1, ws.max_row + 1):
        tx = clean_text(ws.cell(r, header_map["거래종류"]).value)
        dt = ws.cell(r, header_map["주문일자"]).value
        if tx == "" and dt in (None, ""):
            continue

        row = {key: ws.cell(r, col).value for key, col in header_map.items()}

        try:
            calc, note = calculate_row(row, fx_tables)
        except Exception as exc:
            calc, note = [Decimal("0")] * 5, f"계산 실패({tx}): {exc}"

        if calc is None:
            continue

        if note:
            warnings.append(f"[{ws.title} R{r}] {note}")

        out_qty, out_unit, out_amount, out_fee, out_tax = calc

        output_rows.append([
            account_no,
            holder,
            "",
            tx,
            clean_text(row.get("종목명")),
            parse_date_safe(row.get("거래일자")),
            float(out_qty),
            float(out_unit),
            float(out_amount),
            float(out_fee),
            float(out_tax),
        ])

    return output_rows, warnings


def autosize_columns(ws):
    for col_idx, column_cells in enumerate(ws.iter_cols(1, ws.max_column), start=1):
        max_len = 0
        for cell in column_cells:
            if cell.value is None:
                continue
            max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 10), 28)


def save_output(output_path: Path, rows: List[List], warnings: List[str]):
    wb = Workbook()
    ws = wb.active
    ws.title = "정리"
    ws.append(OUTPUT_COLUMNS)
    for row in rows:
        ws.append(row)

    for cell in ws[1]:
        cell.font = Font(bold=True)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        row[5].number_format = "yyyy-mm-dd"
        for idx in [6, 7, 8, 9, 10]:
            row[idx].number_format = "#,##0.00"

    autosize_columns(ws)

    if warnings:
        log_ws = wb.create_sheet("검토필요")
        log_ws.append(["메시지"])
        for msg in warnings:
            log_ws.append([msg])
        log_ws["A1"].font = Font(bold=True)
        log_ws.column_dimensions["A"].width = 120

    wb.save(output_path)


def main():
    input_files = find_input_files(BASE_DIR)
    if not input_files:
        raise FileNotFoundError(
            f"작업폴더({BASE_DIR})에서 입력 파일을 찾지 못했습니다. 예상 파일명: meritz26q1.xlsx"
        )

    exchange_dir = find_exchange_dir(BASE_DIR)
    if exchange_dir is None:
        raise FileNotFoundError(
            f"작업폴더({BASE_DIR}) 안에서 exchange_rate 폴더를 찾지 못했습니다. "
            f"허용 폴더명: {', '.join(EXCHANGE_DIR_NAMES)}"
        )

    fx_tables = load_exchange_rates(exchange_dir)

    for input_file in input_files:
        in_wb = load_workbook(input_file, data_only=True)
        all_rows: List[List] = []
        all_warnings: List[str] = []

        for sheet_name in in_wb.sheetnames:
            ws = in_wb[sheet_name]
            rows, warnings = read_sheet_rows(ws, fx_tables)
            all_rows.extend(rows)
            all_warnings.extend(warnings)

        output_path = input_file.with_name(f"{input_file.stem}{OUTPUT_SUFFIX}")
        save_output(output_path, all_rows, all_warnings)
        print(f"완료: {output_path}")
        if all_warnings:
            print(f"  - 검토필요 건수: {len(all_warnings)}")


if __name__ == "__main__":
    main()
