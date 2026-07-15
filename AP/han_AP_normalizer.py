from collections import defaultdict
from numbers import Number
from pathlib import Path
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


INPUT_FILE_NAME = "한투_AP 및 잔고.xlsx"
OUTPUT_FILE_NAME = "한투_AP 및 잔고_정리.xlsx"

FX_SHEET_NAME = "외화잔고현황"
BALANCE_SHEET_NAME = "잔고내역(원화)"
DETAIL_SHEET_NAME = "해외증권상세내역"

ALL_SHEET_NAME = "전체_내림차순"
OVER_10_SHEET_NAME = "10%초과"

HEADERS = [
    "상품",
    "계좌번호",
    "계약자명",
    "종목명",
    "투자비중",
    "투자사유및전망",
]


def is_number(value):
    """bool을 제외한 숫자인지 확인한다."""
    return isinstance(value, Number) and not isinstance(value, bool)


def clean_text(value):
    """키 비교에 사용할 문자열의 앞뒤 공백을 제거한다."""
    if value is None:
        return ""
    return str(value).strip()


def find_input_file(base_dir: Path) -> Path:
    """
    스크립트와 같은 AP 폴더에서 한투 원본 파일을 찾는다.
    '_정리' 결과 파일과 엑셀 임시 파일은 제외한다.
    """
    preferred = base_dir / INPUT_FILE_NAME
    if preferred.exists():
        return preferred

    candidates = sorted(
        path
        for path in base_dir.glob("한투_AP 및 잔고*.xlsx")
        if "_정리" not in path.stem and not path.name.startswith("~$")
    )

    if not candidates:
        raise FileNotFoundError(
            f"'{INPUT_FILE_NAME}' 파일을 찾지 못했습니다.\n"
            f"확인한 폴더: {base_dir}"
        )

    return candidates[0]


def validate_sheet_headers(fx_ws, balance_ws, detail_ws):
    """사용할 열의 헤더가 예상 구조와 일치하는지 확인한다."""
    expected = [
        (fx_ws, "B1", "계좌번호"),
        (fx_ws, "C1", "계좌명"),
        (balance_ws, "B1", "계좌번호"),
        (balance_ws, "C1", "계좌명"),
        (balance_ws, "O1", "해외증권 총자산"),
        (detail_ws, "B1", "계좌번호"),
        (detail_ws, "C1", "계좌명"),
        (detail_ws, "D1", "종목명"),
        (detail_ws, "M1", "평가금액"),
    ]

    errors = []
    for ws, address, expected_value in expected:
        actual_value = ws[address].value
        if clean_text(actual_value) != expected_value:
            errors.append(
                f"{ws.title}!{address}: "
                f"예상='{expected_value}', 실제={actual_value!r}"
            )

    if errors:
        raise ValueError(
            "원본 파일의 헤더 구조가 예상과 다릅니다.\n- "
            + "\n- ".join(errors)
        )


def build_foreign_balance_keys(fx_ws):
    """
    외화잔고현황에서 (계좌번호, 계좌명) 키 집합을 만든다.
    한 계좌가 통화별로 여러 줄이어도 키는 한 번만 보관한다.
    """
    keys = set()

    for row_no in range(2, fx_ws.max_row + 1):
        account_no = clean_text(fx_ws.cell(row=row_no, column=2).value)
        account_name = clean_text(fx_ws.cell(row=row_no, column=3).value)

        if account_no and account_name:
            keys.add((account_no, account_name))

    return keys


def build_total_asset_map(balance_ws):
    """
    잔고내역(원화)에서 계좌별 해외증권 총자산을 읽는다.

    - B열: 계좌번호
    - C열: 계좌명
    - O열: 해외증권 총자산
    - 데이터 시작: 3행(1~2행은 헤더)
    """
    total_asset_map = {}
    account_order = []

    for row_no in range(3, balance_ws.max_row + 1):
        account_no = clean_text(balance_ws.cell(row=row_no, column=2).value)
        account_name = clean_text(balance_ws.cell(row=row_no, column=3).value)
        total_asset = balance_ws.cell(row=row_no, column=15).value

        if not account_no or not account_name:
            continue

        key = (account_no, account_name)
        account_order.append(key)

        if not is_number(total_asset):
            total_asset_map[key] = None
        else:
            total_asset_map[key] = float(total_asset)

    return total_asset_map, account_order


def extract_records(source_wb):
    """
    세 시트의 계좌번호+계좌명을 키로 연결하여 종목 비중을 계산한다.

    종목명: 해외증권상세내역 D열
    평가금액: 해외증권상세내역 M열
    분모: 잔고내역(원화) O열의 해외증권 총자산

    투자비중 = 해외증권상세내역 평가금액 / 해외증권 총자산
    """
    required_sheets = {
        FX_SHEET_NAME,
        BALANCE_SHEET_NAME,
        DETAIL_SHEET_NAME,
    }
    missing_sheets = required_sheets - set(source_wb.sheetnames)

    if missing_sheets:
        raise KeyError(
            "필수 시트가 없습니다: "
            + ", ".join(sorted(missing_sheets))
        )

    fx_ws = source_wb[FX_SHEET_NAME]
    balance_ws = source_wb[BALANCE_SHEET_NAME]
    detail_ws = source_wb[DETAIL_SHEET_NAME]

    validate_sheet_headers(fx_ws, balance_ws, detail_ws)

    foreign_balance_keys = build_foreign_balance_keys(fx_ws)
    total_asset_map, account_order = build_total_asset_map(balance_ws)

    # 동일 계좌·동일 종목이 여러 행이면 평가금액을 합산한다.
    holdings = defaultdict(lambda: defaultdict(float))

    missing_fx_keys = set()
    missing_balance_keys = set()
    invalid_total_asset_keys = set()

    for row_no in range(2, detail_ws.max_row + 1):
        account_no = clean_text(detail_ws.cell(row=row_no, column=2).value)
        account_name = clean_text(detail_ws.cell(row=row_no, column=3).value)
        stock_name = clean_text(detail_ws.cell(row=row_no, column=4).value)
        valuation = detail_ws.cell(row=row_no, column=13).value

        if not account_no and not account_name and not stock_name:
            continue

        if not account_no or not account_name or not stock_name:
            print(
                f"[건너뜀] {DETAIL_SHEET_NAME} {row_no}행: "
                "계좌번호·계좌명·종목명 중 빈 값이 있습니다."
            )
            continue

        if not is_number(valuation):
            print(
                f"[건너뜀] {DETAIL_SHEET_NAME} {row_no}행: "
                f"평가금액이 숫자가 아닙니다. 값={valuation!r}"
            )
            continue

        key = (account_no, account_name)

        if key not in foreign_balance_keys:
            missing_fx_keys.add(key)
            continue

        if key not in total_asset_map:
            missing_balance_keys.add(key)
            continue

        total_asset = total_asset_map[key]
        if not is_number(total_asset) or float(total_asset) == 0:
            invalid_total_asset_keys.add(key)
            continue

        holdings[key][stock_name] += float(valuation)

    for key in sorted(missing_fx_keys):
        print(f"[건너뜀] 외화잔고현황에서 키를 찾지 못했습니다: {key}")

    for key in sorted(missing_balance_keys):
        print(f"[건너뜀] 잔고내역(원화)에서 키를 찾지 못했습니다: {key}")

    for key in sorted(invalid_total_asset_keys):
        print(f"[건너뜀] 해외증권 총자산이 0 또는 숫자가 아닙니다: {key}")

    all_records = []

    # 잔고내역(원화)의 계좌 순서를 유지하고,
    # 각 계좌 안에서는 투자비중 내림차순으로 정렬한다.
    for key in account_order:
        if key not in holdings:
            continue

        account_no, account_name = key
        total_asset = total_asset_map[key]
        customer_records = []

        for stock_name, valuation in holdings[key].items():
            investment_weight = valuation / total_asset

            customer_records.append([
                "",                    # 상품
                account_no,            # 계좌번호
                account_name,          # 계약자명
                stock_name,            # 종목명
                investment_weight,     # 투자비중
                "",                    # 투자사유및전망
            ])

        customer_records.sort(
            key=lambda record: record[4],
            reverse=True,
        )
        all_records.extend(customer_records)

    return all_records


def write_result_sheet(ws, records):
    """
    결과 시트 작성.

    Excel 복구 경고를 방지하기 위해 Table 객체는 만들지 않고,
    일반 셀 범위와 자동필터만 사용한다.
    """
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin_side = Side(style="thin", color="D9E2F3")
    border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side,
    )

    for column_no, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=column_no, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    for row_no, record in enumerate(records, start=2):
        for column_no, value in enumerate(record, start=1):
            cell = ws.cell(
                row=row_no,
                column=column_no,
                value=value,
            )
            cell.border = border
            cell.alignment = Alignment(
                vertical="center",
                wrap_text=column_no in {4, 6},
            )

        ws.cell(row=row_no, column=5).number_format = "0.00%"

    last_row = max(1, len(records) + 1)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:F{last_row}"
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[1].height = 24

    column_widths = {
        "A": 14,
        "B": 20,
        "C": 16,
        "D": 42,
        "E": 14,
        "F": 36,
    }

    for column_letter, width in column_widths.items():
        ws.column_dimensions[column_letter].width = width


def create_output_workbook(records, output_path: Path):
    """
    하나의 결과 파일에 두 시트를 만든다.

    - 전체_내림차순: 계좌별 전체 종목
    - 10%초과: 투자비중이 정확히 10%를 초과하는 종목
    """
    output_wb = Workbook()
    output_wb.remove(output_wb.active)

    over_10_records = [
        record
        for record in records
        if record[4] > 0.10
    ]

    all_ws = output_wb.create_sheet(ALL_SHEET_NAME)
    write_result_sheet(all_ws, records)

    over_10_ws = output_wb.create_sheet(OVER_10_SHEET_NAME)
    write_result_sheet(over_10_ws, over_10_records)

    output_wb.save(output_path)
    return len(over_10_records)


def main():
    # 이 파일을 REPORT/AP 폴더 안에 두고 실행한다.
    base_dir = Path(__file__).resolve().parent

    # 사용법
    # python hantu_ap_normalizer_openpyxl.py
    # python hantu_ap_normalizer_openpyxl.py "한투_AP 및 잔고.xlsx"
    # python hantu_ap_normalizer_openpyxl.py "원본.xlsx" "결과.xlsx"
    if len(sys.argv) >= 2:
        requested_input = Path(sys.argv[1])
        input_path = (
            requested_input
            if requested_input.is_absolute()
            else base_dir / requested_input
        )
    else:
        input_path = find_input_file(base_dir)

    if not input_path.exists():
        raise FileNotFoundError(
            f"입력 파일을 찾지 못했습니다: {input_path}"
        )

    if len(sys.argv) >= 3:
        requested_output = Path(sys.argv[2])
        output_path = (
            requested_output
            if requested_output.is_absolute()
            else base_dir / requested_output
        )
    else:
        output_path = base_dir / OUTPUT_FILE_NAME

    source_wb = load_workbook(
        input_path,
        data_only=True,
        read_only=False,
    )

    records = extract_records(source_wb)
    over_10_count = create_output_workbook(
        records,
        output_path,
    )

    print(f"원본 파일: {input_path.name}")
    print(f"전체 종목 수: {len(records)}개")
    print(f"10% 초과 종목 수: {over_10_count}개")
    print(f"완료 파일: {output_path}")


if __name__ == "__main__":
    main()
