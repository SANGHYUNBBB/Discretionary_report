# OPENPYXL 전용 버전
from pathlib import Path
from numbers import Number
from collections import defaultdict
import os
import re
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


CONTRACT_LIST_NAME = "ContractList.xlsx"
DEFAULT_OUTPUT_NAME = "SumAPList.xlsx"

ALL_SHEET_NAME = "전체_내림차순"
OVER_10_SHEET_NAME = "10%초과"
CHECK_SHEET_NAME = "미매칭_검토"

OUTPUT_HEADERS = [
    "증권사",
    "상품",
    "계좌번호",
    "계약자명",
    "종목명",
    "투자비중",
    "투자사유및전망",
    "보유수량",
    "취득단가(원화)",
    "매입액원화",
    "시가(분기말기준,원화)",
    "보유잔액(원화)",
]


def is_number(value):
    return isinstance(value, Number) and not isinstance(value, bool)


def clean_text(value):
    if value is None:
        return ""

    return str(value).strip()


def normalize_account(value):
    """
    계좌번호 비교용.
    하이픈, 공백 등 비숫자는 제거한다.
    """
    return re.sub(r"\D", "", clean_text(value))


def normalize_name(value):
    """
    계약자명 비교용.
    공백은 제거한다.
    """
    return re.sub(r"\s+", "", clean_text(value))


def is_cash_or_deposit(stock_name):
    text = clean_text(stock_name).replace(" ", "").lower()

    return (
        "예수금" in text
        or "현금" in text
        or "cash" in text
    )


def locate_dirs():
    """
    실행 위치를 자동 판단한다.

    일반 사용:
    - 이 스크립트가 Report/AP 폴더 안에 있음
    - ContractList.xlsx는 Report 폴더 안에 있음
    - 증권사별 정리 파일들은 Report/AP 폴더 안에 있음

    샌드박스나 테스트:
    - 파일들이 같은 폴더에 있어도 동작
    """
    base_dir = Path(__file__).resolve().parent

    if base_dir.name.lower() == "ap":
        ap_dir = base_dir
        report_dir = base_dir.parent
    elif (base_dir / "AP").exists():
        report_dir = base_dir
        ap_dir = base_dir / "AP"
    else:
        report_dir = base_dir
        ap_dir = base_dir

    return report_dir, ap_dir


def find_contract_list(report_dir, ap_dir):
    candidates = [
        report_dir / CONTRACT_LIST_NAME,
        ap_dir / CONTRACT_LIST_NAME,
    ]

    candidates.extend(sorted(report_dir.glob("ContractList*.xlsx")))
    candidates.extend(sorted(ap_dir.glob("ContractList*.xlsx")))

    for path in candidates:
        if path.exists() and not path.name.startswith("~$"):
            return path

    raise FileNotFoundError(
        "ContractList.xlsx 파일을 찾지 못했습니다.\n"
        f"확인한 폴더: {report_dir}, {ap_dir}"
    )


def read_active_contracts(contract_path):
    """
    ContractList.xlsx에서 살아있는 계약 목록을 읽는다.

    기본 구조:
    - A열: 상품
    - B열: 이름
    - C열: 계좌번호

    이 파일 자체가 살아있는 계약 리스트라는 전제로,
    데이터가 있는 행 전체를 활성 계약으로 본다.
    """
    wb = load_workbook(contract_path, data_only=True, read_only=True)

    try:
        ws = wb[wb.sheetnames[0]]

        active_by_exact_key = {}
        active_by_account = defaultdict(list)
        active_order = []
        duplicate_contracts = []

        for row_no in range(2, ws.max_row + 1):
            product = clean_text(ws.cell(row=row_no, column=1).value)
            name = clean_text(ws.cell(row=row_no, column=2).value)
            account_no = clean_text(ws.cell(row=row_no, column=3).value)

            if not product and not name and not account_no:
                continue

            if not product or not name or not account_no:
                duplicate_contracts.append([
                    "ContractList 불완전 행",
                    row_no,
                    product,
                    name,
                    account_no,
                    "",
                ])
                continue

            account_key = normalize_account(account_no)
            name_key = normalize_name(name)

            if not account_key or not name_key:
                duplicate_contracts.append([
                    "ContractList 키 변환 실패",
                    row_no,
                    product,
                    name,
                    account_no,
                    "",
                ])
                continue

            exact_key = (account_key, name_key)

            info = {
                "product": product,
                "name": name,
                "account_no": account_no,
                "order": len(active_order),
                "source_row": row_no,
            }

            if exact_key in active_by_exact_key:
                duplicate_contracts.append([
                    "ContractList 중복 키",
                    row_no,
                    product,
                    name,
                    account_no,
                    f"기존 행: {active_by_exact_key[exact_key]['source_row']}",
                ])
                continue

            active_by_exact_key[exact_key] = info
            active_by_account[account_key].append(info)
            active_order.append(info)

    finally:
        wb.close()

    return active_by_exact_key, active_by_account, active_order, duplicate_contracts


def workbook_quality_score(path, broker):
    """
    정리 파일 후보의 품질을 점수화한다.

    이름만 보고 고르면 예전 파일을 잡을 수 있어서,
    실제 헤더와 예수금 포함 여부를 같이 확인한다.
    """
    score = 0
    stem = path.stem

    try:
        wb = load_workbook(path, data_only=True, read_only=True)
    except Exception:
        return -10000

    try:
        if ALL_SHEET_NAME in wb.sheetnames:
            ws = wb[ALL_SHEET_NAME]
            headers = [
                clean_text(ws.cell(row=1, column=column_no).value)
                for column_no in range(1, min((ws.max_column or 0), 12) + 1)
            ]

            if "보유잔액(원화)" in headers:
                score += 1000

            elif "투자비중" in headers:
                score += 200

            # 예수금 포함 파일 우선
            stock_col = headers.index("종목명") + 1 if "종목명" in headers else None
            if stock_col:
                for row_no in range(2, min(ws.max_row, 300) + 1):
                    stock_name = ws.cell(row=row_no, column=stock_col).value
                    if is_cash_or_deposit(stock_name):
                        score += 500
                        break

        else:
            # 삼성처럼 고객별 시트 구조인 경우
            non_summary_sheets = [
                sheet_name
                for sheet_name in wb.sheetnames
                if sheet_name not in {"10%초과", "10%초과_모음", "10%이상 종목"}
            ]

            if non_summary_sheets:
                ws = wb[non_summary_sheets[0]]
                headers = [
                    clean_text(ws.cell(row=1, column=column_no).value)
                    for column_no in range(1, min((ws.max_column or 0), 12) + 1)
                ]

                if "보유잔액(원화)" in headers:
                    score += 1000

                elif "투자비중" in headers:
                    score += 200

                score += min(len(non_summary_sheets), 100)

            if "10%초과_모음" in wb.sheetnames:
                score += 100

    finally:
        wb.close()

    # 파일명 힌트는 보조 점수로만 사용한다.
    if "예수금" in stem:
        score += 120
    if "총자산비중" in stem:
        score += 110
    if "확장" in stem:
        score += 100
    if "v5" in stem:
        score += 90
    if "v4" in stem:
        score += 80
    if "v3" in stem:
        score += 70

    # 중간 산출물/요약 전용 파일은 감점
    if "10%요약" in stem:
        score -= 80
    if "현금제외" in stem:
        score -= 40

    return score


def pick_best_file(ap_dir, broker):
    """
    증권사별 정리 파일을 하나만 고른다.
    후보가 여러 개이면 실제 헤더/예수금 포함 여부와 최신 수정시간을 같이 보고 고른다.
    """
    patterns = {
        "삼성": ["삼성_AP 및 잔고*.xlsx"],
        "교보": ["교보_AP 및 잔고_정리*.xlsx"],
        "KB": ["KB_AP 및 잔고_정리*.xlsx", "kb*_정리*.xlsx"],
        "한투": ["한투_AP 및 잔고_정리*.xlsx"],
    }

    candidates = []

    for pattern in patterns[broker]:
        candidates.extend(ap_dir.glob(pattern))

    filtered = []
    for path in candidates:
        name = path.name

        if name.startswith("~$"):
            continue

        if "%20" in name:
            continue

        # 원본 파일 제외
        if "_정리" not in path.stem and "계좌별정리" not in path.stem:
            continue

        # 통합 결과 파일 제외
        if "통합" in path.stem:
            continue

        filtered.append(path)

    if not filtered:
        return None

    def priority(path):
        try:
            mtime = path.stat().st_mtime
        except OSError:
            mtime = 0

        return (
            workbook_quality_score(path, broker),
            mtime,
        )

    return sorted(filtered, key=priority, reverse=True)[0]


def build_header_map(ws):
    header_map = {}

    for column_no in range(1, ws.max_column + 1):
        header = clean_text(ws.cell(row=1, column=column_no).value)

        if header and header not in header_map:
            header_map[header] = column_no

    return header_map


def row_is_empty(values):
    return all(value in (None, "") for value in values)


def parse_row_from_sheet(ws, row_no, header_map):
    """
    각 증권사 정리 파일의 행을 공통 구조로 변환한다.
    6열짜리 과거 파일도 읽을 수 있게 추가 열은 비워 둔다.
    """
    required = [
        "계좌번호",
        "계약자명",
        "종목명",
        "투자비중",
    ]

    if any(header not in header_map for header in required):
        return None

    def get(header):
        col = header_map.get(header)
        if col is None:
            return None
        return ws.cell(row=row_no, column=col).value

    account_no = clean_text(get("계좌번호"))
    name = clean_text(get("계약자명"))
    stock_name = clean_text(get("종목명"))

    if not account_no and not name and not stock_name:
        return None

    weight = get("투자비중")
    if not is_number(weight):
        return None

    return {
        "account_no": account_no,
        "name": name,
        "stock_name": stock_name,
        "weight": float(weight),
        "reason": get("투자사유및전망"),
        "quantity": get("보유수량"),
        "acquisition_unit_krw": get("취득단가(원화)"),
        "buy_amount_krw": get("매입액원화"),
        "market_price_krw": get("시가(분기말기준,원화)"),
        "holding_balance_krw": get("보유잔액(원화)"),
    }


def iter_normalized_rows(file_path, broker):
    """
    증권사 정리 파일에서 전체 행을 읽는다.

    - 교보/KB/한투: 전체_내림차순 시트 우선
    - 삼성: 고객별 시트들 읽기. 10%초과_모음은 제외
    """
    wb = load_workbook(file_path, data_only=True, read_only=True)

    try:
        if ALL_SHEET_NAME in wb.sheetnames:
            target_sheet_names = [ALL_SHEET_NAME]
        else:
            target_sheet_names = [
                name
                for name in wb.sheetnames
                if name not in {
                    "10%초과",
                    "10%초과_모음",
                    "10%이상 종목",
                }
            ]

        for sheet_name in target_sheet_names:
            ws = wb[sheet_name]
            header_map = build_header_map(ws)

            for row_no in range(2, ws.max_row + 1):
                parsed = parse_row_from_sheet(ws, row_no, header_map)

                if parsed is None:
                    continue

                parsed["broker"] = broker
                parsed["source_file"] = file_path.name
                parsed["source_sheet"] = sheet_name
                parsed["source_row"] = row_no

                yield parsed

    finally:
        wb.close()


def resolve_active_contract(row, active_by_exact_key, active_by_account):
    """
    잔고 행이 ContractList의 살아있는 계약에 속하는지 확인한다.

    1순위: 계좌번호 + 계약자명 정확 매칭
    2순위: 계좌번호만 유일하게 매칭되는 경우
    """
    account_key = normalize_account(row["account_no"])
    name_key = normalize_name(row["name"])

    exact_key = (account_key, name_key)

    if exact_key in active_by_exact_key:
        return active_by_exact_key[exact_key], "정확매칭"

    account_matches = active_by_account.get(account_key, [])

    if len(account_matches) == 1:
        return account_matches[0], "계좌번호단독매칭"

    return None, "미매칭"


def collect_records(ap_dir, active_by_exact_key, active_by_account):
    broker_files = {
        broker: pick_best_file(ap_dir, broker)
        for broker in ["삼성", "교보", "KB", "한투"]
    }

    records = []
    unmatched_rows = []
    matched_contract_keys = set()

    for broker, file_path in broker_files.items():
        if file_path is None:
            unmatched_rows.append([
                "정리 파일 없음",
                broker,
                "",
                "",
                "",
                "",
                "",
                "",
                "",
            ])
            continue

        for row in iter_normalized_rows(file_path, broker):
            contract_info, match_type = resolve_active_contract(
                row,
                active_by_exact_key,
                active_by_account,
            )

            if contract_info is None:
                unmatched_rows.append([
                    "ContractList 미매칭으로 제외",
                    broker,
                    row["source_file"],
                    row["source_sheet"],
                    row["account_no"],
                    row["name"],
                    row["stock_name"],
                    row["source_row"],
                    "",
                ])
                continue

            account_key = normalize_account(contract_info["account_no"])
            name_key = normalize_name(contract_info["name"])
            matched_contract_keys.add((account_key, name_key))

            records.append({
                "sort_order": contract_info["order"],
                "broker": broker,
                "product": contract_info["product"],
                "account_no": contract_info["account_no"],
                "name": contract_info["name"],
                "stock_name": row["stock_name"],
                "weight": row["weight"],
                "reason": row["reason"],
                "quantity": row["quantity"],
                "acquisition_unit_krw": row["acquisition_unit_krw"],
                "buy_amount_krw": row["buy_amount_krw"],
                "market_price_krw": row["market_price_krw"],
                "holding_balance_krw": row["holding_balance_krw"],
                "source_file": row["source_file"],
                "source_sheet": row["source_sheet"],
                "match_type": match_type,
            })

    return records, unmatched_rows, matched_contract_keys, broker_files


def build_output_rows(records, over_10_only=False):
    filtered = []

    for record in records:
        if over_10_only:
            if record["weight"] <= 0.10:
                continue

            if is_cash_or_deposit(record["stock_name"]):
                continue

        filtered.append(record)

    filtered.sort(
        key=lambda record: (
            record["sort_order"],
            -float(record["weight"] or 0),
            clean_text(record["stock_name"]),
        )
    )

    rows = []

    for record in filtered:
        rows.append([
            record["broker"],
            record["product"],
            record["account_no"],
            record["name"],
            record["stock_name"],
            record["weight"],
            record["reason"],
            record["quantity"],
            record["acquisition_unit_krw"],
            record["buy_amount_krw"],
            record["market_price_krw"],
            record["holding_balance_krw"],
        ])

    return rows


def style_result_sheet(ws, data_row_count):
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin_side = Side(style="thin", color="D9E2F3")
    border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side,
    )

    for column_no, header in enumerate(OUTPUT_HEADERS, start=1):
        cell = ws.cell(row=1, column=column_no, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    last_row = max(1, data_row_count + 1)

    for row in ws.iter_rows(min_row=2, max_row=last_row, min_col=1, max_col=12):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(
                vertical="center",
                wrap_text=cell.column in {5, 7},
            )

    for row_no in range(2, last_row + 1):
        ws.cell(row=row_no, column=6).number_format = "0.00%"           # 투자비중
        ws.cell(row=row_no, column=8).number_format = "#,##0.########"  # 보유수량
        ws.cell(row=row_no, column=9).number_format = "#,##0.00"        # 취득단가
        ws.cell(row=row_no, column=10).number_format = "#,##0"          # 매입액원화
        ws.cell(row=row_no, column=11).number_format = "#,##0.00"       # 시가
        ws.cell(row=row_no, column=12).number_format = "#,##0"          # 보유잔액

    column_widths = {
        1: 12,  # 증권사
        2: 18,  # 상품
        3: 20,  # 계좌번호
        4: 16,  # 계약자명
        5: 42,  # 종목명
        6: 14,  # 투자비중
        7: 36,  # 투자사유
        8: 14,  # 보유수량
        9: 18,  # 취득단가
        10: 18, # 매입액원화
        11: 22, # 시가
        12: 18, # 보유잔액
    }

    for column_no, width in column_widths.items():
        ws.column_dimensions[get_column_letter(column_no)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:L{last_row}"
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[1].height = 36


def write_result_sheet(wb, sheet_name, rows):
    ws = wb.create_sheet(sheet_name)
    ws.append(OUTPUT_HEADERS)

    for row in rows:
        ws.append(row)

    style_result_sheet(ws, len(rows))


def write_check_sheet(wb, broker_files, unmatched_rows, active_order, matched_contract_keys, duplicate_contracts):
    ws = wb.create_sheet(CHECK_SHEET_NAME)

    header_fill = PatternFill("solid", fgColor="806000")
    header_font = Font(color="FFFFFF", bold=True)
    thin_side = Side(style="thin", color="D9E2F3")
    border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side,
    )

    rows = []

    rows.append(["사용한 정리 파일", "", "", "", "", "", "", "", ""])
    rows.append(["증권사", "파일명", "", "", "", "", "", "", ""])

    for broker, path in broker_files.items():
        rows.append([broker, path.name if path else "찾지 못함", "", "", "", "", "", "", ""])

    rows.append(["", "", "", "", "", "", "", "", ""])
    rows.append(["ContractList에는 있으나 정리 파일에서 발견되지 않은 계약", "", "", "", "", "", "", "", ""])
    rows.append(["상품", "계약자명", "계좌번호", "", "", "", "", "", ""])

    for info in active_order:
        key = (
            normalize_account(info["account_no"]),
            normalize_name(info["name"]),
        )

        if key not in matched_contract_keys:
            rows.append([
                info["product"],
                info["name"],
                info["account_no"],
                "",
                "",
                "",
                "",
                "",
                "",
            ])

    rows.append(["", "", "", "", "", "", "", "", ""])
    rows.append(["제외/검토 행", "", "", "", "", "", "", "", ""])
    rows.append(["사유", "증권사", "파일명", "시트", "계좌번호", "계약자명", "종목명", "원본행", "비고"])

    rows.extend(unmatched_rows)

    rows.append(["", "", "", "", "", "", "", "", ""])
    rows.append(["ContractList 중복/불완전 행", "", "", "", "", "", "", "", ""])
    rows.append(["사유", "원본행", "상품", "계약자명", "계좌번호", "비고", "", "", ""])
    rows.extend(duplicate_contracts)

    for row in rows:
        ws.append(row)

    for row_no in range(1, ws.max_row + 1):
        for column_no in range(1, 10):
            cell = ws.cell(row=row_no, column=column_no)
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)

        first = clean_text(ws.cell(row=row_no, column=1).value)
        if first in {
            "사용한 정리 파일",
            "ContractList에는 있으나 정리 파일에서 발견되지 않은 계약",
            "제외/검토 행",
            "ContractList 중복/불완전 행",
        }:
            for column_no in range(1, 10):
                cell = ws.cell(row=row_no, column=column_no)
                cell.fill = header_fill
                cell.font = header_font

    for column_no, width in {
        1: 26,
        2: 16,
        3: 42,
        4: 18,
        5: 20,
        6: 16,
        7: 42,
        8: 10,
        9: 36,
    }.items():
        ws.column_dimensions[get_column_letter(column_no)].width = width

    ws.freeze_panes = "A1"
    ws.sheet_view.showGridLines = False


def create_output_file(all_rows, over_10_rows, broker_files, unmatched_rows, active_order, matched_contract_keys, duplicate_contracts, output_path):
    wb = Workbook()
    wb.remove(wb.active)

    write_result_sheet(wb, ALL_SHEET_NAME, all_rows)
    write_result_sheet(wb, OVER_10_SHEET_NAME, over_10_rows)
    write_check_sheet(
        wb,
        broker_files=broker_files,
        unmatched_rows=unmatched_rows,
        active_order=active_order,
        matched_contract_keys=matched_contract_keys,
        duplicate_contracts=duplicate_contracts,
    )

    temporary_path = output_path.with_name(f".{output_path.stem}_작성중.xlsx")

    if temporary_path.exists():
        temporary_path.unlink()

    try:
        wb.save(temporary_path)
        wb.close()

        verification_wb = load_workbook(
            temporary_path,
            data_only=False,
            read_only=True,
        )
        verification_wb.close()

        try:
            os.replace(temporary_path, output_path)
        except PermissionError as error:
            raise PermissionError(
                f"'{output_path.name}' 파일이 Excel에서 열려 있습니다. "
                "결과 파일을 완전히 닫고 다시 실행해 주세요."
            ) from error

    except Exception:
        wb.close()
        if temporary_path.exists():
            try:
                temporary_path.unlink()
            except OSError:
                pass
        raise


def main():
    report_dir, ap_dir = locate_dirs()

    if len(sys.argv) >= 2:
        requested_output = Path(sys.argv[1])
        output_path = (
            requested_output
            if requested_output.is_absolute()
            else report_dir / requested_output
        )
    else:
        output_path = report_dir / DEFAULT_OUTPUT_NAME

    contract_path = find_contract_list(report_dir, ap_dir)

    (
        active_by_exact_key,
        active_by_account,
        active_order,
        duplicate_contracts,
    ) = read_active_contracts(contract_path)

    (
        records,
        unmatched_rows,
        matched_contract_keys,
        broker_files,
    ) = collect_records(
        ap_dir,
        active_by_exact_key,
        active_by_account,
    )

    all_rows = build_output_rows(records, over_10_only=False)
    over_10_rows = build_output_rows(records, over_10_only=True)

    create_output_file(
        all_rows=all_rows,
        over_10_rows=over_10_rows,
        broker_files=broker_files,
        unmatched_rows=unmatched_rows,
        active_order=active_order,
        matched_contract_keys=matched_contract_keys,
        duplicate_contracts=duplicate_contracts,
        output_path=output_path,
    )

    print(f"ContractList: {contract_path}")
    print(f"AP 폴더: {ap_dir}")
    print("사용한 정리 파일:")
    for broker, path in broker_files.items():
        print(f"  - {broker}: {path.name if path else '찾지 못함'}")
    print(f"활성 계약 수: {len(active_order)}개")
    print(f"전체 출력 행 수: {len(all_rows)}개")
    print(f"10% 초과 출력 행 수: {len(over_10_rows)}개")
    print(f"검토/제외 행 수: {len(unmatched_rows)}개")
    print(f"완료 파일: {output_path}")


if __name__ == "__main__":
    main()