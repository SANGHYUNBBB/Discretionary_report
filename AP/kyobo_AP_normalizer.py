# OPENPYXL 전용 버전
from pathlib import Path
import sys
from numbers import Number

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


DEFAULT_INPUT_NAME = "교보_AP 및 잔고.xlsx"
DEFAULT_OUTPUT_NAME = "교보_AP 및 잔고_정리.xlsx"

ALL_SHEET_NAME = "전체_내림차순"
OVER_10_SHEET_NAME = "10%초과"

HEADERS = [
    "상품",
    "계좌번호",
    "계약자명",
    "종목명",
    "투자비중",
    "투자사유및전망",
]


def is_number(value):
    """True/False를 제외한 숫자인지 확인한다."""
    return isinstance(value, Number) and not isinstance(value, bool)


def find_input_file(base_dir: Path) -> Path:
    """
    스크립트가 들어 있는 REPORT/AP 폴더에서 교보 원본 파일을 찾는다.
    정확한 파일명이 없으면 '교보_AP 및 잔고*.xlsx' 중 정리 파일이 아닌 것을 찾는다.
    """
    preferred = base_dir / DEFAULT_INPUT_NAME
    if preferred.exists():
        return preferred

    candidates = sorted(
        path
        for path in base_dir.glob("교보_AP 및 잔고*.xlsx")
        if "_정리" not in path.stem and not path.name.startswith("~$")
    )

    if not candidates:
        raise FileNotFoundError(
            f"교보 원본 엑셀 파일을 찾지 못했습니다.\n확인한 폴더: {base_dir}"
        )

    return candidates[0]


def find_stock_header_row(ws):
    """B열='종목명', E열='평가금액'인 종목 헤더 행을 찾는다."""
    for row_no in range(1, ws.max_row + 1):
        if ws.cell(row_no, 2).value == "종목명" and ws.cell(row_no, 5).value == "평가금액":
            return row_no
    return None


def get_asset_values_by_currency(ws, stock_header_row):
    """
    종목 헤더 위쪽에서 통화별 자산평가금액을 수집한다.

    예시:
      A6 = USD
      H6 = 84541.14  (USD 자산평가금액)

    종목의 평가금액(E열)도 USD이므로 같은 통화의 자산평가금액으로 나눈다.
    """
    asset_by_currency = {}

    for row_no in range(1, stock_header_row):
        currency = ws.cell(row_no, 1).value
        asset_value = ws.cell(row_no, 8).value

        if (
            isinstance(currency, str)
            and len(currency.strip()) == 3
            and currency.strip().isalpha()
            and is_number(asset_value)
            and float(asset_value) > 0
        ):
            asset_by_currency[currency.strip().upper()] = float(asset_value)

    return asset_by_currency


def extract_records(input_path: Path):
    """
    각 사람별 시트에서 종목을 추출한다.

    원본 기준:
    - A1: 계약자명
    - B1: 계좌번호
    - 각 종목은 2개 행으로 구성
    - 두 번째 행의 B열: 종목명
    - 두 번째 행의 E열: 평가금액
    - 투자비중 = 평가금액 / 같은 통화의 자산평가금액

    정렬 기준:
    - 원본 시트 순서는 유지
    - 각 고객 안에서는 투자비중 내림차순
    """
    source_wb = load_workbook(input_path, data_only=True)
    all_records = []

    excluded_sheet_names = {
        ALL_SHEET_NAME,
        OVER_10_SHEET_NAME,
        "10%이상 종목",
    }

    for ws in source_wb.worksheets:
        if ws.title in excluded_sheet_names:
            continue

        contractor = ws["A1"].value
        account_no = ws["B1"].value

        if contractor in (None, "") or account_no in (None, ""):
            print(f"[건너뜀] {ws.title}: A1 또는 B1이 비어 있습니다.")
            continue

        stock_header_row = find_stock_header_row(ws)
        if stock_header_row is None:
            print(f"[건너뜀] {ws.title}: 종목 헤더를 찾지 못했습니다.")
            continue

        asset_by_currency = get_asset_values_by_currency(ws, stock_header_row)
        customer_records = []

        # 종목 헤더 다음 행부터 2개 행씩 읽는다.
        # 예: 10~11행이 첫 종목이라면 실제 종목명/평가금액은 11행에 있다.
        for first_row in range(stock_header_row + 1, ws.max_row + 1, 2):
            second_row = first_row + 1
            if second_row > ws.max_row:
                break

            currency = ws.cell(second_row, 1).value
            stock_name = ws.cell(second_row, 2).value
            valuation = ws.cell(second_row, 5).value

            if not isinstance(currency, str):
                continue
            if stock_name in (None, "") or not is_number(valuation):
                continue

            currency = currency.strip().upper()
            total_asset = asset_by_currency.get(currency)

            if total_asset is None or total_asset == 0:
                print(
                    f"[건너뜀] {ws.title} / {stock_name}: "
                    f"{currency} 자산평가금액을 찾지 못했습니다."
                )
                continue

            investment_weight = float(valuation) / total_asset

            customer_records.append(
                [
                    "",                        # 상품
                    str(account_no),           # 계좌번호: B1
                    str(contractor),           # 계약자명: A1
                    str(stock_name),           # 종목명
                    investment_weight,         # 평가금액 / 자산평가금액
                    "",                        # 투자사유및전망
                ]
            )

        customer_records.sort(key=lambda record: record[4], reverse=True)
        all_records.extend(customer_records)

    source_wb.close()
    return all_records


def style_result_sheet(ws, data_row_count):
    """결과 시트에 기본 서식을 적용한다."""
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin_side = Side(style="thin", color="D9E2F3")
    thin_border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side,
    )

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    last_row = data_row_count + 1

    for row in ws.iter_rows(min_row=2, max_row=last_row, min_col=1, max_col=6):
        for cell in row:
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    for row_no in range(2, last_row + 1):
        ws.cell(row_no, 5).number_format = "0.00%"

    widths = {
        1: 14,  # 상품
        2: 18,  # 계좌번호
        3: 18,  # 계약자명
        4: 32,  # 종목명
        5: 14,  # 투자비중
        6: 36,  # 투자사유및전망
    }

    for column_no, width in widths.items():
        ws.column_dimensions[get_column_letter(column_no)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:F{last_row}"


def write_sheet(output_wb, sheet_name, records):
    ws = output_wb.create_sheet(title=sheet_name)
    ws.append(HEADERS)

    for record in records:
        ws.append(record)

    style_result_sheet(ws, len(records))


def create_output_file(records, output_path: Path):
    """
    하나의 엑셀 파일에 두 시트를 만든다.

    1. 전체_내림차순: 고객별 전체 종목을 투자비중 내림차순으로 정리
    2. 10%초과: 투자비중이 정확히 10%를 초과하는 종목만 정리
    """
    output_wb = Workbook()
    default_ws = output_wb.active
    output_wb.remove(default_ws)

    over_10_records = [record for record in records if record[4] > 0.10]

    write_sheet(output_wb, ALL_SHEET_NAME, records)
    write_sheet(output_wb, OVER_10_SHEET_NAME, over_10_records)

    output_wb.save(output_path)
    output_wb.close()

    return len(over_10_records)


def main():
    base_dir = Path(__file__).resolve().parent

    # 실행 예시 1:
    #   python kyobo_ap_normalizer.py
    #
    # 실행 예시 2:
    #   python kyobo_ap_normalizer.py "교보_AP 및 잔고.xlsx"
    #
    # 실행 예시 3:
    #   python kyobo_ap_normalizer.py "원본.xlsx" "결과.xlsx"

    if len(sys.argv) >= 2:
        input_arg = Path(sys.argv[1])
        input_path = input_arg if input_arg.is_absolute() else base_dir / input_arg
    else:
        input_path = find_input_file(base_dir)

    if not input_path.exists():
        raise FileNotFoundError(f"입력 파일을 찾지 못했습니다: {input_path}")

    if len(sys.argv) >= 3:
        output_arg = Path(sys.argv[2])
        output_path = output_arg if output_arg.is_absolute() else base_dir / output_arg
    else:
        output_path = base_dir / DEFAULT_OUTPUT_NAME

    records = extract_records(input_path)
    over_10_count = create_output_file(records, output_path)

    print(f"원본 파일: {input_path.name}")
    print(f"전체 종목 수: {len(records)}개")
    print(f"10% 초과 종목 수: {over_10_count}개")
    print(f"완료 파일: {output_path}")


if __name__ == "__main__":
    main()
