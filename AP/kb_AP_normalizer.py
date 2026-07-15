from pathlib import Path
from numbers import Number
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


INPUT_FILE_NAME = "KB_AP 및 잔고.xlsx"
OUTPUT_FILE_NAME = "KB_AP 및 잔고_정리.xlsx"

ALL_SHEET_NAME = "전체_내림차순"
OVER_10_SHEET_NAME = "10%초과"

HEADERS = [
    "상품",
    "계좌번호",
    "계약자명",
    "종목명",
    "투자비중",
    "투자사유및전망",
]


def is_number(value):
    """bool을 제외한 숫자인지 확인한다."""
    return isinstance(value, Number) and not isinstance(value, bool)


def find_input_file(base_dir: Path) -> Path:
    """
    스크립트와 같은 AP 폴더에서 KB 원본 파일을 찾는다.
    '_정리'가 들어간 결과 파일과 엑셀 임시 파일은 제외한다.
    """
    preferred = base_dir / INPUT_FILE_NAME
    if preferred.exists():
        return preferred

    candidates = sorted(
        path
        for path in base_dir.glob("KB_AP 및 잔고*.xlsx")
        if "_정리" not in path.stem and not path.name.startswith("~$")
    )

    if not candidates:
        raise FileNotFoundError(
            f"'{INPUT_FILE_NAME}' 파일을 찾지 못했습니다.\n"
            f"확인한 폴더: {base_dir}"
        )

    return candidates[0]


def find_net_asset_value(ws):
    """
    1~8행에서 '순자산평가금액'을 찾은 뒤,
    바로 오른쪽 셀의 숫자를 반환한다.
    """
    for row in ws.iter_rows(min_row=1, max_row=min(8, ws.max_row)):
        for cell in row:
            if cell.value == "순자산평가금액":
                value = ws.cell(
                    row=cell.row,
                    column=cell.column + 1,
                ).value

                if is_number(value) and float(value) != 0:
                    return float(value)

    raise ValueError(
        f"'{ws.title}' 시트에서 유효한 순자산평가금액을 찾지 못했습니다."
    )


def extract_records(source_wb):
    """
    각 고객 시트에서 종목 정보를 추출한다.

    원본 기준
    - A1: 계약자명
    - B1: 계좌번호
    - 9행: 종목 헤더
    - H열: 평가금액
    - K열: 종목명
    - 투자비중 = 평가금액 / 순자산평가금액

    정렬 기준
    - 고객 시트 순서는 원본 순서 유지
    - 각 고객 내부에서는 투자비중 내림차순
    """
    all_records = []

    for ws in source_wb.worksheets:
        contractor = ws["A1"].value
        account_no = ws["B1"].value

        if not contractor or not account_no:
            print(
                f"[건너뜀] '{ws.title}' 시트: "
                "A1 계약자명 또는 B1 계좌번호가 비어 있습니다."
            )
            continue

        if ws["H9"].value != "평가금액" or ws["K9"].value != "종목명":
            print(
                f"[건너뜀] '{ws.title}' 시트: "
                "H9='평가금액', K9='종목명' 구조가 아닙니다."
            )
            continue

        try:
            net_asset_value = find_net_asset_value(ws)
        except ValueError as exc:
            print(f"[건너뜀] {exc}")
            continue

        customer_records = []

        for row_no in range(10, ws.max_row + 1):
            valuation = ws.cell(row=row_no, column=8).value    # H열
            stock_name = ws.cell(row=row_no, column=11).value  # K열

            if stock_name in (None, "") and valuation in (None, ""):
                continue

            if stock_name in (None, ""):
                print(
                    f"[건너뜀] '{ws.title}' 시트 {row_no}행: "
                    "종목명이 비어 있습니다."
                )
                continue

            if not is_number(valuation):
                print(
                    f"[건너뜀] '{ws.title}' 시트 {row_no}행: "
                    f"평가금액이 숫자가 아닙니다. 값={valuation!r}"
                )
                continue

            customer_records.append([
                "",                               # 상품
                str(account_no),                  # 계좌번호
                str(contractor),                  # 계약자명
                str(stock_name),                  # 종목명
                float(valuation) / net_asset_value,
                "",                               # 투자사유및전망
            ])

        customer_records.sort(
            key=lambda record: record[4],
            reverse=True,
        )
        all_records.extend(customer_records)

    return all_records


def write_result_sheet(ws, records):
    """
    표(Table) 객체는 만들지 않는다.

    Excel 복구 경고를 피하기 위해
    일반 범위 + 일반 자동필터만 사용한다.
    """
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin_side = Side(style="thin", color="D9E2F3")
    border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side,
    )

    for column_no, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=column_no, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    for row_no, record in enumerate(records, start=2):
        for column_no, value in enumerate(record, start=1):
            cell = ws.cell(
                row=row_no,
                column=column_no,
                value=value,
            )
            cell.border = border
            cell.alignment = Alignment(
                vertical="center",
                wrap_text=column_no in {4, 6},
            )

        ws.cell(row=row_no, column=5).number_format = "0.00%"

    last_row = max(1, len(records) + 1)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:F{last_row}"
    ws.sheet_view.showGridLines = False

    column_widths = {
        "A": 14,
        "B": 20,
        "C": 16,
        "D": 42,
        "E": 14,
        "F": 36,
    }

    for column_letter, width in column_widths.items():
        ws.column_dimensions[column_letter].width = width

    ws.row_dimensions[1].height = 24


def create_output_workbook(records, output_path: Path):
    """
    하나의 파일에 두 시트를 생성한다.

    - 전체_내림차순
    - 10%초과: 투자비중 > 10%
    """
    output_wb = Workbook()
    output_wb.remove(output_wb.active)

    over_10_records = [
        record for record in records
        if record[4] > 0.10
    ]

    all_ws = output_wb.create_sheet(ALL_SHEET_NAME)
    write_result_sheet(all_ws, records)

    over_10_ws = output_wb.create_sheet(OVER_10_SHEET_NAME)
    write_result_sheet(over_10_ws, over_10_records)

    output_wb.save(output_path)

    return len(over_10_records)


def main():
    base_dir = Path(__file__).resolve().parent

    # 사용법
    # python kb_ap_normalizer_openpyxl_fixed.py
    # python kb_ap_normalizer_openpyxl_fixed.py "KB_AP 및 잔고.xlsx"
    # python kb_ap_normalizer_openpyxl_fixed.py "원본.xlsx" "결과.xlsx"
    if len(sys.argv) >= 2:
        requested_input = Path(sys.argv[1])
        input_path = (
            requested_input
            if requested_input.is_absolute()
            else base_dir / requested_input
        )
    else:
        input_path = find_input_file(base_dir)

    if not input_path.exists():
        raise FileNotFoundError(
            f"입력 파일을 찾지 못했습니다: {input_path}"
        )

    if len(sys.argv) >= 3:
        requested_output = Path(sys.argv[2])
        output_path = (
            requested_output
            if requested_output.is_absolute()
            else base_dir / requested_output
        )
    else:
        output_path = base_dir / OUTPUT_FILE_NAME

    source_wb = load_workbook(
        input_path,
        data_only=True,
        read_only=False,
    )

    records = extract_records(source_wb)
    over_10_count = create_output_workbook(
        records,
        output_path,
    )

    print(f"원본 파일: {input_path.name}")
    print(f"고객 시트 수: {len(source_wb.worksheets)}개")
    print(f"전체 종목 수: {len(records)}개")
    print(f"10% 초과 종목 수: {over_10_count}개")
    print(f"완료 파일: {output_path}")


if __name__ == "__main__":
    main()
