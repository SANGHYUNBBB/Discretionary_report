from __future__ import annotations

import bisect
import os
import re
import sys
from collections import OrderedDict, Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from numbers import Number
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


INPUT_FILE_NAME = "삼성_AP 및 잔고.xlsx"
OUTPUT_FILE_NAME = "삼성_AP 및 잔고_정리.xlsx"
SUMMARY_SHEET_NAME = "10%초과_모음"

# None이면 exchange_rate 파일의 가장 최근 날짜 환율을 사용한다.
# 특정 기준일을 강제로 쓰고 싶으면 예: "2026-06-30"
EXCHANGE_LOOKUP_DATE = None

OUTPUT_HEADERS = [
    "상품",
    "계좌번호",
    "계약자명",
    "종목명",
    "투자비중",
    "투자사유및전망",
    "보유수량",
    "취득단가(원화)",
    "매입액원화",
    "시가(분기말기준,원화)",
    "보유잔액(원화)",
]

INVALID_SHEET_CHARS = re.compile(r'[\[\]:*?/\\]')

CURRENCY_CODES = {
    "USD",
    "JPY",
    "HKD",
    "CNY",
    "EUR",
    "GBP",
    "AUD",
    "CAD",
    "CHF",
    "SGD",
    "TWD",
}

ZERO = Decimal("0")
HUNDRED = Decimal("100")


# ============================================================
# 환율 데이터 구조
# ============================================================

@dataclass
class ExchangeTable:
    dates_ord: List[int]
    rates: List[Decimal]

    def lookup(self, target_date: date) -> Decimal:
        """
        기준일과 동일한 날짜의 환율을 우선 사용한다.
        동일 날짜가 없으면 가장 가까운 이전 날짜의 환율을 사용한다.
        """
        target_ordinal = target_date.toordinal()

        position = bisect.bisect_right(
            self.dates_ord,
            target_ordinal,
        ) - 1

        if position < 0:
            raise ValueError(
                f"{target_date} 이전에 사용할 수 있는 환율이 없습니다."
            )

        return self.rates[position]

    def latest(self) -> Decimal:
        """환율 파일에서 가장 최근 날짜의 환율을 사용한다."""
        if not self.rates:
            raise ValueError("환율 데이터가 비어 있습니다.")

        return self.rates[-1]


# ============================================================
# 공통 함수
# ============================================================

def is_number(value):
    return isinstance(value, Number) and not isinstance(value, bool)


def clean_text(value) -> str:
    if value is None:
        return ""

    return str(value).strip()


def to_decimal(value) -> Decimal:
    """
    엑셀 값을 Decimal로 변환한다.
    공란, 숫자가 아닌 값, NaN, 무한대는 0으로 처리한다.
    """
    if value is None or value == "":
        return ZERO

    if isinstance(value, Decimal):
        result = value

    elif isinstance(value, bool):
        result = Decimal(int(value))

    else:
        try:
            result = Decimal(str(value).replace(",", "").strip())

        except (
            InvalidOperation,
            AttributeError,
            ValueError,
        ):
            return ZERO

    if not result.is_finite():
        return ZERO

    return result


def parse_date_safe(value) -> Optional[date]:
    if value is None or value == "":
        return None

    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    text = clean_text(value)

    for date_format in (
        "%Y/%m/%d",
        "%Y-%m-%d",
        "%Y.%m.%d",
        "%Y%m%d",
    ):
        try:
            return datetime.strptime(text, date_format).date()

        except ValueError:
            continue

    return None


def parse_lookup_date(value) -> Optional[date]:
    if value in (None, ""):
        return None

    parsed_date = parse_date_safe(value)

    if parsed_date is None:
        raise ValueError(
            f"EXCHANGE_LOOKUP_DATE 값을 날짜로 해석하지 못했습니다: {value!r}"
        )

    return parsed_date


def is_cash_or_deposit(stock_name):
    """
    10% 초과 통합 시트에서만 제외할 현금성 종목 판별.

    제외 예시:
    - 현금잔고(예수금)
    - USD(외화예수금)
    - CNY(외화예수금)
    - CASH
    """
    normalized = str(stock_name).replace(" ", "").lower()

    return (
        "현금" in normalized
        or "예수금" in normalized
        or "cash" in normalized
    )


# ============================================================
# 파일 및 폴더 검색
# ============================================================

def find_input_file(base_dir: Path) -> Path:
    preferred = base_dir / INPUT_FILE_NAME
    if preferred.exists():
        return preferred

    candidates = sorted(
        path
        for path in base_dir.glob("삼성_AP 및 잔고*.xlsx")
        if "계좌별정리" not in path.stem
        and "_정리" not in path.stem
        and not path.name.startswith("~$")
    )

    if not candidates:
        raise FileNotFoundError(
            f"'{INPUT_FILE_NAME}' 파일을 찾지 못했습니다.\n"
            f"확인한 폴더: {base_dir}"
        )

    return candidates[0]


def find_exchange_dir(base_dir: Path) -> Optional[Path]:
    """
    exchange_rate 폴더를 찾는다.

    스크립트가 REPORT/AP 안에 있어도,
    REPORT/exchange_rate 구조를 자동으로 찾도록 처리한다.
    """
    candidates = [
        base_dir / "exchange_rate",
        base_dir / "Exchange_Rate",
        base_dir / "EXCHANGE_RATE",
        base_dir.parent / "exchange_rate",
        base_dir.parent / "Exchange_Rate",
        base_dir.parent / "EXCHANGE_RATE",
        Path.cwd() / "exchange_rate",
        Path.cwd() / "Exchange_Rate",
        Path.cwd() / "EXCHANGE_RATE",
    ]

    for directory in candidates:
        if directory.exists() and directory.is_dir():
            return directory

    return None


def extract_currency_code_from_filename(file_path: Path) -> str:
    """
    환율 파일명에서 통화코드를 추출한다.

    예:
    USD.xlsx
    USD_환율.xlsx
    환율_USD.xlsx
    """
    stem_upper = file_path.stem.strip().upper()

    match = re.search(
        r"(?<![A-Z])([A-Z]{3})(?![A-Z])",
        stem_upper,
    )

    if match:
        return match.group(1)

    return stem_upper


# ============================================================
# 환율 처리
# ============================================================

def adjust_fx_rate(currency_code: str, rate: Decimal) -> Decimal:
    """
    JPY 환율은 보통 100엔당 환율이므로
    100으로 나누어 1엔당 환율로 바꾼다.
    """
    currency_code = clean_text(currency_code).upper()

    if currency_code == "JPY":
        return rate / HUNDRED

    return rate


def load_xlsx_exchange_file(file_path: Path) -> List[Tuple[int, Decimal]]:
    """
    xlsx/xlsm 환율 파일 읽기.

    파일 구조:
    - 첫 번째 시트 사용
    - A열: 날짜
    - C열: 환율
    - 10행부터 데이터
    """
    workbook = load_workbook(
        file_path,
        data_only=True,
        read_only=True,
    )

    try:
        worksheet = workbook[workbook.sheetnames[0]]
        date_rate_pairs: List[Tuple[int, Decimal]] = []

        for row_number in range(10, worksheet.max_row + 1):
            exchange_date = parse_date_safe(
                worksheet[f"A{row_number}"].value
            )

            exchange_rate = to_decimal(
                worksheet[f"C{row_number}"].value
            )

            if exchange_date is None or exchange_rate == ZERO:
                continue

            date_rate_pairs.append(
                (
                    exchange_date.toordinal(),
                    exchange_rate,
                )
            )

    finally:
        workbook.close()

    return date_rate_pairs


def load_xls_exchange_file(file_path: Path) -> List[Tuple[int, Decimal]]:
    """
    xls 환율 파일 읽기.

    xlrd가 설치되어 있을 때만 동작한다.
    설치가 안 되어 있으면:
        pip install xlrd
    """
    try:
        import xlrd

    except ImportError as error:
        raise ImportError(
            f"{file_path.name} 파일은 .xls 형식입니다. "
            "이 파일을 읽으려면 xlrd가 필요합니다.\n"
            "설치 명령어: pip install xlrd\n"
            "또는 환율 파일을 .xlsx로 저장해 주세요."
        ) from error

    workbook = xlrd.open_workbook(str(file_path))
    worksheet = workbook.sheet_by_index(0)

    date_rate_pairs: List[Tuple[int, Decimal]] = []

    # Excel 기준 10행부터이므로 xlrd zero-based index는 9부터.
    for row_index in range(9, worksheet.nrows):
        raw_date = worksheet.cell_value(row_index, 0)
        raw_rate = worksheet.cell_value(row_index, 2)

        exchange_date = None

        cell_type = worksheet.cell_type(row_index, 0)
        if cell_type == xlrd.XL_CELL_DATE:
            exchange_date = xlrd.xldate_as_datetime(
                raw_date,
                workbook.datemode,
            ).date()
        else:
            exchange_date = parse_date_safe(raw_date)

        exchange_rate = to_decimal(raw_rate)

        if exchange_date is None or exchange_rate == ZERO:
            continue

        date_rate_pairs.append(
            (
                exchange_date.toordinal(),
                exchange_rate,
            )
        )

    return date_rate_pairs


def load_exchange_rates(exchange_dir: Optional[Path]) -> Dict[str, ExchangeTable]:
    """
    exchange_rate 폴더의 환율 파일을 읽는다.

    중요:
    - 같은 통화에 .xlsx/.xlsm 파일과 .xls 파일이 같이 있으면 .xlsx/.xlsm을 우선 사용한다.
    - .xls 파일이 실제 구형 Excel 파일이 아니면 xlrd에서 오류가 날 수 있으므로,
      같은 통화의 .xlsx/.xlsm 대체 파일이 있으면 해당 .xls는 건너뛴다.
    """
    rate_map: Dict[str, ExchangeTable] = {}

    if exchange_dir is None:
        return rate_map

    candidate_files: List[Path] = []

    for extension in (
        "*.xlsx",
        "*.xlsm",
        "*.xls",
    ):
        candidate_files.extend(exchange_dir.glob(extension))
        candidate_files.extend(exchange_dir.glob(f"**/{extension}"))

    unique_files = []
    seen_paths = set()

    for file_path in candidate_files:
        resolved_path = str(file_path.resolve())

        if resolved_path in seen_paths:
            continue

        if file_path.name.startswith("~$"):
            continue

        seen_paths.add(resolved_path)
        unique_files.append(file_path)

    files_by_currency: Dict[str, List[Path]] = {}

    for file_path in unique_files:
        currency_code = extract_currency_code_from_filename(file_path)

        if currency_code not in CURRENCY_CODES:
            continue

        files_by_currency.setdefault(currency_code, []).append(file_path)

    for currency_code, files in sorted(files_by_currency.items()):
        # 같은 통화에 여러 파일이 있으면 xlsx/xlsm을 먼저 사용한다.
        files_sorted = sorted(
            files,
            key=lambda path: (
                0 if path.suffix.lower() in {".xlsx", ".xlsm"} else 1,
                path.name.lower(),
            ),
        )

        last_error = None

        for file_path in files_sorted:
            suffix = file_path.suffix.lower()

            try:
                if suffix in {".xlsx", ".xlsm"}:
                    date_rate_pairs = load_xlsx_exchange_file(file_path)

                elif suffix == ".xls":
                    date_rate_pairs = load_xls_exchange_file(file_path)

                else:
                    continue

                if not date_rate_pairs:
                    raise ValueError(
                        "환율 파일에서 날짜와 환율을 찾지 못했습니다: "
                        f"{file_path}"
                    )

                date_rate_pairs.sort(key=lambda item: item[0])

                rate_map[currency_code] = ExchangeTable(
                    dates_ord=[item[0] for item in date_rate_pairs],
                    rates=[item[1] for item in date_rate_pairs],
                )

                break

            except Exception as error:
                last_error = error
                print(f"[환율파일 건너뜀] {file_path.name}: {error}")
                continue

        if currency_code not in rate_map:
            checked_files = ", ".join(path.name for path in files_sorted)
            raise ValueError(
                f"{currency_code} 환율 파일을 읽지 못했습니다. "
                f"확인한 파일: {checked_files}. "
                f"마지막 오류: {last_error}"
            )

    return rate_map


def get_rate_for_currency(currency_code: str, fx_tables: Dict[str, ExchangeTable]) -> Optional[float]:
    """
    통화코드의 환율을 가져온다.
    EXCHANGE_LOOKUP_DATE가 None이면 가장 최근 환율을 사용한다.
    """
    currency_code = clean_text(currency_code).upper()

    if currency_code in ("", "KRW"):
        return 1.0

    exchange_table = fx_tables.get(currency_code)

    if exchange_table is None:
        return None

    lookup_date = parse_lookup_date(EXCHANGE_LOOKUP_DATE)

    if lookup_date is None:
        rate = exchange_table.latest()
    else:
        rate = exchange_table.lookup(lookup_date)

    return float(adjust_fx_rate(currency_code, rate))


def build_fx_rate_map(fx_tables: Dict[str, ExchangeTable]) -> Dict[str, float]:
    rate_map = {}

    for currency_code in sorted(fx_tables.keys()):
        rate = get_rate_for_currency(currency_code, fx_tables)
        if rate is not None and rate > 0:
            rate_map[currency_code] = rate

    return rate_map


def implied_fx_rate(quantity, current_price, valuation) -> Optional[float]:
    """
    삼성 AP 파일에는 통화코드가 없으므로,
    평가금액 / (잔고수량 * 현재가)로 원화 환산율을 역산한다.

    이 값은 통화 판별에만 사용하고,
    실제 취득단가/시가 계산에는 exchange_rate 폴더의 환율을 우선 사용한다.
    """
    if not (
        is_number(quantity)
        and is_number(current_price)
        and is_number(valuation)
    ):
        return None

    denominator = float(quantity) * float(current_price)

    if denominator == 0:
        return None

    return float(valuation) / denominator


def infer_currency_and_rate(
    stock_name,
    product_name,
    quantity,
    current_price,
    valuation,
    fx_rate_map: Dict[str, float],
):
    """
    삼성 AP 파일에는 통화코드가 따로 없으므로 다음 순서로 처리한다.

    1. 현금/예수금:
       - USD(외화예수금), CNY(외화예수금) 등의 현재가는 이미 원화 환율이므로 1을 사용
    2. 국내 상품:
       - 고객거래상품명이 거래소/ETF/원화 등으로 표시되면 KRW, 환율 1 사용
    3. 해외주식:
       - 평가금액 / (잔고수량 * 현재가)로 암묵적 환율을 계산
       - exchange_rate 폴더의 통화별 환율 중 가장 가까운 통화를 선택
       - 계산 자체는 선택된 통화의 exchange_rate 환율로 수행
    4. exchange_rate 폴더가 없거나 해당 통화를 못 찾으면 암묵적 환율을 fallback으로 사용
    """
    stock_text = clean_text(stock_name).upper()
    product_text = clean_text(product_name)

    if is_cash_or_deposit(stock_name):
        # 외화예수금의 현재가는 이미 원화 환율 형태이므로 여기서는 추가 환산하지 않는다.
        return "CASH", 1.0

    if product_text in {"거래소", "ETF", "원화", "국내주식", "국내ETF"}:
        return "KRW", 1.0

    implied = implied_fx_rate(
        quantity,
        current_price,
        valuation,
    )

    if implied is None or implied <= 0:
        return "KRW", 1.0

    # 원화 상품인데 고객거래상품명이 애매한 경우
    if 0.5 <= implied <= 1.5:
        return "KRW", 1.0

    # 종목명에 통화코드가 직접 들어간 경우 우선 사용
    for currency_code, rate in fx_rate_map.items():
        if currency_code in stock_text:
            return currency_code, rate

    if fx_rate_map:
        best_currency = None
        best_rate = None
        best_score = None

        for currency_code, rate in fx_rate_map.items():
            if rate <= 0:
                continue

            score = abs(implied - rate) / rate

            if best_score is None or score < best_score:
                best_score = score
                best_currency = currency_code
                best_rate = rate

        # 환율 파일 기준과 30% 이내로 비슷하면 해당 통화 환율 사용
        if best_currency is not None and best_score is not None and best_score <= 0.30:
            return best_currency, best_rate

    # fallback:
    # exchange_rate가 없거나 통화 판별이 안 되면,
    # 삼성 파일의 평가금액으로 역산한 환율을 사용해서 계산은 계속 진행한다.
    return "IMPLIED", implied


# ============================================================
# 삼성 AP 처리
# ============================================================

def find_source_sheet(workbook):
    required = {"계좌번호", "계좌명", "종목명", "평가금액"}

    for ws in workbook.worksheets:
        headers = {
            str(cell.value).strip()
            for cell in ws[1]
            if cell.value is not None
        }

        if required.issubset(headers):
            return ws

    raise ValueError(
        "계좌번호, 계좌명, 종목명, 평가금액 헤더가 있는 시트를 찾지 못했습니다."
    )


def build_header_map(ws):
    header_map = {}

    for cell in ws[1]:
        if cell.value is not None:
            header_map[str(cell.value).strip()] = cell.column

    required = [
        "계좌번호",
        "계좌명",
        "종목명",
        "잔고수량",
        "매수단가",
        "현재가",
        "평가금액",
        "매수금액",
    ]

    missing = [header for header in required if header not in header_map]

    if missing:
        raise ValueError(f"필수 헤더를 찾지 못했습니다: {missing}")

    return header_map


def new_holding_bucket():
    return {
        "valuation": 0.0,                 # 평가금액
        "quantity": 0.0,                  # 잔고수량
        "buy_amount_krw": 0.0,            # 매수금액
        "unit_cost_weight_sum": 0.0,      # (매수단가 * 환율) * 수량
        "market_price_weight_sum": 0.0,   # (현재가 * 환율) * 수량
        "price_weight_quantity": 0.0,
        "unit_cost_simple_sum": 0.0,
        "unit_cost_simple_count": 0,
        "market_price_simple_sum": 0.0,
        "market_price_simple_count": 0,
    }


def add_holding_row(
    bucket,
    quantity,
    buy_unit_price,
    current_price,
    valuation,
    buy_amount,
    exchange_rate,
):
    if is_number(valuation):
        bucket["valuation"] += float(valuation)

    if is_number(quantity):
        bucket["quantity"] += float(quantity)

    if is_number(buy_amount):
        bucket["buy_amount_krw"] += float(buy_amount)

    unit_cost_krw = None
    market_price_krw = None

    if is_number(buy_unit_price):
        unit_cost_krw = float(buy_unit_price) * float(exchange_rate)
        bucket["unit_cost_simple_sum"] += unit_cost_krw
        bucket["unit_cost_simple_count"] += 1

    if is_number(current_price):
        market_price_krw = float(current_price) * float(exchange_rate)
        bucket["market_price_simple_sum"] += market_price_krw
        bucket["market_price_simple_count"] += 1

    if is_number(quantity) and float(quantity) != 0:
        q = float(quantity)

        if unit_cost_krw is not None:
            bucket["unit_cost_weight_sum"] += unit_cost_krw * q

        if market_price_krw is not None:
            bucket["market_price_weight_sum"] += market_price_krw * q

        bucket["price_weight_quantity"] += q


def finalize_holding(bucket):
    if bucket["price_weight_quantity"] != 0:
        acquisition_unit_krw = (
            bucket["unit_cost_weight_sum"]
            / bucket["price_weight_quantity"]
        )
        market_price_krw = (
            bucket["market_price_weight_sum"]
            / bucket["price_weight_quantity"]
        )
    else:
        acquisition_unit_krw = (
            bucket["unit_cost_simple_sum"]
            / bucket["unit_cost_simple_count"]
            if bucket["unit_cost_simple_count"]
            else None
        )
        market_price_krw = (
            bucket["market_price_simple_sum"]
            / bucket["market_price_simple_count"]
            if bucket["market_price_simple_count"]
            else None
        )

    return {
        "quantity": bucket["quantity"],
        "acquisition_unit_krw": acquisition_unit_krw,
        "buy_amount_krw": bucket["buy_amount_krw"],
        "market_price_krw": market_price_krw,
        "valuation": bucket["valuation"],
    }


def group_account_rows(ws, header_map, fx_rate_map):
    account_col = header_map["계좌번호"]
    name_col = header_map["계좌명"]
    stock_col = header_map["종목명"]
    quantity_col = header_map["잔고수량"]
    buy_unit_price_col = header_map["매수단가"]
    current_price_col = header_map["현재가"]
    valuation_col = header_map["평가금액"]
    buy_amount_col = header_map["매수금액"]
    product_col = header_map.get("고객거래상품명")

    groups = OrderedDict()

    for row_no in range(2, ws.max_row + 1):
        account_no = ws.cell(row=row_no, column=account_col).value
        account_name = ws.cell(row=row_no, column=name_col).value
        stock_name = ws.cell(row=row_no, column=stock_col).value
        quantity = ws.cell(row=row_no, column=quantity_col).value
        buy_unit_price = ws.cell(row=row_no, column=buy_unit_price_col).value
        current_price = ws.cell(row=row_no, column=current_price_col).value
        valuation = ws.cell(row=row_no, column=valuation_col).value
        buy_amount = ws.cell(row=row_no, column=buy_amount_col).value
        product_name = (
            ws.cell(row=row_no, column=product_col).value
            if product_col
            else ""
        )

        if account_no in (None, "") or account_name in (None, ""):
            continue

        if stock_name in (None, ""):
            continue

        if not is_number(valuation):
            print(
                f"[건너뜀] {row_no}행: 평가금액이 숫자가 아닙니다. "
                f"값={valuation!r}"
            )
            continue

        key = (
            str(account_no).strip(),
            str(account_name).strip(),
        )

        if key not in groups:
            groups[key] = OrderedDict()

        normalized_stock_name = str(stock_name).strip()

        if normalized_stock_name not in groups[key]:
            groups[key][normalized_stock_name] = new_holding_bucket()

        currency_code, exchange_rate = infer_currency_and_rate(
            stock_name=normalized_stock_name,
            product_name=product_name,
            quantity=quantity,
            current_price=current_price,
            valuation=valuation,
            fx_rate_map=fx_rate_map,
        )

        add_holding_row(
            bucket=groups[key][normalized_stock_name],
            quantity=quantity,
            buy_unit_price=buy_unit_price,
            current_price=current_price,
            valuation=valuation,
            buy_amount=buy_amount,
            exchange_rate=exchange_rate,
        )

    return groups


def safe_sheet_name(raw_name, used_names):
    cleaned = INVALID_SHEET_CHARS.sub("_", str(raw_name))
    cleaned = cleaned.strip().strip("'")

    if not cleaned:
        cleaned = "계좌"

    cleaned = cleaned[:31]
    candidate = cleaned
    suffix_no = 2

    while candidate in used_names:
        suffix = f"_{suffix_no}"
        candidate = f"{cleaned[:31-len(suffix)]}{suffix}"
        suffix_no += 1

    used_names.add(candidate)
    return candidate


def build_sheet_names(groups):
    name_counts = Counter(
        account_name
        for _, account_name in groups.keys()
    )
    name_running = defaultdict(int)
    used_names = {SUMMARY_SHEET_NAME}
    result = OrderedDict()

    for account_no, account_name in groups.keys():
        if name_counts[account_name] > 1:
            name_running[account_name] += 1
            raw_name = f"{account_name}{name_running[account_name]}"
        else:
            raw_name = account_name

        result[(account_no, account_name)] = safe_sheet_name(
            raw_name,
            used_names,
        )

    return result


def build_account_rows(groups):
    """
    개별 고객 시트:
    - 현금과 예수금도 그대로 포함
    - 투자비중 계산의 분모에도 포함

    10%초과_모음 시트:
    - 투자비중이 10%를 초과하는 종목만 포함
    - 종목명에 현금, 예수금, CASH가 들어간 행은 제외
    """
    account_rows = OrderedDict()
    summary_rows = []
    excluded_cash_count = 0

    for key, stock_buckets in groups.items():
        account_no, account_name = key
        total_account_value = sum(
            bucket["valuation"]
            for bucket in stock_buckets.values()
        )

        if total_account_value == 0:
            raise ValueError(
                f"{account_name}({account_no}) 계좌의 "
                "총 평가금액이 0이라 투자비중을 계산할 수 없습니다."
            )

        rows = []

        for stock_name, bucket in stock_buckets.items():
            final = finalize_holding(bucket)
            weight = final["valuation"] / total_account_value

            rows.append([
                "",
                account_no,
                account_name,
                stock_name,
                weight,
                "",
                final["quantity"],                 # 보유수량 = 잔고수량
                final["acquisition_unit_krw"],     # 취득단가 = 매수단가 * 환율
                final["buy_amount_krw"],           # 매입액원화 = 매수금액
                final["market_price_krw"],         # 시가 = 현재가 * 환율
                final["valuation"],                # 보유잔액 = 평가금액
            ])

        rows.sort(key=lambda row: row[4], reverse=True)
        account_rows[key] = rows

        for row in rows:
            if row[4] <= 0.10:
                continue

            if is_cash_or_deposit(row[3]):
                excluded_cash_count += 1
                continue

            summary_rows.append(row)

    return account_rows, summary_rows, excluded_cash_count


# ============================================================
# 결과 파일 작성
# ============================================================

def write_result_sheet(ws, rows):
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")
    border = Border(
        left=thin,
        right=thin,
        top=thin,
        bottom=thin,
    )

    for column_no, header in enumerate(OUTPUT_HEADERS, start=1):
        cell = ws.cell(row=1, column=column_no, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    for row_no, row_values in enumerate(rows, start=2):
        for column_no, value in enumerate(row_values, start=1):
            cell = ws.cell(
                row=row_no,
                column=column_no,
                value=value,
            )
            cell.border = border
            cell.alignment = Alignment(
                vertical="center",
                wrap_text=column_no in {4, 6},
            )

        ws.cell(row=row_no, column=5).number_format = "0.00%"          # 투자비중
        ws.cell(row=row_no, column=7).number_format = "#,##0.########" # 보유수량
        ws.cell(row=row_no, column=8).number_format = "#,##0.00"       # 취득단가
        ws.cell(row=row_no, column=9).number_format = "#,##0"          # 매입액원화
        ws.cell(row=row_no, column=10).number_format = "#,##0.00"      # 시가
        ws.cell(row=row_no, column=11).number_format = "#,##0"         # 보유잔액

    last_row = max(1, len(rows) + 1)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:K{last_row}"
    ws.sheet_view.showGridLines = False

    column_widths = {
        "A": 14,
        "B": 20,
        "C": 16,
        "D": 38,
        "E": 14,
        "F": 34,
        "G": 14,
        "H": 18,
        "I": 18,
        "J": 22,
        "K": 18,
    }

    for column_letter, width in column_widths.items():
        ws.column_dimensions[column_letter].width = width

    ws.row_dimensions[1].height = 36


def create_output(groups, output_path):
    output_wb = Workbook()
    output_wb.remove(output_wb.active)

    sheet_names = build_sheet_names(groups)
    account_rows, summary_rows, excluded_cash_count = build_account_rows(groups)

    # 맨 앞의 최종 요약 시트
    summary_ws = output_wb.create_sheet(SUMMARY_SHEET_NAME)
    write_result_sheet(summary_ws, summary_rows)

    # 고객별 시트는 기존과 동일
    for key, rows in account_rows.items():
        ws = output_wb.create_sheet(sheet_names[key])
        write_result_sheet(ws, rows)

    temporary_path = output_path.with_name(
        f".{output_path.stem}_작성중.xlsx"
    )

    if temporary_path.exists():
        temporary_path.unlink()

    try:
        output_wb.save(temporary_path)
        output_wb.close()

        # 파일이 실제로 열리는지 검증
        verification_wb = load_workbook(
            temporary_path,
            data_only=False,
            read_only=True,
        )
        verification_wb.close()

        try:
            os.replace(temporary_path, output_path)

        except PermissionError as error:
            raise PermissionError(
                f"'{output_path.name}' 파일이 Excel에서 열려 있습니다. "
                "결과 파일을 완전히 닫고 다시 실행해 주세요."
            ) from error

    except Exception:
        output_wb.close()

        if temporary_path.exists():
            try:
                temporary_path.unlink()
            except OSError:
                pass

        raise

    return (
        len(account_rows),
        len(summary_rows),
        excluded_cash_count,
    )


# ============================================================
# 메인 실행
# ============================================================

def main():
    base_dir = Path(__file__).resolve().parent

    # 사용법
    # python samsung_ap_normalizer_openpyxl_extended.py
    # python samsung_ap_normalizer_openpyxl_extended.py "삼성_AP 및 잔고.xlsx"
    # python samsung_ap_normalizer_openpyxl_extended.py "원본.xlsx" "결과.xlsx"

    if len(sys.argv) >= 2:
        requested_input = Path(sys.argv[1])
        input_path = (
            requested_input
            if requested_input.is_absolute()
            else base_dir / requested_input
        )
    else:
        input_path = find_input_file(base_dir)

    if not input_path.exists():
        raise FileNotFoundError(
            f"입력 파일을 찾지 못했습니다: {input_path}"
        )

    if len(sys.argv) >= 3:
        requested_output = Path(sys.argv[2])
        output_path = (
            requested_output
            if requested_output.is_absolute()
            else base_dir / requested_output
        )
    else:
        output_path = base_dir / OUTPUT_FILE_NAME

    exchange_dir = find_exchange_dir(base_dir)
    fx_tables = load_exchange_rates(exchange_dir)
    fx_rate_map = build_fx_rate_map(fx_tables)

    source_wb = load_workbook(
        input_path,
        data_only=True,
        read_only=False,
    )

    try:
        source_ws = find_source_sheet(source_wb)
        header_map = build_header_map(source_ws)
        groups = group_account_rows(source_ws, header_map, fx_rate_map)

    finally:
        source_wb.close()

    if not groups:
        raise ValueError("처리할 계좌 데이터가 없습니다.")

    (
        account_count,
        summary_count,
        excluded_cash_count,
    ) = create_output(groups, output_path)

    print(f"원본 파일: {input_path.name}")
    print(f"환율 폴더: {exchange_dir if exchange_dir else '찾지 못함'}")
    print(f"사용 가능 환율: {', '.join(sorted(fx_rate_map.keys())) if fx_rate_map else '없음'}")
    print(f"원본 시트: {source_ws.title}")
    print(f"생성 계좌 시트 수: {account_count}개")
    print(f"10% 초과 통합 행 수: {summary_count}개")
    print(f"요약 시트에서 제외한 현금·예수금 행 수: {excluded_cash_count}개")
    print(f"완료 파일: {output_path}")


if __name__ == "__main__":
    main()