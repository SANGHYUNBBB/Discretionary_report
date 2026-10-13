"""계약 이력 생성부터 취합까지 기간·상품 선택 회귀 검사."""
from datetime import date
from pathlib import Path
import json
import sys
import tempfile
import unittest
from unittest.mock import patch
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "Scripts"))
import contract_report_list as contracts
import report_paths as paths
import SumAP as ap
import SumTradeList as trades


def day(text):
    return date.fromisoformat(text)


def row(start, end, termination=None, status="신규", source_row=2):
    return contracts.ContractRow(
        "상품", "테스트", "123-456", "123456", day(start), day(end), status,
        day(termination) if termination else None, "상품", source_row,
    )


def periods(rows):
    records, notes = contracts.build_output_records(
        {("상품", "123456"): rows}, day("2026-07-01"), day("2026-09-30"))
    return [(r.report_start_date.isoformat(), r.report_end_date.isoformat()) for r in records], notes


class ContractHistoryTests(unittest.TestCase):
    def test_new_and_terminated_within_quarter(self):
        actual, _ = periods([row("2026-08-05", "2027-08-04", "2026-09-12")])
        self.assertEqual(actual, [("2026-08-05", "2026-09-12")])

    def test_start_and_end_dates_are_inclusive(self):
        for start, end, termination, expected in [
            ("2025-07-01", "2026-12-31", "2026-07-01", [("2026-07-01", "2026-07-01")]),
            ("2026-09-30", "2027-09-29", None, [("2026-09-30", "2026-09-30")]),
            ("2025-10-01", "2026-09-30", "2026-09-30", [("2026-07-01", "2026-09-30")]),
            ("2025-07-01", "2026-12-31", "2026-06-30", []),
            ("2026-10-01", "2027-09-30", None, []),
        ]:
            with self.subTest(start=start, termination=termination):
                self.assertEqual(periods([row(start,end,termination)])[0], expected)

    def test_rejoin_marked_renewal_preserves_termination_gap(self):
        actual, notes = periods([
            row("2025-08-05", "2026-08-04", "2026-07-15"),
            row("2026-08-05", "2027-08-04", status="연장", source_row=3),
        ])
        self.assertEqual(actual, [("2026-07-01", "2026-07-15"), ("2026-08-05", "2026-09-30")])
        self.assertTrue(any("생략" in note for note in notes))

    def test_maturity_gap_is_not_filled(self):
        actual, _ = periods([
            row("2025-07-16", "2026-07-15"),
            row("2026-08-05", "2027-08-04", status="연장", source_row=3),
        ])
        self.assertEqual(actual, [("2026-07-01", "2026-07-15"), ("2026-08-05", "2026-09-30")])

    def test_continuous_renewal_merges(self):
        actual, _ = periods([
            row("2025-08-05", "2026-08-04"),
            row("2026-08-05", "2027-08-04", status="연장", source_row=3),
        ])
        self.assertEqual(actual, [("2026-07-01", "2026-09-30")])

    def test_missing_previous_row_can_still_be_inferred(self):
        actual, notes = periods([row("2026-07-11", "2027-07-10", status="연장")])
        self.assertEqual(actual, [("2026-07-01", "2026-09-30")])
        self.assertTrue(any("추정하여 포함" in note for note in notes))

    def test_future_renewal_with_past_termination_retains_end(self):
        actual, _ = periods([row("2026-10-05", "2027-10-04", "2026-08-28", "연장")])
        self.assertEqual(actual, [("2026-07-01", "2026-08-28")])


class ContractResultTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.override = patch.object(paths, "ROOT", self.root)
        self.override.start()
        self.write_settings("2026-07-01", "2026-09-30")

    def tearDown(self):
        self.override.stop()
        self.temporary.cleanup()

    def write_settings(self, start, end):
        (self.root / "settings.json").write_text(json.dumps({
            "report_start":start, "report_end":end, "ap_exchange_date":None}))

    def workbook(self, rows=(), stamp=True):
        wb = Workbook()
        ws = wb.active
        ws.title = "260930"
        ws.append(contracts.OUTPUT_HEADERS)
        for product, start, end in rows:
            ws.append([product,"테스트","123-456",day(start),day(end)])
        if stamp:
            paths.stamp_contract_period(wb,day("2026-07-01"),day("2026-09-30"))
        return wb

    def save(self, wb):
        target = self.root / "ContractList.xlsx"
        wb.save(target)
        wb.close()
        return target

    def test_period_change_narrower_or_wider_is_rejected(self):
        p = self.save(self.workbook([("상품","2026-07-01","2026-09-30")]))
        for start in ("2026-08-01", "2026-06-01"):
            self.write_settings(start,"2026-09-30")
            for loader in (trades.load_report_contracts, ap.read_active_contracts):
                with self.subTest(start=start,loader=loader.__name__), self.assertRaisesRegex(ValueError,"생성 기간"):
                    loader(p)

    def test_missing_metadata_is_rejected(self):
        wb = self.workbook(stamp=False)
        with self.assertRaisesRegex(ValueError,"생성 기간"):
            paths.contract_sheet(wb)
        wb.close()

    def test_old_four_column_schema_is_rejected(self):
        wb = self.workbook()
        wb.active.cell(1,4,"보고서작성기준일")
        with self.assertRaisesRegex(ValueError,"A~E"):
            paths.contract_sheet(wb)
        wb.close()

    def test_trade_switches_product_and_excludes_gap(self):
        p = self.save(self.workbook([
            ("이전상품","2026-07-01","2026-07-09"),
            ("새상품","2026-07-15","2026-09-30"),
        ]))
        data = trades.load_report_contracts(p)
        key = ("123456","테스트")
        self.assertEqual(trades.product_for_trade_date(key,day("2026-07-09"),data),"이전상품")
        self.assertIsNone(trades.product_for_trade_date(key,day("2026-07-10"),data))
        self.assertEqual(trades.product_for_trade_date(key,day("2026-07-15"),data),"새상품")

    def test_ap_chooses_end_date_product_regardless_of_row_order(self):
        rows = [("테크차이나","2026-07-01","2026-07-09"),("글로벌임펄스","2026-07-10","2026-09-30")]
        for ordered in (rows, list(reversed(rows))):
            p = self.save(self.workbook(ordered))
            exact, _, order, reviews = ap.read_active_contracts(p)
            self.assertEqual(exact[("123456","테스트")]["product"],"글로벌임펄스")
            self.assertEqual(len(order[0]["contract_periods"]),2)
            self.assertEqual(reviews,[])

    def test_ap_retains_ended_contract_and_flags_date_review(self):
        p = self.save(self.workbook([("상품","2026-07-01","2026-08-20")]))
        _, _, order, reviews = ap.read_active_contracts(p)
        self.assertEqual(len(order),1)
        self.assertEqual(reviews[0][0],"분기 중 종료 계약 AP 기준일 확인")

    def test_conflicting_product_on_same_date_is_rejected(self):
        p = self.save(self.workbook([
            ("이전상품","2026-07-01","2026-09-30"),
            ("새상품","2026-07-10","2026-09-30"),
        ]))
        with self.assertRaisesRegex(ValueError,"겹칩니다"):
            ap.read_active_contracts(p)
        data = trades.load_report_contracts(p)
        with self.assertRaisesRegex(ValueError,"겹칩니다"):
            trades.product_for_trade_date(("123456","테스트"),day("2026-09-30"),data)

    def test_invalid_row_period_is_rejected(self):
        p = self.save(self.workbook([("상품","2026-06-30","2026-09-30")]))
        for loader in (trades.load_report_contracts, ap.read_active_contracts):
            with self.subTest(loader=loader.__name__), self.assertRaisesRegex(ValueError,"벗어납니다"):
                loader(p)

    def test_calculation_failures_stop_final_merge(self):
        wb = Workbook()
        ws = wb.create_sheet("검토필요")
        ws.append(["메시지"])
        ws.append(["계산 실패(매수): 거래일 이전 환율이 없습니다."])
        p = self.save(wb)
        with self.assertRaisesRegex(ValueError,"계산 실패 1건"):
            paths.assert_trade_calculations_valid(p)

    def test_unspecified_rule_remains_reviewable(self):
        wb = Workbook()
        wb.create_sheet("검토필요").append(["규칙 미지정 거래종류: 테스트"])
        paths.assert_trade_calculations_valid(self.save(wb))


if __name__ == "__main__":
    unittest.main()
