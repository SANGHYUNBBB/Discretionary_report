"""실행: Report 폴더에서 python -m unittest discover -s Tests -v"""
from pathlib import Path
from decimal import Decimal as D
from datetime import date, datetime
import importlib
import json
import sys
import tempfile
import unittest
from unittest.mock import patch
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "Scripts"))
import report_paths as paths
from tax_refunds import is_tax_refund, negative_refund


class RefundTests(unittest.TestCase):
    def test_broker_currency_and_sign(self):
        for prefix in ("han", "kb", "samsung", "kyobo", "meritz"):
            suffix = "_v2" if prefix == "meritz" else ""
            m = importlib.import_module(prefix + "_trade_normalizer" + suffix)
            tables = {c: m.ExchangeTable([date(2026, 4, 1).toordinal()], [v])
                      for c, v in (("USD", D(1500)), ("JPY", D(950)))}
            for currency, rate in (("KRW", D(1)), ("USD", D(1500)), ("JPY", D('9.5'))):
                for sign in (1, -1):
                    with self.subTest(broker=prefix, currency=currency, sign=sign):
                        amount = D(10 * sign)
                        row = {"거래금액": amount, "거래일자": "2026-04-02", "거래일": "2026-04-02",
                               "통화구분": currency, "통화코드": currency}
                        if prefix == "han":
                            row["구분기준거래종류"] = "현지세금환급"
                            result = m.calculate_row(row, tables)
                        elif prefix == "kb":
                            row.update(거래종류="배당세금환급 입금", 외화정산금액=amount)
                            result = m.calculate_row(row, None, tables)
                        elif prefix == "samsung":
                            row.update(거래명="외화배당세금환급", 거래수량=amount)
                            result = m.calculate_row(row, tables)
                        elif prefix == "kyobo":
                            row["적요명"] = "외화배당세금환급입금"
                            result, _ = m.calculate_row(row, tables)
                        else:
                            row["거래종류"] = "배당세금환급"
                            result, _ = m.calculate_row(row, tables)
                        self.assertEqual(result[2], 0)
                        self.assertEqual(result[4], -abs(amount * rate))

    def test_jpy_original_rate_priority(self):
        import han_trade_normalizer as han
        import kb_trade_normalizer as kb
        self.assertEqual(han.calculate_row({"구분기준거래종류": "현지세금환급", "거래금액": 10,
                                            "통화코드": "JPY", "환율": 900}, {})[4], D(-90))
        self.assertEqual(kb.calculate_row({"거래종류": "배당세금환급입금", "외화정산금액": 10,
                                          "통화구분": "JPY", "환율": 9}, None, {})[4], D(-90))

    def test_not_a_tax_refund(self):
        for text in ("현지세금재징수", "배당세금추징출금", "외화배당세금환급취소", "수수료환급", "보증금환급"):
            self.assertFalse(is_tax_refund(text), text)
        with self.assertRaises(ValueError):
            negative_refund(D(0))


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.override = patch.object(paths, "ROOT", self.root)
        self.override.start()
        (self.root / "settings.json").write_text(json.dumps({"report_start": "2026-04-01", "report_end": "2026-06-30", "ap_exchange_date": None}))

    def tearDown(self):
        self.override.stop()
        self.temporary.cleanup()

    def test_current_input_only_and_duplicate_stem(self):
        p = paths.trade_input_dir("KB")
        p.mkdir(parents=True)
        (p / "새이름26q3.xlsx").touch()
        (p / "~$잠금.xlsx").touch()
        (p / "과거_정리.xlsx").touch()
        self.assertEqual([f.name for f in paths.trade_files("KB")], ["새이름26q3.xlsx"])
        (p / "새이름26q3.xlsm").touch()
        with self.assertRaises(ValueError):
            paths.trade_files("KB")

    def test_select_only_report_sheet(self):
        wb = Workbook()
        wb.active.title = "260331"
        for name in ("260331", "260630"):
            ws = wb[name] if name in wb.sheetnames else wb.create_sheet(name)
            ws.append(["상품명", "이름", "계좌번호", "보고서작성시작일", "보고서작성마감일"])
        paths.stamp_contract_period(wb, date(2026,4,1), date(2026,6,30))
        self.assertEqual(paths.contract_sheet(wb).title, "260630")
        wb.close()

    def test_trade_period_is_inclusive(self):
        import SumTradeList as trade_sum
        wb = Workbook()
        ws = wb.active
        ws.append(["계좌번호", "계약자명", "PF명", "구분", "종목명", "매매일자", "수량", "매매단가", "매매금액", "위탁매매수수료", "각종세금"])
        for day in (datetime(2026,3,31), datetime(2026,4,1), datetime(2026,6,30), datetime(2026,7,1)):
            ws.append(["123456", "테스트", None, "현지세금환급", "예시", day, 0, 0, 0, 0, -10])
        p = self.root / "sample.xlsx"
        wb.save(p)
        wb.close()
        _, rows, _ = trade_sum.read_filtered_rows(
            p,
            {("123456", "테스트"): [("상품", date(2026,4,1), date(2026,6,30))]},
        )
        self.assertEqual(len(rows), 2)
        self.assertEqual([r[5].date() for r in rows], [date(2026,4,1), date(2026,6,30)])
        self.assertTrue(all(r[2] == "상품" and r[10] == -10 for r in rows))

    def test_trade_uses_customer_specific_end_date(self):
        import SumTradeList as trade_sum
        wb = Workbook()
        ws = wb.active
        ws.append(["계좌번호", "계약자명", "PF명", "구분", "종목명", "매매일자", "수량", "매매단가", "매매금액", "위탁매매수수료", "각종세금"])
        for day in (datetime(2026,7,1), datetime(2026,8,20), datetime(2026,8,21), datetime(2026,9,30)):
            ws.append(["123456", "테스트", None, "매수", "예시", day, 1, 100, 100, 0, 0])
        p = self.root / "customer_end.xlsx"
        wb.save(p)
        wb.close()

        _, rows, _ = trade_sum.read_filtered_rows(
            p,
            {("123456", "테스트"): [("상품", date(2026,7,1), date(2026,8,20))]},
        )
        self.assertEqual([r[5].date() for r in rows], [date(2026,7,1), date(2026,8,20)])

    def test_trade_supports_rejoin_windows(self):
        import SumTradeList as trade_sum
        wb = Workbook()
        ws = wb.active
        ws.append(["계좌번호", "계약자명", "PF명", "구분", "종목명", "매매일자", "수량", "매매단가", "매매금액", "위탁매매수수료", "각종세금"])
        for day in (datetime(2026,7,10), datetime(2026,7,20), datetime(2026,8,5), datetime(2026,8,20)):
            ws.append(["123456", "테스트", None, "매수", "예시", day, 1, 100, 100, 0, 0])
        p = self.root / "rejoin.xlsx"
        wb.save(p)
        wb.close()

        _, rows, _ = trade_sum.read_filtered_rows(
            p,
            {
                ("123456", "테스트"): [
                    ("상품", date(2026,7,1), date(2026,7,15)),
                    ("상품", date(2026,8,1), date(2026,9,30)),
                ]
            },
        )
        self.assertEqual([r[5].date() for r in rows], [date(2026,7,10), date(2026,8,5), date(2026,8,20)])


if __name__ == "__main__":
    unittest.main()
