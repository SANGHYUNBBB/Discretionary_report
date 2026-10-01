from __future__ import annotations

import argparse
import os
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


# 이 파이썬 파일과 같은 폴더에 있어야 하는 원본 파일
CONTRACT_FILE_NAME = "ContractList.xlsx"

# 계약 자료를 읽을 시트 순서
TARGET_SHEETS = [
    "테크차이나",
    "글로벌임펄스",
    "일본온앤오프",
    "미국온앤오프",
]

# ContractList 각 열의 위치
COL_PRODUCT = 1       # A열: 상품명
COL_NAME = 3          # C열: 고객명
COL_ACCOUNT = 6       # F열: 계좌번호
COL_CONTRACT_DATE = 7 # G열: 계약일
COL_MATURITY_DATE = 8 # H열: 만기일
COL_STATUS = 11       # K열: 계약상태
COL_TERMINATION = 12  # L열: 해지일 또는 운용중단일

OUTPUT_HEADERS = ["상품명", "이름", "계좌번호", "보고서작성기준일"]


@dataclass
class ContractRow:
    product: str
    name: str
    account_no: str
    account_key: str
    contract_date: date
    maturity_date: date
    status: str
    termination_date: date | None
    source_sheet: str
    source_row: int


@dataclass
class ContractEpisode:
    """재가입 전까지 이어지는 하나의 연속 계약 구간."""

    rows: list[ContractRow] = field(default_factory=list)
    start_date: date | None = None
    scheduled_end_date: date | None = None
    termination_date: date | None = None

    @property
    def actual_end_date(self) -> date | None:
        # 해지일/중단일이 있으면 그 날을 실제 종료일로 본다.
        return self.termination_date or self.scheduled_end_date

    def add(self, row: ContractRow) -> None:
        self.rows.append(row)

        if self.start_date is None or row.contract_date < self.start_date:
            self.start_date = row.contract_date

        if (
            self.scheduled_end_date is None
            or row.maturity_date > self.scheduled_end_date
        ):
            self.scheduled_end_date = row.maturity_date

        if row.termination_date is not None:
            if (
                self.termination_date is None
                or row.termination_date < self.termination_date
            ):
                self.termination_date = row.termination_date


@dataclass
class OutputRecord:
    product: str
    name: str
    account_no: str
    report_base_date: date
    source_type: str  # 실제 계약구간 / 연장추정


# ---------------------------------------------------------------------------
# 값 정리 함수
# ---------------------------------------------------------------------------

def clean_text(value: Any) -> str:
    if value is None:
        return ""

    if isinstance(value, float) and value.is_integer():
        return str(int(value))

    return str(value).strip()


def normalize_account(value: Any) -> str:
    """계좌번호 비교용으로 숫자만 남긴다."""
    return re.sub(r"\D", "", clean_text(value))


def parse_input_date(text: str) -> date:
    try:
        return datetime.strptime(text, "%Y-%m-%d").date()
    except ValueError as error:
        raise argparse.ArgumentTypeError(
            f"날짜는 YYYY-MM-DD 형식으로 입력해 주세요: {text}"
        ) from error


def to_date(value: Any) -> date | None:
    """
    엑셀 날짜, datetime, date, 문자열, 엑셀 일련번호를 date로 변환한다.
    """
    if value in (None, ""):
        return None

    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    # 엑셀 1900 날짜 시스템 기준 일련번호
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return date(1899, 12, 30) + timedelta(days=int(value))

    text = clean_text(value)
    if not text:
        return None

    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y%m%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass

    return None


def one_year_before(value: date) -> date:
    """연장 계약의 직전 연도 시작일 추정용. 윤년 2월 29일도 처리한다."""
    try:
        return value.replace(year=value.year - 1)
    except ValueError:
        return value.replace(year=value.year - 1, day=28)


# ---------------------------------------------------------------------------
# 계약 자료 읽기
# ---------------------------------------------------------------------------

def read_contract_rows(value_workbook) -> tuple[
    dict[tuple[str, str], list[ContractRow]],
    list[str],
]:
    """
    네 개 상품 시트를 읽어서 상품명 + 계좌번호 기준으로 묶는다.

    반환값:
    - 계약 행 그룹
    - 검토가 필요한 경고 문구
    """
    grouped_rows: dict[tuple[str, str], list[ContractRow]] = defaultdict(list)
    warnings: list[str] = []

    for sheet_name in TARGET_SHEETS:
        if sheet_name not in value_workbook.sheetnames:
            warnings.append(f"시트 없음: {sheet_name}")
            continue

        ws = value_workbook[sheet_name]

        for row_no in range(2, ws.max_row + 1):
            product = clean_text(ws.cell(row_no, COL_PRODUCT).value) or sheet_name
            name = clean_text(ws.cell(row_no, COL_NAME).value)
            account_no = clean_text(ws.cell(row_no, COL_ACCOUNT).value)
            account_key = normalize_account(account_no)
            contract_date = to_date(ws.cell(row_no, COL_CONTRACT_DATE).value)
            maturity_date = to_date(ws.cell(row_no, COL_MATURITY_DATE).value)
            status = clean_text(ws.cell(row_no, COL_STATUS).value)
            termination_date = to_date(ws.cell(row_no, COL_TERMINATION).value)

            # 완전히 빈 행은 건너뛴다.
            if not any((product, name, account_no, contract_date, maturity_date, status)):
                continue

            if not name or not account_key:
                warnings.append(
                    f"[{sheet_name} {row_no}행] 이름 또는 계좌번호가 없어 제외"
                )
                continue

            if contract_date is None or maturity_date is None:
                warnings.append(
                    f"[{sheet_name} {row_no}행] 계약일 또는 만기일을 읽지 못해 제외"
                )
                continue

            if maturity_date < contract_date:
                warnings.append(
                    f"[{sheet_name} {row_no}행] 만기일이 계약일보다 빨라 제외"
                )
                continue

            if termination_date is not None and termination_date < contract_date:
                # 예: 미래 연장행에 과거 운용중단일이 적힌 경우.
                # 이 값은 같은 연속 계약 전체의 실제 종료일로 사용할 수 있으므로
                # 제외하지 않고 경고만 남긴다.
                warnings.append(
                    f"[{sheet_name} {row_no}행] 해지일({termination_date})이 "
                    f"계약일({contract_date})보다 빠름: 연속 계약 종료일로 반영"
                )

            grouped_rows[(sheet_name, account_key)].append(
                ContractRow(
                    product=product,
                    name=name,
                    account_no=account_no,
                    account_key=account_key,
                    contract_date=contract_date,
                    maturity_date=maturity_date,
                    status=status,
                    termination_date=termination_date,
                    source_sheet=sheet_name,
                    source_row=row_no,
                )
            )

    return grouped_rows, warnings


# ---------------------------------------------------------------------------
# 연속 계약 구간 만들기
# ---------------------------------------------------------------------------

def split_into_episodes(rows: list[ContractRow]) -> list[ContractEpisode]:
    """
    같은 상품·계좌의 행들을 연속 계약 단위로 묶는다.

    핵심 규칙:
    1. 만기 다음 날 바로 계약이 시작되면 같은 연속 계약으로 본다.
    2. 기존 계약에 해지일이 있고, 다음 계약일이 그 해지일보다 늦으면
       재가입으로 보고 새로운 계약 구간을 만든다.

    따라서 과거 해지 후 다시 가입한 고객은 과거 해지 때문에
    현재 계약이 잘못 제외되지 않는다.
    """
    sorted_rows = sorted(
        rows,
        key=lambda row: (row.contract_date, row.maturity_date, row.source_row),
    )

    episodes: list[ContractEpisode] = []
    current: ContractEpisode | None = None

    for row in sorted_rows:
        if current is None:
            current = ContractEpisode()
            current.add(row)
            continue

        assert current.scheduled_end_date is not None

        # 명시적인 해지 이후 다시 시작한 계약이면 재가입으로 분리한다.
        is_rejoined_after_termination = (
            current.termination_date is not None
            and row.contract_date > current.termination_date
        )

        # 만기 다음 날까지 이어지는 계약은 같은 연속 계약이다.
        is_continuous = (
            row.contract_date
            <= current.scheduled_end_date + timedelta(days=1)
        )

        if is_rejoined_after_termination or not is_continuous:
            episodes.append(current)
            current = ContractEpisode()

        current.add(row)

    if current is not None:
        episodes.append(current)

    return episodes


def find_active_episode(
    episodes: list[ContractEpisode],
    report_end: date,
) -> ContractEpisode | None:
    """보고서 종료일에 실제로 살아 있는 계약 구간을 찾는다."""
    for episode in episodes:
        if episode.start_date is None or episode.actual_end_date is None:
            continue

        # 해지일이 보고서 종료일과 같으면 이미 해지된 것으로 본다.
        if episode.termination_date is not None:
            if episode.start_date <= report_end < episode.termination_date:
                return episode
        elif episode.start_date <= report_end <= episode.scheduled_end_date:
            return episode

    return None


def infer_previous_period_from_renewal(
    rows: list[ContractRow],
    report_end: date,
) -> tuple[date, ContractRow] | None:
    """
    과거 계약행이 누락되어 있고 연장행만 남은 경우의 보완 로직.

    예:
    - 연장 계약일: 2026-07-11
    - 보고서 종료일: 2026-06-30
    - 직전 계약을 2025-07-11 ~ 2026-07-10으로 추정

    실제 과거 계약행이 보고서 종료일을 덮고 있으면 이 함수는 사용되지 않는다.
    """
    renewal_rows = sorted(
        (row for row in rows if row.status == "연장"),
        key=lambda row: row.contract_date,
    )

    for row in renewal_rows:
        if row.contract_date <= report_end:
            continue

        inferred_start = one_year_before(row.contract_date)
        inferred_end = row.contract_date - timedelta(days=1)

        if not (inferred_start <= report_end <= inferred_end):
            continue

        # 해지/중단일이 보고서 종료일까지 도래했다면 대상이 아니다.
        if row.termination_date is not None and row.termination_date <= report_end:
            continue

        return inferred_start, row

    return None


# ---------------------------------------------------------------------------
# 출력 대상 만들기
# ---------------------------------------------------------------------------

def build_output_records(
    grouped_rows: dict[tuple[str, str], list[ContractRow]],
    report_start: date,
    report_end: date,
) -> tuple[list[OutputRecord], list[str]]:
    records: list[OutputRecord] = []
    inferred_messages: list[str] = []

    for (_sheet_name, _account_key), rows in grouped_rows.items():
        episodes = split_into_episodes(rows)
        active_episode = find_active_episode(episodes, report_end)

        if active_episode is not None:
            assert active_episode.start_date is not None

            # 보고서 시작 전에 가입한 고객은 입력한 기간 시작일이 기준일,
            # 기간 중 신규 가입자는 실제 계약일이 기준일이다.
            report_base_date = max(report_start, active_episode.start_date)

            # 이름과 계좌번호는 해당 연속 계약의 가장 최근 행에서 가져온다.
            display_row = max(
                active_episode.rows,
                key=lambda row: (row.contract_date, row.source_row),
            )

            records.append(
                OutputRecord(
                    product=display_row.product,
                    name=display_row.name,
                    account_no=display_row.account_no,
                    report_base_date=report_base_date,
                    source_type="실제 계약구간",
                )
            )
            continue

        # 연장행만 있고 직전 계약행이 누락된 경우를 위한 보완 처리
        inferred = infer_previous_period_from_renewal(rows, report_end)
        if inferred is None:
            continue

        inferred_start, renewal_row = inferred
        records.append(
            OutputRecord(
                product=renewal_row.product,
                name=renewal_row.name,
                account_no=renewal_row.account_no,
                report_base_date=max(report_start, inferred_start),
                source_type="연장추정",
            )
        )
        inferred_messages.append(
            f"[{renewal_row.source_sheet} {renewal_row.source_row}행] "
            f"{renewal_row.name} / {renewal_row.account_no}: "
            f"직전 계약기간을 추정하여 포함"
        )

    product_order = {name: index for index, name in enumerate(TARGET_SHEETS)}

    records.sort(
        key=lambda record: (
            product_order.get(record.product, len(product_order)),
            record.report_base_date,
            record.name,
            normalize_account(record.account_no),
        )
    )

    return records, inferred_messages


# ---------------------------------------------------------------------------
# 엑셀 시트 작성
# ---------------------------------------------------------------------------

def write_output_sheet(workbook, sheet_name: str, records: list[OutputRecord]) -> None:
    # 같은 이름의 기존 시트가 있으면 같은 위치에 새로 만든다.
    if sheet_name in workbook.sheetnames:
        old_sheet = workbook[sheet_name]
        old_index = workbook.sheetnames.index(sheet_name)
        workbook.remove(old_sheet)
        ws = workbook.create_sheet(sheet_name, old_index)
    else:
        ws = workbook.create_sheet(sheet_name, 0)

    ws.append(OUTPUT_HEADERS)

    for record in records:
        ws.append(
            [
                record.product,
                record.name,
                record.account_no,
                record.report_base_date,
            ]
        )

    # 서식
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin_side = Side(style="thin", color="D9E2F3")
    border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side,
    )

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    for row in ws.iter_rows(min_row=2, max_row=max(ws.max_row, 2), min_col=1, max_col=4):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center")

    for row_no in range(2, ws.max_row + 1):
        ws.cell(row_no, 3).number_format = "@"          # 계좌번호는 텍스트
        ws.cell(row_no, 4).number_format = "yyyy-mm-dd" # 보고서작성기준일

    widths = {1: 20, 2: 16, 3: 22, 4: 22}
    for column_no, width in widths.items():
        ws.column_dimensions[get_column_letter(column_no)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:D{max(ws.max_row, 1)}"
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[1].height = 24


# ---------------------------------------------------------------------------
# 저장 및 실행
# ---------------------------------------------------------------------------

def save_safely(workbook, contract_path: Path) -> None:
    temporary_path = contract_path.with_name(
        f".{contract_path.stem}_작성중{contract_path.suffix}"
    )

    if temporary_path.exists():
        temporary_path.unlink()

    try:
        workbook.save(temporary_path)
        workbook.close()

        # 저장된 파일이 정상적으로 다시 열리는지 간단히 검증한다.
        verification_wb = load_workbook(temporary_path, read_only=True, data_only=False)
        verification_wb.close()

        try:
            os.replace(temporary_path, contract_path)
        except PermissionError as error:
            raise PermissionError(
                f"'{contract_path.name}' 파일이 Excel에서 열려 있습니다. "
                "파일을 완전히 닫고 다시 실행해 주세요."
            ) from error

    except Exception:
        try:
            workbook.close()
        except Exception:
            pass

        if temporary_path.exists():
            try:
                temporary_path.unlink()
            except OSError:
                pass
        raise


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "ContractList.xlsx에서 보고서 종료일 현재 살아 있는 일임계약을 "
            "추출하여 종료일 기준 YYMMDD 시트에 정리합니다."
        )
    )
    parser.add_argument(
        "report_start",
        nargs="?",
        type=parse_input_date,
        help="보고서 작성기간 시작일: YYYY-MM-DD",
    )
    parser.add_argument(
        "report_end",
        nargs="?",
        type=parse_input_date,
        help="보고서 작성기간 종료일: YYYY-MM-DD",
    )
    args = parser.parse_args()

    # VS Code의 실행 버튼처럼 인수를 주지 않고 실행한 경우 직접 입력받는다.
    if args.report_start is None and args.report_end is None:
        args.report_start = parse_input_date(
            input("보고서 작성기간 시작일을 입력하세요 (YYYY-MM-DD): ").strip()
        )
        args.report_end = parse_input_date(
            input("보고서 작성기간 종료일을 입력하세요 (YYYY-MM-DD): ").strip()
        )
    elif args.report_start is None or args.report_end is None:
        parser.error("시작일과 종료일을 모두 입력해야 합니다.")

    return args


def main() -> None:
    args = parse_args()
    report_start: date = args.report_start
    report_end: date = args.report_end

    if report_start > report_end:
        raise ValueError("보고서 시작일은 종료일보다 늦을 수 없습니다.")

    base_dir = Path(__file__).resolve().parent
    contract_path = base_dir / CONTRACT_FILE_NAME

    if not contract_path.exists():
        raise FileNotFoundError(
            f"파이썬 파일과 같은 폴더에서 '{CONTRACT_FILE_NAME}'을 찾지 못했습니다.\n"
            f"확인한 위치: {contract_path}"
        )

    # 값 읽기용: 수식이 있으면 계산된 값을 읽는다.
    value_workbook = load_workbook(contract_path, data_only=True, read_only=True)

    # 저장용: 기존 수식과 서식을 유지한다.
    write_workbook = load_workbook(contract_path, data_only=False)

    try:
        grouped_rows, warnings = read_contract_rows(value_workbook)
        records, inferred_messages = build_output_records(
            grouped_rows,
            report_start,
            report_end,
        )

        output_sheet_name = report_end.strftime("%y%m%d")
        write_output_sheet(write_workbook, output_sheet_name, records)

    except Exception:
        write_workbook.close()
        raise
    finally:
        value_workbook.close()

    save_safely(write_workbook, contract_path)

    # 실행 결과 출력
    print(f"원본 파일: {contract_path}")
    print(f"보고서 기간: {report_start} ~ {report_end}")
    print(f"작성 시트: {output_sheet_name}")
    print(f"총 작성 대상: {len(records)}명")

    product_counts: dict[str, int] = defaultdict(int)
    for record in records:
        product_counts[record.product] += 1

    for product in TARGET_SHEETS:
        print(f"  - {product}: {product_counts.get(product, 0)}명")

    if inferred_messages:
        print("\n[연장 계약 추정 포함]")
        for message in inferred_messages:
            print(f"  - {message}")

    if warnings:
        print("\n[검토 필요]")
        for warning in warnings:
            print(f"  - {warning}")


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print(f"오류: {error}", file=sys.stderr)
        sys.exit(1)