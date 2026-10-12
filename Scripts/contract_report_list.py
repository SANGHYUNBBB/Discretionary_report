from __future__ import annotations

import argparse
import os
import re
import sys

from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

from report_paths import contract_input, contract_output, settings, update_period

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


# ---------------------------------------------------------------------------
# 기본 설정
# ---------------------------------------------------------------------------

# Input/계약에서 읽고 Output/계약에 별도로 저장한다.
CONTRACT_FILE_NAME = "ContractList.xlsx"


# 계약 자료를 읽을 시트 순서
TARGET_SHEETS = [
    "테크차이나",
    "글로벌임펄스",
    "일본온앤오프",
    "미국온앤오프",
]


# ContractList 각 열의 위치
COL_PRODUCT = 1        # A열: 상품명
COL_NAME = 3           # C열: 고객명
COL_ACCOUNT = 6        # F열: 계좌번호
COL_CONTRACT_DATE = 7  # G열: 계약일
COL_MATURITY_DATE = 8  # H열: 만기일
COL_STATUS = 11        # K열: 계약상태
COL_TERMINATION = 12   # L열: 해지일 또는 운용중단일


# 출력 시트
OUTPUT_HEADERS = [
    "상품명",
    "이름",
    "계좌번호",
    "보고서작성시작일",
    "보고서작성마감일",
]


# ---------------------------------------------------------------------------
# 데이터 구조
# ---------------------------------------------------------------------------

@dataclass
class ContractRow:
    product: str
    name: str
    account_no: str
    account_key: str
    contract_date: date
    maturity_date: date
    status: str
    termination_date: date | None
    source_sheet: str
    source_row: int


@dataclass
class ContractEpisode:
    """
    재가입 전까지 이어지는 하나의 연속 계약 구간.
    """

    rows: list[ContractRow] = field(default_factory=list)

    start_date: date | None = None
    scheduled_end_date: date | None = None
    termination_date: date | None = None

    @property
    def actual_end_date(self) -> date | None:
        """
        실제 계약 종료일.

        해지일/운용중단일이 있으면 해당 날짜를 실제 종료일로 보고,
        없으면 계약상 만기일을 사용한다.
        """
        return self.termination_date or self.scheduled_end_date

    def add(self, row: ContractRow) -> None:
        self.rows.append(row)

        if self.start_date is None or row.contract_date < self.start_date:
            self.start_date = row.contract_date

        if (
            self.scheduled_end_date is None
            or row.maturity_date > self.scheduled_end_date
        ):
            self.scheduled_end_date = row.maturity_date

        if row.termination_date is not None:
            if (
                self.termination_date is None
                or row.termination_date < self.termination_date
            ):
                self.termination_date = row.termination_date


@dataclass
class OutputRecord:
    product: str
    name: str
    account_no: str

    report_start_date: date
    report_end_date: date

    source_type: str  # 실제 계약구간 / 연장추정


# ---------------------------------------------------------------------------
# 값 정리 함수
# ---------------------------------------------------------------------------

def clean_text(value: Any) -> str:
    if value is None:
        return ""

    if isinstance(value, float) and value.is_integer():
        return str(int(value))

    return str(value).strip()


def normalize_account(value: Any) -> str:
    """
    계좌번호 비교용으로 숫자만 남긴다.
    """
    return re.sub(r"\D", "", clean_text(value))


def parse_input_date(text: str) -> date:
    try:
        return datetime.strptime(text, "%Y-%m-%d").date()

    except ValueError as error:
        raise argparse.ArgumentTypeError(
            f"날짜는 YYYY-MM-DD 형식으로 입력해 주세요: {text}"
        ) from error


def to_date(value: Any) -> date | None:
    """
    엑셀 날짜, datetime, date, 문자열,
    엑셀 일련번호를 date로 변환한다.
    """

    if value in (None, ""):
        return None

    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    # 엑셀 1900 날짜 시스템 기준 일련번호
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return date(1899, 12, 30) + timedelta(days=int(value))

    text = clean_text(value)

    if not text:
        return None

    for fmt in (
        "%Y-%m-%d",
        "%Y/%m/%d",
        "%Y.%m.%d",
        "%Y%m%d",
    ):
        try:
            return datetime.strptime(text, fmt).date()

        except ValueError:
            pass

    return None


def one_year_before(value: date) -> date:
    """
    연장 계약의 직전 연도 시작일 추정용.
    윤년 2월 29일도 처리한다.
    """

    try:
        return value.replace(year=value.year - 1)

    except ValueError:
        return value.replace(
            year=value.year - 1,
            day=28,
        )


# ---------------------------------------------------------------------------
# 계약 자료 읽기
# ---------------------------------------------------------------------------

def read_contract_rows(
    value_workbook,
) -> tuple[
    dict[tuple[str, str], list[ContractRow]],
    list[str],
]:
    """
    네 개 상품 시트를 읽어서
    상품명 + 계좌번호 기준으로 묶는다.

    반환값:
    - 계약 행 그룹
    - 검토가 필요한 경고 문구
    """

    grouped_rows: dict[
        tuple[str, str],
        list[ContractRow],
    ] = defaultdict(list)

    warnings: list[str] = []

    for sheet_name in TARGET_SHEETS:

        if sheet_name not in value_workbook.sheetnames:
            warnings.append(
                f"시트 없음: {sheet_name}"
            )
            continue

        ws = value_workbook[sheet_name]

        for row_no in range(
            2,
            ws.max_row + 1,
        ):

            product = (
                clean_text(
                    ws.cell(
                        row_no,
                        COL_PRODUCT,
                    ).value
                )
                or sheet_name
            )

            name = clean_text(
                ws.cell(
                    row_no,
                    COL_NAME,
                ).value
            )

            account_no = clean_text(
                ws.cell(
                    row_no,
                    COL_ACCOUNT,
                ).value
            )

            account_key = normalize_account(
                account_no
            )

            contract_date = to_date(
                ws.cell(
                    row_no,
                    COL_CONTRACT_DATE,
                ).value
            )

            maturity_date = to_date(
                ws.cell(
                    row_no,
                    COL_MATURITY_DATE,
                ).value
            )

            status = clean_text(
                ws.cell(
                    row_no,
                    COL_STATUS,
                ).value
            )

            termination_date = to_date(
                ws.cell(
                    row_no,
                    COL_TERMINATION,
                ).value
            )

            # 완전히 빈 행은 건너뛴다.
            if not any(
                (
                    product,
                    name,
                    account_no,
                    contract_date,
                    maturity_date,
                    status,
                )
            ):
                continue

            if not name or not account_key:
                warnings.append(
                    f"[{sheet_name} {row_no}행] "
                    f"이름 또는 계좌번호가 없어 제외"
                )
                continue

            if (
                contract_date is None
                or maturity_date is None
            ):
                warnings.append(
                    f"[{sheet_name} {row_no}행] "
                    f"계약일 또는 만기일을 읽지 못해 제외"
                )
                continue

            if maturity_date < contract_date:
                warnings.append(
                    f"[{sheet_name} {row_no}행] "
                    f"만기일이 계약일보다 빨라 제외"
                )
                continue

            if (
                termination_date is not None
                and termination_date < contract_date
            ):
                # 예:
                # 미래 연장행에 과거 운용중단일이 적혀 있는 경우.
                #
                # 이 값은 같은 연속 계약 전체의
                # 실제 종료일로 사용할 수 있으므로
                # 제외하지 않고 경고만 남긴다.
                warnings.append(
                    f"[{sheet_name} {row_no}행] "
                    f"해지일({termination_date})이 "
                    f"계약일({contract_date})보다 빠름: "
                    f"연속 계약 종료일로 반영"
                )

            grouped_rows[
                (
                    sheet_name,
                    account_key,
                )
            ].append(
                ContractRow(
                    product=product,
                    name=name,
                    account_no=account_no,
                    account_key=account_key,
                    contract_date=contract_date,
                    maturity_date=maturity_date,
                    status=status,
                    termination_date=termination_date,
                    source_sheet=sheet_name,
                    source_row=row_no,
                )
            )

    return grouped_rows, warnings


# ---------------------------------------------------------------------------
# 연속 계약 구간 만들기
# ---------------------------------------------------------------------------

def split_into_episodes(
    rows: list[ContractRow],
) -> list[ContractEpisode]:
    """
    같은 상품·계좌의 행들을
    연속 계약 단위로 묶는다.

    핵심 규칙:

    1. 만기 다음 날 바로 계약이 시작되면
       같은 연속 계약으로 본다.

    2. 기존 계약에 해지일이 있고,
       다음 계약일이 그 해지일보다 늦으면
       재가입으로 보고 새로운 계약 구간을 만든다.

    따라서 과거 해지 후 다시 가입한 고객은
    과거 해지 때문에 현재 계약이 잘못 제외되지 않는다.
    """

    sorted_rows = sorted(
        rows,
        key=lambda row: (
            row.contract_date,
            row.maturity_date,
            row.source_row,
        ),
    )

    episodes: list[ContractEpisode] = []

    current: ContractEpisode | None = None

    for row in sorted_rows:

        if current is None:
            current = ContractEpisode()
            current.add(row)
            continue

        assert current.scheduled_end_date is not None

        # 명시적인 해지 이후 다시 시작한 계약이면
        # 재가입으로 분리한다.
        is_rejoined_after_termination = (
            current.termination_date is not None
            and row.contract_date
            > current.termination_date
        )

        # 만기 다음 날까지 이어지는 계약은
        # 같은 연속 계약이다.
        is_continuous = (
            row.contract_date
            <= current.scheduled_end_date
            + timedelta(days=1)
        )

        if (
            is_rejoined_after_termination
            or not is_continuous
        ):
            episodes.append(current)
            current = ContractEpisode()

        current.add(row)

    if current is not None:
        episodes.append(current)

    return episodes


# ---------------------------------------------------------------------------
# 보고서 기간과 계약기간 겹침 판단
# ---------------------------------------------------------------------------

def find_overlapping_episodes(
    episodes: list[ContractEpisode],
    report_start: date,
    report_end: date,
) -> list[ContractEpisode]:
    """
    보고서 기간과 실제 계약기간이
    하루라도 겹치는 계약구간을 찾는다.

    예:
    보고서기간
    2026-07-01 ~ 2026-09-30

    계속 운용 중
    -> 포함

    2026-08-20 해지
    -> 포함

    2026-07-15 신규가입
    -> 포함

    2026-06-30 종료
    -> 제외

    2026-10-01 신규가입
    -> 제외
    """

    overlapping: list[ContractEpisode] = []

    for episode in episodes:

        if (
            episode.start_date is None
            or episode.actual_end_date is None
        ):
            continue

        contract_start = episode.start_date
        contract_end = episode.actual_end_date

        # 두 기간이 하루라도 겹치는 조건
        if (
            contract_start <= report_end
            and contract_end >= report_start
        ):
            overlapping.append(episode)

    return overlapping


# ---------------------------------------------------------------------------
# 기존 계약행 누락 시 연장행으로 직전 계약기간 추정
# ---------------------------------------------------------------------------

def subtract_existing_episode_ranges(
    start_date: date,
    end_date: date,
    episodes: list[ContractEpisode],
) -> list[tuple[date, date]]:
    """
    start_date ~ end_date 중에서
    실제 ContractList 계약행으로 이미 확인되는 기간을 제거한다.

    반환값은 실제 계약행으로 확인되지 않는 기간들이다.

    이 함수는 '연장' 행만 존재하고
    직전 계약행이 누락된 경우의 중복 방지용이다.
    """

    remaining = [
        (
            start_date,
            end_date,
        )
    ]

    for episode in episodes:

        if (
            episode.start_date is None
            or episode.actual_end_date is None
        ):
            continue

        covered_start = episode.start_date
        covered_end = episode.actual_end_date

        new_remaining: list[
            tuple[date, date]
        ] = []

        for segment_start, segment_end in remaining:

            # 겹치지 않는 경우
            if (
                covered_end < segment_start
                or covered_start > segment_end
            ):
                new_remaining.append(
                    (
                        segment_start,
                        segment_end,
                    )
                )
                continue

            # 왼쪽에 남는 기간
            if segment_start < covered_start:
                left_end = (
                    covered_start
                    - timedelta(days=1)
                )

                if segment_start <= left_end:
                    new_remaining.append(
                        (
                            segment_start,
                            left_end,
                        )
                    )

            # 오른쪽에 남는 기간
            if covered_end < segment_end:
                right_start = (
                    covered_end
                    + timedelta(days=1)
                )

                if right_start <= segment_end:
                    new_remaining.append(
                        (
                            right_start,
                            segment_end,
                        )
                    )

        remaining = new_remaining

        if not remaining:
            break

    return remaining


def infer_missing_previous_periods_from_renewal(
    rows: list[ContractRow],
    episodes: list[ContractEpisode],
    report_start: date,
    report_end: date,
) -> list[
    tuple[
        date,
        date,
        ContractRow,
    ]
]:
    """
    연장행만 있고 직전 계약행이 누락된 경우
    직전 계약기간을 추정한다.

    예:
    연장 계약일: 2026-07-11

    직전 계약기간 추정:
    2025-07-11 ~ 2026-07-10

    보고서기간:
    2026-07-01 ~ 2026-09-30

    실제 기존 계약행이 7/11부터만 존재한다면
    누락된 7/1 ~ 7/10을 연장추정으로 보완한다.

    실제 계약행으로 이미 확인되는 기간은
    중복 출력하지 않는다.
    """

    inferred_periods: list[
        tuple[
            date,
            date,
            ContractRow,
        ]
    ] = []

    renewal_rows = sorted(
        (
            row
            for row in rows
            if row.status == "연장"
        ),
        key=lambda row: row.contract_date,
    )

    for row in renewal_rows:

        inferred_start = one_year_before(
            row.contract_date
        )

        inferred_end = (
            row.contract_date
            - timedelta(days=1)
        )

        # 과거 계약기간 안에 해지/운용중단일이 있으면
        # 실제 종료일로 사용한다.
        if row.termination_date is not None:
            if (
                inferred_start
                <= row.termination_date
                <= inferred_end
            ):
                inferred_end = row.termination_date

        # 보고서 기간과 겹치는 부분만 본다.
        overlap_start = max(
            report_start,
            inferred_start,
        )

        overlap_end = min(
            report_end,
            inferred_end,
        )

        if overlap_start > overlap_end:
            continue

        # 실제 ContractList 계약행으로 확인되는 기간은 빼고
        # 정말 누락된 부분만 보완한다.
        missing_segments = (
            subtract_existing_episode_ranges(
                overlap_start,
                overlap_end,
                episodes,
            )
        )

        for (
            missing_start,
            missing_end,
        ) in missing_segments:

            inferred_periods.append(
                (
                    missing_start,
                    missing_end,
                    row,
                )
            )

    return inferred_periods


# ---------------------------------------------------------------------------
# 동일 계약의 연속 출력기간 합치기
# ---------------------------------------------------------------------------

def merge_contiguous_records(
    records: list[OutputRecord],
) -> list[OutputRecord]:
    """
    동일 상품 + 동일 계좌가
    연속된 날짜로 여러 번 만들어진 경우
    하나의 출력행으로 합친다.

    예:

    2026-07-01 ~ 2026-07-10  연장추정
    2026-07-11 ~ 2026-09-30  실제 계약구간

    ->

    2026-07-01 ~ 2026-09-30

    단, 실제 해지 후 며칠 뒤 재가입처럼
    기간 사이에 공백이 있으면 별도 행으로 유지한다.
    """

    if not records:
        return []

    records_sorted = sorted(
        records,
        key=lambda record: (
            record.product,
            normalize_account(
                record.account_no
            ),
            record.report_start_date,
            record.report_end_date,
        ),
    )

    merged: list[OutputRecord] = []

    for record in records_sorted:

        if not merged:
            merged.append(record)
            continue

        previous = merged[-1]

        same_contract = (
            previous.product == record.product
            and normalize_account(
                previous.account_no
            )
            == normalize_account(
                record.account_no
            )
        )

        is_contiguous = (
            record.report_start_date
            <= previous.report_end_date
            + timedelta(days=1)
        )

        if same_contract and is_contiguous:

            previous.report_end_date = max(
                previous.report_end_date,
                record.report_end_date,
            )

            previous.report_start_date = min(
                previous.report_start_date,
                record.report_start_date,
            )

            if (
                previous.source_type
                != record.source_type
            ):
                previous.source_type = (
                    "실제 계약구간+연장추정"
                )

            # 더 최근 이름/계좌 표현 사용
            previous.name = record.name
            previous.account_no = record.account_no

        else:
            merged.append(record)

    return merged


# ---------------------------------------------------------------------------
# 출력 대상 만들기
# ---------------------------------------------------------------------------

def build_output_records(
    grouped_rows: dict[
        tuple[str, str],
        list[ContractRow],
    ],
    report_start: date,
    report_end: date,
) -> tuple[
    list[OutputRecord],
    list[str],
]:
    """
    보고서 기간과 계약기간이 하루라도 겹치는 계약을
    모두 출력대상으로 만든다.

    보고서작성시작일:
    max(보고서 시작일, 실제 계약 시작일)

    보고서작성마감일:
    min(보고서 종료일, 실제 계약 종료일)

    따라서 예를 들어

    보고서:
    2026-07-01 ~ 2026-09-30

    계약:
    2025-01-01 ~ 계속

    -> 2026-07-01 ~ 2026-09-30


    계약:
    2026-07-20 ~ 계속

    -> 2026-07-20 ~ 2026-09-30


    계약:
    2025-01-01 ~ 2026-08-20 해지

    -> 2026-07-01 ~ 2026-08-20
    """

    records: list[OutputRecord] = []

    inferred_messages: list[str] = []

    for (
        _sheet_name,
        _account_key,
    ), rows in grouped_rows.items():

        episodes = split_into_episodes(
            rows
        )

        # ---------------------------------------------------------------
        # 1. 실제 계약구간
        # ---------------------------------------------------------------

        overlapping_episodes = (
            find_overlapping_episodes(
                episodes,
                report_start,
                report_end,
            )
        )

        for episode in overlapping_episodes:

            if (
                episode.start_date is None
                or episode.actual_end_date is None
            ):
                continue

            report_record_start = max(
                report_start,
                episode.start_date,
            )

            report_record_end = min(
                report_end,
                episode.actual_end_date,
            )

            if (
                report_record_start
                > report_record_end
            ):
                continue

            # 이름 / 계좌번호 / 상품명은
            # 해당 연속 계약의 가장 최근 행에서 가져온다.
            display_row = max(
                episode.rows,
                key=lambda row: (
                    row.contract_date,
                    row.source_row,
                ),
            )

            records.append(
                OutputRecord(
                    product=display_row.product,
                    name=display_row.name,
                    account_no=display_row.account_no,
                    report_start_date=report_record_start,
                    report_end_date=report_record_end,
                    source_type="실제 계약구간",
                )
            )

        # ---------------------------------------------------------------
        # 2. 연장행만 있고 직전 계약행이 누락된 경우 보완
        # ---------------------------------------------------------------

        inferred_periods = (
            infer_missing_previous_periods_from_renewal(
                rows,
                episodes,
                report_start,
                report_end,
            )
        )

        for (
            inferred_start,
            inferred_end,
            renewal_row,
        ) in inferred_periods:

            records.append(
                OutputRecord(
                    product=renewal_row.product,
                    name=renewal_row.name,
                    account_no=renewal_row.account_no,
                    report_start_date=inferred_start,
                    report_end_date=inferred_end,
                    source_type="연장추정",
                )
            )

            inferred_messages.append(
                f"[{renewal_row.source_sheet} "
                f"{renewal_row.source_row}행] "
                f"{renewal_row.name} / "
                f"{renewal_row.account_no}: "
                f"{inferred_start} ~ {inferred_end} "
                f"직전 계약기간을 추정하여 포함"
            )

    # 실제구간 + 연장추정이 날짜상 이어지면
    # 한 행으로 합친다.
    records = merge_contiguous_records(
        records
    )

    product_order = {
        name: index
        for index, name
        in enumerate(TARGET_SHEETS)
    }

    records.sort(
        key=lambda record: (
            product_order.get(
                record.product,
                len(product_order),
            ),
            record.report_start_date,
            record.name,
            normalize_account(
                record.account_no
            ),
        )
    )

    return records, inferred_messages


# ---------------------------------------------------------------------------
# 엑셀 시트 작성
# ---------------------------------------------------------------------------

def write_output_sheet(
    workbook,
    sheet_name: str,
    records: list[OutputRecord],
) -> None:
    """
    결과를

    상품명
    이름
    계좌번호
    보고서작성시작일
    보고서작성마감일

    형식으로 저장한다.
    """

    # 같은 이름의 기존 시트가 있으면
    # 같은 위치에 새로 만든다.
    if sheet_name in workbook.sheetnames:

        old_sheet = workbook[sheet_name]

        old_index = (
            workbook.sheetnames.index(
                sheet_name
            )
        )

        workbook.remove(
            old_sheet
        )

        ws = workbook.create_sheet(
            sheet_name,
            old_index,
        )

    else:
        ws = workbook.create_sheet(
            sheet_name,
            0,
        )

    ws.append(
        OUTPUT_HEADERS
    )

    for record in records:

        ws.append(
            [
                record.product,
                record.name,
                record.account_no,
                record.report_start_date,
                record.report_end_date,
            ]
        )

    # -------------------------------------------------------------------
    # 서식
    # -------------------------------------------------------------------

    header_fill = PatternFill(
        "solid",
        fgColor="1F4E78",
    )

    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    thin_side = Side(
        style="thin",
        color="D9E2F3",
    )

    border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side,
    )

    for cell in ws[1]:

        cell.fill = header_fill
        cell.font = header_font

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        cell.border = border

    for row in ws.iter_rows(
        min_row=2,
        max_row=max(
            ws.max_row,
            2,
        ),
        min_col=1,
        max_col=5,
    ):

        for cell in row:

            cell.border = border

            cell.alignment = Alignment(
                vertical="center"
            )

    for row_no in range(
        2,
        ws.max_row + 1,
    ):

        # 계좌번호는 텍스트
        ws.cell(
            row_no,
            3,
        ).number_format = "@"

        # 보고서작성시작일
        ws.cell(
            row_no,
            4,
        ).number_format = "yyyy-mm-dd"

        # 보고서작성마감일
        ws.cell(
            row_no,
            5,
        ).number_format = "yyyy-mm-dd"

    widths = {
        1: 20,
        2: 16,
        3: 22,
        4: 22,
        5: 22,
    }

    for (
        column_no,
        width,
    ) in widths.items():

        ws.column_dimensions[
            get_column_letter(
                column_no
            )
        ].width = width

    ws.freeze_panes = "A2"

    ws.auto_filter.ref = (
        f"A1:E{max(ws.max_row, 1)}"
    )

    ws.sheet_view.showGridLines = False

    ws.row_dimensions[1].height = 24


# ---------------------------------------------------------------------------
# 저장
# ---------------------------------------------------------------------------

def save_safely(
    workbook,
    contract_path: Path,
) -> None:

    temporary_path = (
        contract_path.with_name(
            f".{contract_path.stem}_작성중"
            f"{contract_path.suffix}"
        )
    )

    if temporary_path.exists():
        temporary_path.unlink()

    try:

        workbook.save(
            temporary_path
        )

        workbook.close()

        # 저장된 파일이 정상적으로
        # 다시 열리는지 검증
        verification_wb = load_workbook(
            temporary_path,
            read_only=True,
            data_only=False,
        )

        verification_wb.close()

        try:
            os.replace(
                temporary_path,
                contract_path,
            )

        except PermissionError as error:

            raise PermissionError(
                f"'{contract_path.name}' 파일이 "
                f"Excel에서 열려 있습니다. "
                f"파일을 완전히 닫고 다시 실행해 주세요."
            ) from error

    except Exception:

        try:
            workbook.close()

        except Exception:
            pass

        if temporary_path.exists():

            try:
                temporary_path.unlink()

            except OSError:
                pass

        raise


# ---------------------------------------------------------------------------
# 실행 인수
# ---------------------------------------------------------------------------

def parse_args() -> argparse.Namespace:

    parser = argparse.ArgumentParser(
        description=(
            "ContractList.xlsx에서 "
            "보고서 기간과 계약기간이 하루라도 겹치는 "
            "일임계약을 추출하여 "
            "종료일 기준 YYMMDD 시트에 정리합니다."
        )
    )

    parser.add_argument(
        "report_start",
        nargs="?",
        type=parse_input_date,
        help=(
            "보고서 작성기간 시작일: "
            "YYYY-MM-DD"
        ),
    )

    parser.add_argument(
        "report_end",
        nargs="?",
        type=parse_input_date,
        help=(
            "보고서 작성기간 종료일: "
            "YYYY-MM-DD"
        ),
    )

    args = parser.parse_args()

    # 인수가 없으면
    # Report/settings.json의 보고서 기간 사용
    if (
        args.report_start is None
        and args.report_end is None
    ):

        config = settings()

        args.report_start = (
            parse_input_date(
                config["report_start"]
            )
        )

        args.report_end = (
            parse_input_date(
                config["report_end"]
            )
        )

    elif (
        args.report_start is None
        or args.report_end is None
    ):

        parser.error(
            "시작일과 종료일을 모두 입력해야 합니다."
        )

    return args


# ---------------------------------------------------------------------------
# 메인
# ---------------------------------------------------------------------------

def main() -> None:

    args = parse_args()

    report_start: date = (
        args.report_start
    )

    report_end: date = (
        args.report_end
    )

    if report_start > report_end:

        raise ValueError(
            "보고서 시작일은 "
            "종료일보다 늦을 수 없습니다."
        )

    contract_path = (
        contract_input()
    )

    output_path = (
        contract_output()
    )

    if not contract_path.exists():

        raise FileNotFoundError(
            f"Input/계약에서 "
            f"'{CONTRACT_FILE_NAME}'을 "
            f"찾지 못했습니다.\n"
            f"확인한 위치: {contract_path}"
        )

    # 값 읽기용:
    # 수식이 있으면 계산된 값을 읽는다.
    value_workbook = load_workbook(
        contract_path,
        data_only=True,
        read_only=True,
    )

    # 저장용:
    # 기존 수식과 서식을 유지한다.
    write_workbook = load_workbook(
        contract_path,
        data_only=False,
    )

    try:

        grouped_rows, warnings = (
            read_contract_rows(
                value_workbook
            )
        )

        records, inferred_messages = (
            build_output_records(
                grouped_rows,
                report_start,
                report_end,
            )
        )

        output_sheet_name = (
            report_end.strftime(
                "%y%m%d"
            )
        )

        write_output_sheet(
            write_workbook,
            output_sheet_name,
            records,
        )

    except Exception:

        write_workbook.close()
        raise

    finally:

        value_workbook.close()

    write_workbook.active = (
        write_workbook.sheetnames.index(
            output_sheet_name
        )
    )

    save_safely(
        write_workbook,
        output_path,
    )

    update_period(
        report_start,
        report_end,
    )

    # -------------------------------------------------------------------
    # 실행 결과 출력
    # -------------------------------------------------------------------

    print(
        f"결과 파일: {output_path}"
    )

    print(
        f"원본 파일: {contract_path}"
    )

    print(
        f"보고서 기간: "
        f"{report_start} ~ {report_end}"
    )

    print(
        f"작성 시트: "
        f"{output_sheet_name}"
    )

    print(
        f"총 작성 대상: "
        f"{len(records)}건"
    )

    product_counts: dict[
        str,
        int,
    ] = defaultdict(int)

    for record in records:
        product_counts[
            record.product
        ] += 1

    for product in TARGET_SHEETS:

        print(
            f"  - {product}: "
            f"{product_counts.get(product, 0)}건"
        )

    if inferred_messages:

        print(
            "\n[연장 계약 추정 포함]"
        )

        for message in inferred_messages:

            print(
                f"  - {message}"
            )

    if warnings:

        print(
            "\n[검토 필요]"
        )

        for warning in warnings:

            print(
                f"  - {warning}"
            )


if __name__ == "__main__":

    try:
        main()

    except Exception as error:

        print(
            f"오류: {error}",
            file=sys.stderr,
        )

        sys.exit(1)