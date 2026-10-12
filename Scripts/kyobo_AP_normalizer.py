# OPENPYXL 전용 버전
from pathlib import Path
import os
import sys
from numbers import Number

from report_paths import ap_input_dir, ap_output_dir, exchange_dir, settings, single_ap_input

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


DEFAULT_INPUT_NAME = "교보_AP 및 잔고.xlsx"
DEFAULT_OUTPUT_NAME = "교보_AP 및 잔고_정리.xlsx"

ALL_SHEET_NAME = "전체_내림차순"
OVER_10_SHEET_NAME = "10%초과"

HEADERS = [
    "상품",
    "계좌번호",
    "계약자명",
    "종목명",
    "투자비중",
    "투자사유및전망",
    "보유수량",
    "취득단가(원화)",
    "매입액원화",
    "시가(분기말기준,원화)",
    "보유잔액(원화)",
]


def is_number(value):
    """True/False를 제외한 숫자인지 확인한다."""
    return isinstance(value, Number) and not isinstance(value, bool)


def multiply_if_number(value_1, value_2):
    """두 값이 모두 숫자면 곱하고, 아니면 None을 반환한다."""
    if is_number(value_1) and is_number(value_2):
        return float(value_1) * float(value_2)
    return None


def is_cash_record(record):
    """10%초과 시트에서 제외할 예수금 행인지 판별한다."""
    if len(record) < 4:
        return False

    stock_name = str(record[3]).replace(" ", "")

    return "예수금" in stock_name


def find_input_file(base_dir):
    return single_ap_input("교보")


def find_stock_header_row(ws):
    """B열='종목명', E열='평가금액'인 종목 헤더 행을 찾는다."""
    for row_no in range(1, ws.max_row + 1):
        if (
            ws.cell(row_no, 2).value == "종목명"
            and ws.cell(row_no, 5).value == "평가금액"
        ):
            return row_no
    return None


def build_cash_rows(ws, stock_header_row, account_no, contractor):
    """
    종목 헤더 위쪽의 통화별 추정예수금을 결과 행으로 만든다.

    원본 구조:
    - 통화 행: A열=KRW/USD 등, F열=적용환율
    - 바로 다음 행: B열=추정예수금, F열=추정예수금(￦)

    출력:
    - KRW: 원화예수금
    - USD 등 외화: 외화예수금(USD)
    - 취득단가(원화): 0
    - 시가(분기말기준,원화): 0

    투자비중은 여기서 계산하지 않고,
    같은 계좌의 보유종목+보유예수금 전체 보유잔액 합계로 나중에 계산한다.
    """
    records = []

    for row_no in range(1, stock_header_row - 1):
        currency = ws.cell(row_no, 1).value

        if not (
            isinstance(currency, str)
            and len(currency.strip()) == 3
            and currency.strip().isalpha()
        ):
            continue

        currency = currency.strip().upper()
        next_row = row_no + 1

        estimated_cash = ws.cell(next_row, 2).value       # 추정예수금
        exchange_rate = ws.cell(row_no, 6).value          # 적용환율
        estimated_cash_krw = ws.cell(next_row, 6).value   # 추정예수금(￦)

        if not is_number(estimated_cash) or float(estimated_cash) == 0:
            continue

        if not is_number(estimated_cash_krw):
            estimated_cash_krw = multiply_if_number(estimated_cash, exchange_rate)

        if not is_number(estimated_cash_krw):
            print(
                f"[건너뜀] {ws.title} / {currency} 예수금: "
                "추정예수금(￦)을 계산하지 못했습니다."
            )
            continue

        if currency == "KRW":
            cash_name = "원화예수금"
        else:
            cash_name = f"외화예수금({currency})"

        records.append(
            [
                "",                                             # 상품
                str(account_no),                                # 계좌번호
                str(contractor),                                # 계약자명
                cash_name,                                      # 종목명
                0,                                              # 투자비중: 계좌 총자산 기준으로 나중에 재계산
                "",                                             # 투자사유및전망
                float(estimated_cash),                          # 보유수량 = 추정예수금
                0,                                              # 취득단가(원화)
                float(estimated_cash_krw),                      # 매입액원화
                0,                                              # 시가(분기말기준,원화)
                float(estimated_cash_krw),                      # 보유잔액(원화)
            ]
        )

    return records


def apply_account_total_weight(records, account_name):
    """
    계좌별 총자산 기준으로 투자비중을 계산한다.

    총자산평가금액 =
        모든 보유종목의 보유잔액(원화)
        + 모든 예수금의 보유잔액(원화)

    투자비중 =
        각 행의 보유잔액(원화) / 총자산평가금액
    """
    total_asset_krw = sum(
        float(record[10])
        for record in records
        if is_number(record[10])
    )

    if total_asset_krw == 0:
        raise ValueError(
            f"{account_name} 계좌의 총자산평가금액이 0이라 투자비중을 계산할 수 없습니다."
        )

    for record in records:
        if is_number(record[10]):
            record[4] = float(record[10]) / total_asset_krw
        else:
            record[4] = 0

    return total_asset_krw


def extract_records(input_path: Path):
    """
    각 사람별 시트에서 종목과 예수금을 추출한다.

    종목:
    - 보유수량 = 매도가능수량
    - 취득단가(원화) = 평균단가 * 환율
    - 매입액원화 = 원화매입금액
    - 시가(분기말기준,원화) = 최근종가 * 환율
    - 보유잔액(원화) = 원화평가금액

    예수금:
    - KRW는 원화예수금
    - USD 등 외화는 외화예수금(USD) 형식으로 표시
    - 취득단가(원화), 시가(분기말기준,원화)는 0

    투자비중:
    - 종목과 예수금을 모두 더한 계좌별 총자산평가금액 기준
    - 각 행의 보유잔액(원화) / 계좌 총자산평가금액
    """
    source_wb = load_workbook(input_path, data_only=True)
    all_records = []

    excluded_sheet_names = {
        ALL_SHEET_NAME,
        OVER_10_SHEET_NAME,
        "10%이상 종목",
    }

    for ws in source_wb.worksheets:
        if ws.title in excluded_sheet_names:
            continue

        contractor = ws["A1"].value
        account_no = ws["B1"].value

        if contractor in (None, "") or account_no in (None, ""):
            print(f"[건너뜀] {ws.title}: A1 또는 B1이 비어 있습니다.")
            continue

        stock_header_row = find_stock_header_row(ws)
        if stock_header_row is None:
            print(f"[건너뜀] {ws.title}: 종목 헤더를 찾지 못했습니다.")
            continue

        customer_records = []

        # 예수금 행 추가
        customer_records.extend(
            build_cash_rows(
                ws=ws,
                stock_header_row=stock_header_row,
                account_no=account_no,
                contractor=contractor,
            )
        )

        # 종목 헤더 다음 행부터 2개 행씩 읽는다.
        # 예: 10~11행이 첫 종목이면 실제 종목명/평가금액은 11행에 있다.
        for first_row in range(stock_header_row + 1, ws.max_row + 1, 2):
            second_row = first_row + 1
            if second_row > ws.max_row:
                break

            currency = ws.cell(second_row, 1).value
            stock_name = ws.cell(second_row, 2).value
            sellable_quantity = ws.cell(second_row, 3).value        # 매도가능수량
            recent_close = ws.cell(second_row, 4).value             # 최근종가
            krw_valuation = ws.cell(second_row, 7).value            # 원화평가금액
            exchange_rate = ws.cell(second_row, 8).value            # 적용환율

            average_price = ws.cell(first_row, 4).value             # 평균단가
            krw_buy_amount = ws.cell(first_row, 7).value            # 원화매입금액

            if not isinstance(currency, str):
                continue

            if stock_name in (None, "") or not is_number(krw_valuation):
                continue

            acquisition_unit_krw = multiply_if_number(average_price, exchange_rate)
            market_price_krw = multiply_if_number(recent_close, exchange_rate)

            customer_records.append(
                [
                    "",                                         # 상품
                    str(account_no),                            # 계좌번호: B1
                    str(contractor),                            # 계약자명: A1
                    str(stock_name),                            # 종목명
                    0,                                          # 투자비중: 계좌 총자산 기준으로 나중에 재계산
                    "",                                         # 투자사유및전망
                    sellable_quantity if is_number(sellable_quantity) else None,
                    acquisition_unit_krw,                       # 평균단가 * 환율
                    krw_buy_amount if is_number(krw_buy_amount) else None,
                    market_price_krw,                           # 최근종가 * 환율
                    krw_valuation if is_number(krw_valuation) else None,
                ]
            )

        apply_account_total_weight(customer_records, contractor)

        customer_records.sort(key=lambda record: record[4], reverse=True)
        all_records.extend(customer_records)

    source_wb.close()
    return all_records


def style_result_sheet(ws, data_row_count):
    """결과 시트에 기본 서식을 적용한다."""
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin_side = Side(style="thin", color="D9E2F3")
    thin_border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side,
    )

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = thin_border

    last_row = data_row_count + 1

    for row in ws.iter_rows(min_row=2, max_row=last_row, min_col=1, max_col=11):
        for cell in row:
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    for row_no in range(2, last_row + 1):
        ws.cell(row_no, 5).number_format = "0.00%"           # 투자비중
        ws.cell(row_no, 7).number_format = "#,##0.########"  # 보유수량
        ws.cell(row_no, 8).number_format = "#,##0.00"        # 취득단가(원화)
        ws.cell(row_no, 9).number_format = "#,##0"           # 매입액원화
        ws.cell(row_no, 10).number_format = "#,##0.00"       # 시가
        ws.cell(row_no, 11).number_format = "#,##0"          # 보유잔액

    widths = {
        1: 14,  # 상품
        2: 18,  # 계좌번호
        3: 18,  # 계약자명
        4: 32,  # 종목명
        5: 14,  # 투자비중
        6: 36,  # 투자사유및전망
        7: 14,  # 보유수량
        8: 18,  # 취득단가
        9: 18,  # 매입액원화
        10: 22, # 시가
        11: 18, # 보유잔액
    }

    for column_no, width in widths.items():
        ws.column_dimensions[get_column_letter(column_no)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:K{last_row}"
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[1].height = 36


def write_sheet(output_wb, sheet_name, records):
    ws = output_wb.create_sheet(title=sheet_name)
    ws.append(HEADERS)

    for record in records:
        ws.append(record)

    style_result_sheet(ws, len(records))


def create_output_file(records, output_path: Path):
    """
    하나의 엑셀 파일에 두 시트를 만든다.

    1. 전체_내림차순: 고객별 전체 종목과 예수금을 투자비중 내림차순으로 정리
    2. 10%초과: 투자비중이 정확히 10%를 초과하는 종목만 정리
       단, 예수금 행은 포함하지 않는다.
    """
    output_wb = Workbook()
    default_ws = output_wb.active
    output_wb.remove(default_ws)

    over_10_records = [
        record
        for record in records
        if record[4] > 0.10 and not is_cash_record(record)
    ]

    write_sheet(output_wb, ALL_SHEET_NAME, records)
    write_sheet(output_wb, OVER_10_SHEET_NAME, over_10_records)

    temporary_path = output_path.with_name(f".{output_path.stem}_작성중.xlsx")

    if temporary_path.exists():
        temporary_path.unlink()

    try:
        output_wb.save(temporary_path)
        output_wb.close()

        # 파일이 실제로 열리는지 검증
        verification_wb = load_workbook(
            temporary_path,
            data_only=False,
            read_only=True,
        )
        verification_wb.close()

        try:
            os.replace(temporary_path, output_path)
        except PermissionError as error:
            raise PermissionError(
                f"'{output_path.name}' 파일이 Excel에서 열려 있습니다. "
                "결과 파일을 완전히 닫고 다시 실행해 주세요."
            ) from error

    except Exception:
        output_wb.close()
        if temporary_path.exists():
            try:
                temporary_path.unlink()
            except OSError:
                pass
        raise

    return len(over_10_records)


def main():
    base_dir = ap_input_dir("교보")
    output_dir = ap_output_dir("교보")

    # 실행 예시 1:
    #   python kyobo_ap_normalizer_openpyxl_extended_v4_cash_total_weight.py
    #
    # 실행 예시 2:
    #   python kyobo_ap_normalizer_openpyxl_extended_v4_cash_total_weight.py "교보_AP 및 잔고.xlsx"
    #
    # 실행 예시 3:
    #   python kyobo_ap_normalizer_openpyxl_extended_v4_cash_total_weight.py "원본.xlsx" "결과.xlsx"

    if len(sys.argv) >= 2:
        input_arg = Path(sys.argv[1])
        input_path = input_arg if input_arg.is_absolute() else base_dir / input_arg
    else:
        input_path = find_input_file(base_dir)

    if not input_path.exists():
        raise FileNotFoundError(f"입력 파일을 찾지 못했습니다: {input_path}")

    if len(sys.argv) >= 3:
        output_arg = Path(sys.argv[2])
        output_path = output_arg if output_arg.is_absolute() else output_dir / output_arg
    else:
        output_path = output_dir / DEFAULT_OUTPUT_NAME

    records = extract_records(input_path)
    over_10_count = create_output_file(records, output_path)

    cash_count = sum(1 for record in records if is_cash_record(record))

    print(f"원본 파일: {input_path.name}")
    print(f"전체 행 수: {len(records)}개")
    print(f"예수금 행 수: {cash_count}개")
    print(f"10% 초과 종목 수: {over_10_count}개")
    print(f"완료 파일: {output_path}")


if __name__ == "__main__":
    main()