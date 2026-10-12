from collections import defaultdict
from numbers import Number
from pathlib import Path
import sys

from report_paths import ap_input_dir, ap_output_dir, exchange_dir, settings, single_ap_input

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


INPUT_FILE_NAME = "한투_AP 및 잔고.xlsx"
OUTPUT_FILE_NAME = "한투_AP 및 잔고_정리.xlsx"

FX_SHEET_NAME = "외화잔고현황"
BALANCE_SHEET_NAME = "잔고내역(원화)"
DETAIL_SHEET_NAME = "해외증권상세내역"

ALL_SHEET_NAME = "전체_내림차순"
OVER_10_SHEET_NAME = "10%초과"

HEADERS = [
    "상품",
    "계좌번호",
    "계약자명",
    "종목명",
    "투자비중",
    "투자사유및전망",
    "보유수량",
    "취득단가(원화)",
    "매입액원화",
    "시가(분기말기준,원화)",
    "보유잔액(원화)",
]


def is_number(value):
    """bool을 제외한 숫자인지 확인한다."""
    return isinstance(value, Number) and not isinstance(value, bool)


def clean_text(value):
    """키 비교에 사용할 문자열의 앞뒤 공백을 제거한다."""
    if value is None:
        return ""
    return str(value).strip()


def is_cash_record(record):
    """10%초과 시트에서 제외할 예수금 행인지 판별한다."""
    if len(record) < 4:
        return False

    stock_name = str(record[3]).replace(" ", "")

    return "예수금" in stock_name


def find_input_file(base_dir):
    return single_ap_input("한투")


def validate_sheet_headers(fx_ws, balance_ws, detail_ws):
    """사용할 열의 헤더가 예상 구조와 일치하는지 확인한다."""
    expected = [
        (fx_ws, "B1", "계좌번호"),
        (fx_ws, "C1", "계좌명"),
        (fx_ws, "D1", "통화코드"),
        (fx_ws, "E1", "외화잔고"),
        (fx_ws, "J1", "출금가능외화"),
        (fx_ws, "L1", "출금가능금액원화평가"),
        (balance_ws, "B1", "계좌번호"),
        (balance_ws, "C1", "계좌명"),
        (balance_ws, "O1", "해외증권 총자산"),
        (detail_ws, "B1", "계좌번호"),
        (detail_ws, "C1", "계좌명"),
        (detail_ws, "D1", "종목명"),
        (detail_ws, "H1", "체결보유수량"),
        (detail_ws, "J1", "매입단가"),
        (detail_ws, "K1", "매입금액"),
        (detail_ws, "L1", "현재가"),
        (detail_ws, "M1", "평가금액"),
    ]

    errors = []
    for ws, address, expected_value in expected:
        actual_value = ws[address].value
        if clean_text(actual_value) != expected_value:
            errors.append(
                f"{ws.title}!{address}: "
                f"예상='{expected_value}', 실제={actual_value!r}"
            )

    if errors:
        raise ValueError(
            "원본 파일의 헤더 구조가 예상과 다릅니다.\n- "
            + "\n- ".join(errors)
        )


def build_foreign_balance_keys(fx_ws):
    """
    외화잔고현황에서 (계좌번호, 계좌명) 키 집합을 만든다.
    한 계좌가 통화별로 여러 줄이어도 키는 한 번만 보관한다.
    """
    keys = set()

    for row_no in range(2, fx_ws.max_row + 1):
        account_no = clean_text(fx_ws.cell(row=row_no, column=2).value)
        account_name = clean_text(fx_ws.cell(row=row_no, column=3).value)

        if account_no and account_name:
            keys.add((account_no, account_name))

    return keys


def build_total_asset_map(balance_ws):
    """
    잔고내역(원화)에서 계좌별 해외증권 총자산을 읽는다.

    - B열: 계좌번호
    - C열: 계좌명
    - O열: 해외증권 총자산
    - 데이터 시작: 3행(1~2행은 헤더)
    """
    total_asset_map = {}
    account_order = []

    for row_no in range(3, balance_ws.max_row + 1):
        account_no = clean_text(balance_ws.cell(row=row_no, column=2).value)
        account_name = clean_text(balance_ws.cell(row=row_no, column=3).value)
        total_asset = balance_ws.cell(row=row_no, column=15).value

        if not account_no or not account_name:
            continue

        key = (account_no, account_name)
        account_order.append(key)

        if not is_number(total_asset):
            total_asset_map[key] = None
        else:
            total_asset_map[key] = float(total_asset)

    return total_asset_map, account_order


def build_cash_records(fx_ws, total_asset_map):
    """
    외화잔고현황에서 원화예수금과 외화예수금을 만든다.

    원화:
    - 통화코드가 KRW인 행
    - E열 외화잔고 값을 그대로 사용

    외화:
    - 통화코드가 KRW가 아닌 행
    - 종목명은 외화예수금(USD), 외화예수금(JPY) 형식
    - 보유수량은 J열 출금가능외화
    - 매입액원화와 보유잔액(원화)은 L열 출금가능금액원화평가

    공통:
    - 취득단가(원화), 시가(분기말기준,원화)는 0
    - 투자비중 = 보유잔액(원화) / 잔고내역(원화) O열 해외증권 총자산
    - 10%초과 시트에는 포함하지 않는다.
    """
    cash_records_by_key = defaultdict(list)

    for row_no in range(2, fx_ws.max_row + 1):
        account_no = clean_text(fx_ws.cell(row=row_no, column=2).value)
        account_name = clean_text(fx_ws.cell(row=row_no, column=3).value)
        currency_code = clean_text(fx_ws.cell(row=row_no, column=4).value).upper()

        if not account_no or not account_name or not currency_code:
            continue

        key = (account_no, account_name)
        total_asset = total_asset_map.get(key)

        if not is_number(total_asset) or float(total_asset) == 0:
            continue

        if currency_code == "KRW":
            cash_name = "원화예수금"
            cash_quantity = fx_ws.cell(row=row_no, column=5).value              # 외화잔고
            cash_balance_krw = fx_ws.cell(row=row_no, column=5).value           # 외화잔고 그대로

        else:
            cash_name = f"외화예수금({currency_code})"
            cash_quantity = fx_ws.cell(row=row_no, column=10).value             # 출금가능외화
            cash_balance_krw = fx_ws.cell(row=row_no, column=12).value          # 출금가능금액원화평가

        if not is_number(cash_balance_krw) or float(cash_balance_krw) == 0:
            continue

        cash_balance_krw = float(cash_balance_krw)

        cash_records_by_key[key].append([
            "",                                      # 상품
            account_no,                              # 계좌번호
            account_name,                            # 계약자명
            cash_name,                               # 종목명
            cash_balance_krw / float(total_asset),   # 투자비중
            "",                                      # 투자사유및전망
            float(cash_quantity) if is_number(cash_quantity) else cash_balance_krw,
            0,                                       # 취득단가(원화)
            cash_balance_krw,                        # 매입액원화
            0,                                       # 시가(분기말기준,원화)
            cash_balance_krw,                        # 보유잔액(원화)
        ])

    return cash_records_by_key


def new_holding_bucket():
    """
    계좌·종목별 집계용 빈 구조를 만든다.

    동일 계좌·동일 종목이 여러 행이면:
    - 보유수량, 매입액원화, 보유잔액원화는 합산
    - 취득단가와 시가는 보유수량 기준 가중평균
    """
    return {
        "valuation_for_weight": 0.0,    # M열 평가금액
        "quantity": 0.0,                # H열 체결보유수량
        "buy_amount_krw": 0.0,          # K열 매입금액
        "holding_balance_krw": 0.0,     # M열 평가금액
        "unit_cost_weight_sum": 0.0,    # J열 매입단가 * 수량
        "market_price_weight_sum": 0.0, # L열 현재가 * 수량
        "price_weight_quantity": 0.0,
        "unit_cost_simple_sum": 0.0,
        "unit_cost_simple_count": 0,
        "market_price_simple_sum": 0.0,
        "market_price_simple_count": 0,
    }


def add_holding_row(bucket, quantity, buy_unit_price, buy_amount, current_price, valuation):
    """
    해외증권상세내역 한 행을 계좌·종목 집계값에 더한다.

    사용자 수정 규칙:
    - 보유수량: H열 체결보유수량
    - 취득단가(원화): J열 매입단가
    - 매입액원화: K열 매입금액
    - 시가(분기말기준,원화): L열 현재가
    - 보유잔액(원화): M열 평가금액
    - 투자비중: M열 평가금액 / 잔고내역(원화) O열 해외증권 총자산
    """
    if is_number(valuation):
        bucket["valuation_for_weight"] += float(valuation)
        bucket["holding_balance_krw"] += float(valuation)

    if is_number(quantity):
        bucket["quantity"] += float(quantity)

    if is_number(buy_amount):
        bucket["buy_amount_krw"] += float(buy_amount)

    if is_number(buy_unit_price):
        bucket["unit_cost_simple_sum"] += float(buy_unit_price)
        bucket["unit_cost_simple_count"] += 1

    if is_number(current_price):
        bucket["market_price_simple_sum"] += float(current_price)
        bucket["market_price_simple_count"] += 1

    # 같은 종목이 여러 행이면 수량 기준 가중평균.
    # 수량이 없거나 0이면 simple average fallback.
    if is_number(quantity) and float(quantity) != 0:
        q = float(quantity)

        if is_number(buy_unit_price):
            bucket["unit_cost_weight_sum"] += float(buy_unit_price) * q

        if is_number(current_price):
            bucket["market_price_weight_sum"] += float(current_price) * q

        bucket["price_weight_quantity"] += q


def finalize_holding(bucket):
    """집계 bucket을 최종 출력용 값으로 변환한다."""
    if bucket["price_weight_quantity"] != 0:
        acquisition_unit_krw = (
            bucket["unit_cost_weight_sum"]
            / bucket["price_weight_quantity"]
        )
        market_price = (
            bucket["market_price_weight_sum"]
            / bucket["price_weight_quantity"]
        )
    else:
        acquisition_unit_krw = (
            bucket["unit_cost_simple_sum"]
            / bucket["unit_cost_simple_count"]
            if bucket["unit_cost_simple_count"]
            else None
        )
        market_price = (
            bucket["market_price_simple_sum"]
            / bucket["market_price_simple_count"]
            if bucket["market_price_simple_count"]
            else None
        )

    return {
        "valuation_for_weight": bucket["valuation_for_weight"],
        "quantity": bucket["quantity"],
        "acquisition_unit_krw": acquisition_unit_krw,
        "buy_amount_krw": bucket["buy_amount_krw"],
        "market_price": market_price,
        "holding_balance_krw": bucket["holding_balance_krw"],
    }


def extract_records(source_wb):
    """
    세 시트의 계좌번호+계좌명을 키로 연결하여 종목 비중과 추가 정보를 계산한다.

    종목명: 해외증권상세내역 D열
    투자비중: 해외증권상세내역 M열 평가금액 / 잔고내역(원화) O열 해외증권 총자산

    추가 열:
    - 보유수량: H열 체결보유수량
    - 취득단가(원화): J열 매입단가
    - 매입액원화: K열 매입금액
    - 시가(분기말기준,원화): L열 현재가
    - 보유잔액(원화): M열 평가금액

    예수금:
    - 원화예수금: 외화잔고현황 시트의 KRW 행 E열 외화잔고
    - 외화예수금(XXX): 외화잔고현황 시트의 해당 통화 행 L열 출금가능금액원화평가
    - 예수금 취득단가(원화), 시가는 0
    - 예수금은 전체_내림차순에만 포함하고 10%초과 시트에서는 제외
    """
    required_sheets = {
        FX_SHEET_NAME,
        BALANCE_SHEET_NAME,
        DETAIL_SHEET_NAME,
    }
    missing_sheets = required_sheets - set(source_wb.sheetnames)

    if missing_sheets:
        raise KeyError(
            "필수 시트가 없습니다: "
            + ", ".join(sorted(missing_sheets))
        )

    fx_ws = source_wb[FX_SHEET_NAME]
    balance_ws = source_wb[BALANCE_SHEET_NAME]
    detail_ws = source_wb[DETAIL_SHEET_NAME]

    validate_sheet_headers(fx_ws, balance_ws, detail_ws)

    foreign_balance_keys = build_foreign_balance_keys(fx_ws)
    total_asset_map, account_order = build_total_asset_map(balance_ws)
    cash_records_by_key = build_cash_records(fx_ws, total_asset_map)

    holdings = defaultdict(lambda: defaultdict(new_holding_bucket))

    missing_fx_keys = set()
    missing_balance_keys = set()
    invalid_total_asset_keys = set()

    for row_no in range(2, detail_ws.max_row + 1):
        account_no = clean_text(detail_ws.cell(row=row_no, column=2).value)
        account_name = clean_text(detail_ws.cell(row=row_no, column=3).value)
        stock_name = clean_text(detail_ws.cell(row=row_no, column=4).value)

        quantity = detail_ws.cell(row=row_no, column=8).value       # H 체결보유수량
        buy_unit_price = detail_ws.cell(row=row_no, column=10).value # J 매입단가
        buy_amount = detail_ws.cell(row=row_no, column=11).value     # K 매입금액
        current_price = detail_ws.cell(row=row_no, column=12).value  # L 현재가
        valuation = detail_ws.cell(row=row_no, column=13).value      # M 평가금액

        if not account_no and not account_name and not stock_name:
            continue

        if not account_no or not account_name or not stock_name:
            print(
                f"[건너뜀] {DETAIL_SHEET_NAME} {row_no}행: "
                "계좌번호·계좌명·종목명 중 빈 값이 있습니다."
            )
            continue

        if not is_number(valuation):
            print(
                f"[건너뜀] {DETAIL_SHEET_NAME} {row_no}행: "
                f"평가금액이 숫자가 아닙니다. 값={valuation!r}"
            )
            continue

        key = (account_no, account_name)

        if key not in foreign_balance_keys:
            missing_fx_keys.add(key)
            continue

        if key not in total_asset_map:
            missing_balance_keys.add(key)
            continue

        total_asset = total_asset_map[key]
        if not is_number(total_asset) or float(total_asset) == 0:
            invalid_total_asset_keys.add(key)
            continue

        bucket = holdings[key][stock_name]
        add_holding_row(
            bucket=bucket,
            quantity=quantity,
            buy_unit_price=buy_unit_price,
            buy_amount=buy_amount,
            current_price=current_price,
            valuation=valuation,
        )

    for key in sorted(missing_fx_keys):
        print(f"[건너뜀] 외화잔고현황에서 키를 찾지 못했습니다: {key}")

    for key in sorted(missing_balance_keys):
        print(f"[건너뜀] 잔고내역(원화)에서 키를 찾지 못했습니다: {key}")

    for key in sorted(invalid_total_asset_keys):
        print(f"[건너뜀] 해외증권 총자산이 0 또는 숫자가 아닙니다: {key}")

    all_records = []

    # 잔고내역(원화)의 계좌 순서를 유지하고,
    # 각 계좌 안에서는 투자비중 내림차순으로 정렬한다.
    for key in account_order:
        if key not in holdings and key not in cash_records_by_key:
            continue

        account_no, account_name = key
        total_asset = total_asset_map[key]
        customer_records = []

        # 예수금 행 추가
        customer_records.extend(cash_records_by_key.get(key, []))

        for stock_name, bucket in holdings.get(key, {}).items():
            final = finalize_holding(bucket)
            investment_weight = final["valuation_for_weight"] / total_asset

            customer_records.append([
                "",                                  # 상품
                account_no,                          # 계좌번호
                account_name,                        # 계약자명
                stock_name,                          # 종목명
                investment_weight,                   # 투자비중
                "",                                  # 투자사유및전망
                final["quantity"],                   # 보유수량
                final["acquisition_unit_krw"],       # 취득단가(원화)
                final["buy_amount_krw"],             # 매입액원화
                final["market_price"],               # 시가(분기말기준,원화)
                final["holding_balance_krw"],        # 보유잔액(원화)
            ])

        customer_records.sort(
            key=lambda record: record[4],
            reverse=True,
        )
        all_records.extend(customer_records)

    return all_records


def write_result_sheet(ws, records):
    """
    결과 시트 작성.

    Excel 복구 경고를 방지하기 위해 Table 객체는 만들지 않고,
    일반 셀 범위와 자동필터만 사용한다.
    """
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin_side = Side(style="thin", color="D9E2F3")
    border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side,
    )

    for column_no, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=column_no, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    for row_no, record in enumerate(records, start=2):
        for column_no, value in enumerate(record, start=1):
            cell = ws.cell(
                row=row_no,
                column=column_no,
                value=value,
            )
            cell.border = border
            cell.alignment = Alignment(
                vertical="center",
                wrap_text=column_no in {4, 6},
            )

        ws.cell(row=row_no, column=5).number_format = "0.00%"          # 투자비중
        ws.cell(row=row_no, column=7).number_format = "#,##0.########" # 보유수량
        ws.cell(row=row_no, column=8).number_format = "#,##0.00"       # 취득단가
        ws.cell(row=row_no, column=9).number_format = "#,##0"          # 매입액원화
        ws.cell(row=row_no, column=10).number_format = "#,##0.00"      # 시가
        ws.cell(row=row_no, column=11).number_format = "#,##0"         # 보유잔액

    last_row = max(1, len(records) + 1)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:K{last_row}"
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[1].height = 36

    column_widths = {
        "A": 14,
        "B": 20,
        "C": 16,
        "D": 42,
        "E": 14,
        "F": 36,
        "G": 14,
        "H": 18,
        "I": 18,
        "J": 22,
        "K": 18,
    }

    for column_letter, width in column_widths.items():
        ws.column_dimensions[column_letter].width = width


def create_output_workbook(records, output_path: Path):
    """
    하나의 결과 파일에 두 시트를 만든다.

    - 전체_내림차순: 계좌별 전체 종목
    - 10%초과: 투자비중이 정확히 10%를 초과하는 종목
    """
    output_wb = Workbook()
    output_wb.remove(output_wb.active)

    over_10_records = [
        record
        for record in records
        if record[4] > 0.10 and not is_cash_record(record)
    ]

    all_ws = output_wb.create_sheet(ALL_SHEET_NAME)
    write_result_sheet(all_ws, records)

    over_10_ws = output_wb.create_sheet(OVER_10_SHEET_NAME)
    write_result_sheet(over_10_ws, over_10_records)

    temporary_path = output_path.with_name(f".{output_path.stem}_작성중.xlsx")

    if temporary_path.exists():
        temporary_path.unlink()

    try:
        output_wb.save(temporary_path)
        output_wb.close()

        verification_wb = load_workbook(
            temporary_path,
            data_only=False,
            read_only=True,
        )
        verification_wb.close()

        try:
            import os
            os.replace(temporary_path, output_path)
        except PermissionError as error:
            raise PermissionError(
                f"'{output_path.name}' 파일이 Excel에서 열려 있습니다. "
                "결과 파일을 완전히 닫고 다시 실행해 주세요."
            ) from error

    except Exception:
        output_wb.close()
        if temporary_path.exists():
            try:
                temporary_path.unlink()
            except OSError:
                pass
        raise

    return len(over_10_records)


def main():
    # Input/AP에서 읽고 Output/AP에 저장한다.
    base_dir = ap_input_dir("한투")
    output_dir = ap_output_dir("한투")

    # 사용법
    # python han_AP_normalizer.py
    # python han_AP_normalizer.py "한투_AP 및 잔고.xlsx"
    # python han_AP_normalizer.py "원본.xlsx" "결과.xlsx"
    if len(sys.argv) >= 2:
        requested_input = Path(sys.argv[1])
        input_path = (
            requested_input
            if requested_input.is_absolute()
            else base_dir / requested_input
        )
    else:
        input_path = find_input_file(base_dir)

    if not input_path.exists():
        raise FileNotFoundError(
            f"입력 파일을 찾지 못했습니다: {input_path}"
        )

    if len(sys.argv) >= 3:
        requested_output = Path(sys.argv[2])
        output_path = (
            requested_output
            if requested_output.is_absolute()
            else output_dir / requested_output
        )
    else:
        output_path = output_dir / OUTPUT_FILE_NAME

    source_wb = load_workbook(
        input_path,
        data_only=True,
        read_only=False,
    )

    records = extract_records(source_wb)
    over_10_count = create_output_workbook(
        records,
        output_path,
    )

    cash_count = sum(1 for record in records if is_cash_record(record))

    print(f"원본 파일: {input_path.name}")
    print(f"전체 행 수: {len(records)}개")
    print(f"예수금 행 수: {cash_count}개")
    print(f"10% 초과 종목 수: {over_10_count}개")
    print(f"완료 파일: {output_path}")


if __name__ == "__main__":
    main()