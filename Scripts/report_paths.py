"""Report 폴더 기준 경로와 실행 설정. 현재 터미널 위치에 의존하지 않는다."""
from datetime import date
from pathlib import Path
import json

ROOT = Path(__file__).resolve().parents[1]
BROKERS = ("삼성", "한투", "KB", "교보", "메리츠")
CONTRACT_METADATA_SHEET = "_report_metadata"


def stamp_contract_period(workbook, start, end):
    """계약목록을 생성한 전체 기간을 결과 파일 안에 보관한다."""
    if CONTRACT_METADATA_SHEET in workbook.sheetnames:
        del workbook[CONTRACT_METADATA_SHEET]
    ws = workbook.create_sheet(CONTRACT_METADATA_SHEET)
    ws.append(["schema", "contract_periods_v2"])
    ws.append(["report_start", start.isoformat()])
    ws.append(["report_end", end.isoformat()])
    ws.sheet_state = "hidden"


def settings():
    path = ROOT / "settings.json"
    with path.open(encoding="utf-8-sig") as stream:
        value = json.load(stream)
    start = date.fromisoformat(value["report_start"])
    end = date.fromisoformat(value["report_end"])
    if start > end:
        raise ValueError("settings.json의 시작일이 종료일보다 늦습니다.")
    if value.get("ap_exchange_date"):
        date.fromisoformat(value["ap_exchange_date"])
    return value


def period():
    value = settings()
    return date.fromisoformat(value["report_start"]), date.fromisoformat(value["report_end"])


def update_period(start, end):
    value = settings()
    value.update(report_start=start.isoformat(), report_end=end.isoformat())
    path = ROOT / "settings.json"
    temporary = path.with_suffix(".tmp")
    temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    temporary.replace(path)


def output_dir(*parts):
    path = ROOT / "Output"
    for part in parts:
        path /= part
    path.mkdir(parents=True, exist_ok=True)
    return path


def trade_input_dir(broker):
    return ROOT / "Input" / "거래내역" / broker


def ap_input_dir(broker):
    return ROOT / "Input" / "AP" / broker


def ap_output_dir(broker=None):
    return output_dir("AP", broker) if broker else output_dir("AP")


def exchange_dir():
    return ROOT / "Input" / "exchange_rate"


def final_dir():
    return output_dir("최종")


def contract_input():
    return ROOT / "Input" / "계약" / "ContractList.xlsx"


def contract_output():
    return output_dir("계약") / "ContractList.xlsx"


def excel_files(folder, extensions):
    if not folder.is_dir():
        return []
    return sorted((p for p in folder.iterdir() if p.is_file()
                   and p.suffix.lower() in extensions
                   and not p.name.startswith(("~$", "."))
                   and "_정리" not in p.stem and "_작성중" not in p.stem),
                  key=lambda p: p.name.casefold())


def trade_files(broker):
    folder = trade_input_dir(broker)
    extensions = {".xlsx", ".xlsm", ".xls"} if broker == "메리츠" else {".xlsx", ".xlsm"}
    unsupported = excel_files(folder, {".xls"}) if broker != "메리츠" else []
    if unsupported:
        raise ValueError(f"{broker} 거래내역은 Excel에서 .xlsx로 저장하세요: {unsupported[0].name}")
    files = excel_files(folder, extensions)
    # .xlsx와 .xlsm이 같은 stem이면 결과가 덮어써지므로 중단한다.
    stems = [p.stem.casefold() for p in files]
    if len(stems) != len(set(stems)):
        raise ValueError(f"{folder}에 이름이 같은 원본이 둘 이상 있습니다. 파일명을 구분하세요.")
    return files


def trade_output_path(broker, input_file):
    return output_dir("거래내역", broker) / f"{input_file.stem}_정리.xlsx"


def normalized_trade_files(broker):
    sources = trade_files(broker)
    if not sources:
        raise FileNotFoundError(f"{trade_input_dir(broker)}에 거래 원본을 넣으세요.")
    outputs = [trade_output_path(broker, p) for p in sources]
    for source, result in zip(sources, outputs):
        if not result.exists():
            raise FileNotFoundError(f"거래 정리를 먼저 실행하세요: {result}")
        if source.stat().st_mtime > result.stat().st_mtime:
            raise ValueError(f"원본이 정리 결과보다 최신입니다. 거래 정리를 다시 실행하세요: {source.name}")
        assert_trade_calculations_valid(result)
    return outputs


def assert_trade_calculations_valid(path):
    """계산 실패로 0이 기록된 정리 결과를 최종 취합하지 않는다."""
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        if "검토필요" not in wb.sheetnames:
            return
        failures = sum(
            1 for row in wb["검토필요"].iter_rows(values_only=True)
            if any("계산 실패" in str(value or "") for value in row)
        )
    finally:
        wb.close()
    if failures:
        raise ValueError(
            f"{path.name}: 계산 실패 {failures}건으로 취합을 중단합니다. "
            "검토필요 시트와 거래일별 환율 자료를 확인하고 거래 정리를 다시 실행하세요."
        )


def single_ap_input(broker):
    folder = ap_input_dir(broker)
    files = excel_files(folder, {".xlsx"})
    if len(files) != 1:
        raise ValueError(f"{folder}에 현재 기준일의 AP 원본 .xlsx를 정확히 1개 넣으세요. 현재 {len(files)}개")
    return files[0]


def contract_sheet(workbook):
    """생성 기간이 현재 설정과 같은 5열 계약목록만 사용한다."""
    start, end = period()
    if CONTRACT_METADATA_SHEET not in workbook.sheetnames:
        raise ValueError("계약 결과에 생성 기간 정보가 없습니다. 계약 대상 추출을 다시 실행하세요.")
    metadata = dict(workbook[CONTRACT_METADATA_SHEET].iter_rows(
        min_col=1, max_col=2, values_only=True))
    if (metadata.get("schema") != "contract_periods_v2"
            or metadata.get("report_start") != start.isoformat()
            or metadata.get("report_end") != end.isoformat()):
        raise ValueError("계약 결과의 생성 기간과 현재 설정이 다릅니다. 계약 대상 추출을 다시 실행하세요.")
    name = end.strftime("%y%m%d")
    if name not in workbook.sheetnames:
        raise ValueError(f"계약 결과에 {name} 시트가 없습니다. 계약 대상 추출을 먼저 실행하세요.")
    ws = workbook[name]
    expected = [
        "상품명",
        "이름",
        "계좌번호",
        "보고서작성시작일",
        "보고서작성마감일",
    ]
    headers = [str(ws.cell(1, i).value or "").strip() for i in range(1, 6)]
    if headers != expected:
        raise ValueError(
            "계약 결과의 A~E열이 올바르지 않습니다. "
            "contract_report_list.py를 다시 실행하세요."
        )
    return ws
