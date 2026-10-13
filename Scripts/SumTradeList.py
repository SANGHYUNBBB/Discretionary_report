from __future__ import annotations

import math
import re
import unicodedata
from datetime import date, datetime
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Set, Tuple

from report_paths import ROOT, final_dir, contract_output, contract_sheet, normalized_trade_files, period

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel


# ============================================================
# 기본 설정
# ============================================================

BASE_DIR = ROOT
OUTPUT_FILE_NAME = "SumTradeList.xlsx"

# 결과 시트명, 정리 거래파일의 증권사 폴더명
SOURCE_CONFIG: Sequence[Tuple[str, str]] = (
    ("삼성", "삼성"),
    ("메리츠", "메리츠"),
    ("KB", "KB"),
    ("교보", "교보"),
    ("한국투자", "한투"),
)

ACCOUNT_HEADER_ALIASES = {
    "계좌번호",
    "계좌 번호",
    "증권계좌번호",
    "증권 계좌번호",
    "위탁계좌번호",
    "위탁 계좌번호",
}

TRADE_DATE_HEADER_ALIASES = {
    "거래일자",
    "거래 일자",
    "거래일",
    "거래 일",
    "매매일자",
    "매매 일자",
    "매매일",
    "매매 일",
    "체결일자",
    "체결 일자",
    "체결일",
    "체결 일",
    "처리일자",
    "처리 일자",
    "처리일",
    "처리 일",
    "거래일시",
    "거래 일시",
    "체결일시",
    "체결 일시",
}

PF_HEADER_ALIASES = {
    "PF명",
    "PF 명",
    "PF",
    "포트폴리오명",
    "포트폴리오 명",
}

CONTRACT_NAME_HEADER_ALIASES = {
    "계약자명",
    "계약자 명",
    "고객명",
    "고객 명",
    "계좌명",
    "계좌 명",
    "성명",
    "이름",
}

REPORT_START_HEADER_ALIASES = {
    "보고서작성시작일",
    "보고서 작성 시작일",
    "보고서시작일",
    "보고서 시작일",
}

REPORT_END_HEADER_ALIASES = {
    "보고서작성마감일",
    "보고서 작성 마감일",
    "보고서마감일",
    "보고서 마감일",
    "보고서작성종료일",
    "보고서 작성 종료일",
}

HEADER_SCAN_ROWS = 30


# ============================================================
# 문자열 / 계좌번호 / 날짜 처리
# ============================================================

def clean_text(value) -> str:
    if value is None:
        return ""

    text = unicodedata.normalize("NFKC", str(value))
    text = ILLEGAL_CHARACTERS_RE.sub("", text)
    text = text.replace("\u00A0", " ").replace("\u202F", " ")
    return text.strip()


def normalize_header(value) -> str:
    text = clean_text(value)
    return re.sub(r"\s+", "", text).casefold()


def normalized_aliases(values: Iterable[str]) -> Set[str]:
    return {normalize_header(value) for value in values}


ACCOUNT_HEADER_KEYS = normalized_aliases(ACCOUNT_HEADER_ALIASES)
TRADE_DATE_HEADER_KEYS = normalized_aliases(TRADE_DATE_HEADER_ALIASES)
PF_HEADER_KEYS = normalized_aliases(PF_HEADER_ALIASES)
CONTRACT_NAME_HEADER_KEYS = normalized_aliases(CONTRACT_NAME_HEADER_ALIASES)
REPORT_START_HEADER_KEYS = normalized_aliases(REPORT_START_HEADER_ALIASES)
REPORT_END_HEADER_KEYS = normalized_aliases(REPORT_END_HEADER_ALIASES)


def normalize_account(value) -> str:
    """계좌번호 비교용 키를 만든다. 하이픈/공백 등은 제거한다."""
    if value is None or isinstance(value, bool):
        return ""

    if isinstance(value, int):
        text = str(value)
    elif isinstance(value, float):
        if not math.isfinite(value):
            return ""
        text = str(int(value)) if value.is_integer() else format(value, ".15g")
    else:
        text = clean_text(value)
        if text.startswith("'"):
            text = text[1:]
        if re.fullmatch(r"[+-]?\d+\.0+", text):
            text = text.split(".", 1)[0]

    text = unicodedata.normalize("NFKC", text).upper()
    return re.sub(r"[^0-9A-Z]", "", text)


def normalize_name(value) -> str:
    """이름 비교용 키. 공백·대소문자·전각/반각 차이를 제거한다."""
    text = unicodedata.normalize("NFKC", clean_text(value)).casefold()
    return re.sub(r"\s+", "", text)


def parse_date_value(value) -> Optional[date]:
    """Excel 날짜, 날짜 문자열, 숫자형 날짜를 date로 변환한다."""
    if value is None or isinstance(value, bool):
        return None

    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    if isinstance(value, (int, float)):
        if isinstance(value, float) and not math.isfinite(value):
            return None

        numeric_value = float(value)
        integer_value = int(numeric_value)

        # 20260401 같은 YYYYMMDD 숫자형 날짜
        if numeric_value == integer_value and 19000101 <= integer_value <= 99991231:
            try:
                return datetime.strptime(str(integer_value), "%Y%m%d").date()
            except ValueError:
                return None

        # 일반적인 Excel 일련번호 날짜
        try:
            converted = from_excel(numeric_value)
            if isinstance(converted, datetime):
                return converted.date()
            if isinstance(converted, date):
                return converted
        except (OverflowError, TypeError, ValueError):
            return None

        return None

    text = clean_text(value)
    if not text:
        return None

    if text.startswith("'"):
        text = text[1:].strip()

    try:
        return datetime.fromisoformat(text).date()
    except ValueError:
        pass

    separated_formats = (
        (r"\d{4}-\d{1,2}-\d{1,2}", "%Y-%m-%d"),
        (r"\d{4}\.\d{1,2}\.\d{1,2}", "%Y.%m.%d"),
        (r"\d{4}/\d{1,2}/\d{1,2}", "%Y/%m/%d"),
        (r"\d{2}-\d{1,2}-\d{1,2}", "%y-%m-%d"),
        (r"\d{2}\.\d{1,2}\.\d{1,2}", "%y.%m.%d"),
        (r"\d{2}/\d{1,2}/\d{1,2}", "%y/%m/%d"),
    )

    for pattern, date_format in separated_formats:
        if not re.fullmatch(pattern, text):
            continue
        try:
            return datetime.strptime(text, date_format).date()
        except ValueError:
            return None

    if re.fullmatch(r"\d{8}", text):
        try:
            return datetime.strptime(text, "%Y%m%d").date()
        except ValueError:
            return None

    if re.fullmatch(r"\d{6}", text):
        try:
            return datetime.strptime(text, "%y%m%d").date()
        except ValueError:
            return None

    korean_match = re.fullmatch(
        r"(\d{4})\s*년\s*(\d{1,2})\s*월\s*(\d{1,2})\s*일?",
        text,
    )
    if korean_match:
        try:
            return date(
                int(korean_match.group(1)),
                int(korean_match.group(2)),
                int(korean_match.group(3)),
            )
        except ValueError:
            return None

    if re.fullmatch(r"\d+(?:\.0+)?", text):
        try:
            converted = from_excel(float(text))
            if isinstance(converted, datetime):
                return converted.date()
            if isinstance(converted, date):
                return converted
        except (OverflowError, TypeError, ValueError):
            return None

    return None


# ============================================================
# 헤더 찾기
# ============================================================

def find_header_row_and_columns(
    worksheet,
    required_aliases: Set[str],
) -> Optional[Tuple[int, Dict[str, int]]]:
    max_scan_row = min(worksheet.max_row, HEADER_SCAN_ROWS)

    for row_number in range(1, max_scan_row + 1):
        header_map: Dict[str, int] = {}

        for column_number in range(1, worksheet.max_column + 1):
            key = normalize_header(worksheet.cell(row_number, column_number).value)
            if key and key not in header_map:
                header_map[key] = column_number

        if any(alias in header_map for alias in required_aliases):
            return row_number, header_map

    return None


def find_column(header_map: Dict[str, int], aliases: Set[str]) -> Optional[int]:
    for alias in aliases:
        if alias in header_map:
            return header_map[alias]
    return None


def last_nonempty_header_column(worksheet, header_row: int) -> int:
    last_column = 0
    for column_number in range(1, worksheet.max_column + 1):
        if clean_text(worksheet.cell(header_row, column_number).value):
            last_column = column_number
    return last_column


# ============================================================
# 보고서 작성 대상 계약 목록 읽기
# ============================================================

ReportWindow = Tuple[str, date, date]
ReportContracts = Dict[Tuple[str, str], List[ReportWindow]]


def load_report_contracts(contract_file: Path) -> ReportContracts:
    """
    Output/계약/ContractList.xlsx의 현재 보고서 시트 A~E열을 읽는다.

      A열: 상품명
      B열: 이름
      C열: 계좌번호
      D열: 보고서작성시작일
      E열: 보고서작성마감일

    같은 고객/계좌가 분기 중 해지 후 재가입해서 작성기간이 둘 이상이면
    덮어쓰지 않고 모든 기간을 보관한다.
    """
    report_contracts: ReportContracts = {}

    workbook = load_workbook(contract_file, data_only=True, read_only=False)
    try:
        worksheet = contract_sheet(workbook)
        overall_start, overall_end = period()

        if normalize_header(worksheet.cell(1, 4).value) not in REPORT_START_HEADER_KEYS:
            raise ValueError(
                f"{contract_file.name}의 '{worksheet.title}' 시트 D1이 "
                "'보고서작성시작일'이 아닙니다."
            )

        if normalize_header(worksheet.cell(1, 5).value) not in REPORT_END_HEADER_KEYS:
            raise ValueError(
                f"{contract_file.name}의 '{worksheet.title}' 시트 E1이 "
                "'보고서작성마감일'이 아닙니다."
            )

        last_product_name = ""

        for row_number in range(2, worksheet.max_row + 1):
            raw_product = worksheet.cell(row_number, 1).value
            raw_name = worksheet.cell(row_number, 2).value
            raw_account = worksheet.cell(row_number, 3).value
            raw_report_start = worksheet.cell(row_number, 4).value
            raw_report_end = worksheet.cell(row_number, 5).value

            current_product_name = clean_text(raw_product)
            if current_product_name:
                last_product_name = current_product_name
            product_name = current_product_name or last_product_name

            name_key = normalize_name(raw_name)
            account_key = normalize_account(raw_account)

            if not account_key or not name_key:
                continue

            report_start_date = parse_date_value(raw_report_start)
            if report_start_date is None:
                raise ValueError(
                    f"{contract_file.name}의 '{worksheet.title}' 시트 "
                    f"{row_number}행 D열 보고서작성시작일을 날짜로 읽지 못했습니다. "
                    f"이름={clean_text(raw_name)}, 계좌번호={clean_text(raw_account)}, "
                    f"D열 값={clean_text(raw_report_start)}"
                )

            report_end_date = parse_date_value(raw_report_end)
            if report_end_date is None:
                raise ValueError(
                    f"{contract_file.name}의 '{worksheet.title}' 시트 "
                    f"{row_number}행 E열 보고서작성마감일을 날짜로 읽지 못했습니다. "
                    f"이름={clean_text(raw_name)}, 계좌번호={clean_text(raw_account)}, "
                    f"E열 값={clean_text(raw_report_end)}"
                )

            if report_start_date > report_end_date:
                raise ValueError(
                    f"{contract_file.name}의 '{worksheet.title}' 시트 {row_number}행에서 "
                    f"보고서작성시작일({report_start_date})이 "
                    f"보고서작성마감일({report_end_date})보다 늦습니다."
                )

            if not overall_start <= report_start_date <= report_end_date <= overall_end:
                raise ValueError(
                    f"계약 결과 {row_number}행의 작성기간이 현재 보고서 기간을 벗어납니다. "
                    "계약 대상 추출을 다시 실행하세요."
                )

            if not product_name:
                raise ValueError(
                    f"{contract_file.name}의 '{worksheet.title}' 시트 "
                    f"{row_number}행 A열 상품명을 찾지 못했습니다. "
                    f"이름={clean_text(raw_name)}, 계좌번호={clean_text(raw_account)}"
                )

            contract_key = (account_key, name_key)
            report_contracts.setdefault(contract_key, []).append(
                (product_name, report_start_date, report_end_date)
            )
    finally:
        workbook.close()

    if not report_contracts:
        raise ValueError(
            f"{contract_file.name}의 A~E열에서 보고서 작성 대상 계약을 찾지 못했습니다. "
            "A=상품명, B=이름, C=계좌번호, D=보고서작성시작일, "
            "E=보고서작성마감일인지 확인해 주세요."
        )

    for windows in report_contracts.values():
        windows.sort(key=lambda item: (item[1], item[2], item[0]))

    return report_contracts


def product_for_trade_date(
    contract_key: Tuple[str, str],
    trade_date: date,
    report_contracts: ReportContracts,
) -> Optional[str]:
    """
    거래일이 고객별 보고서 작성기간 안이면 해당 상품명을 반환한다.
    작성기간 밖이면 None을 반환한다.
    """
    matches = [
        (product_name, start_date, end_date)
        for product_name, start_date, end_date in report_contracts.get(contract_key, [])
        if start_date <= trade_date <= end_date
    ]

    if not matches:
        return None

    products = {
        clean_text(product_name)
        for product_name, _start, _end in matches
    }

    if len(products) > 1:
        raise ValueError(
            "같은 계좌/이름의 동일 거래일에 서로 다른 상품의 보고서 기간이 겹칩니다. "
            f"거래일={trade_date}, 상품={sorted(products)}"
        )

    return clean_text(matches[0][0])


# ============================================================
# 거래내역 필터링
# ============================================================

def read_filtered_rows(
    source_file: Path,
    report_contracts: ReportContracts,
) -> Tuple[List, List[List], Set[Tuple[str, str]]]:
    """
    고객별 보고서작성시작일~보고서작성마감일에 해당하는 거래만 남긴다.

    예:
    - 보고서 전체 기간: 7/1 ~ 9/30
    - A 고객: 7/1 ~ 9/30 -> 9/30까지 포함
    - B 고객: 7/1 ~ 8/20 해지 -> 8/20까지만 포함
    - C 고객: 8/5 가입 ~ 9/12 해지 -> 8/5~9/12만 포함
    """
    headers: Optional[List] = None
    filtered_rows: List[List] = []
    matched_accounts: Set[Tuple[str, str]] = set()

    workbook = load_workbook(source_file, data_only=True, read_only=False)
    try:
        for worksheet in workbook.worksheets:
            found = find_header_row_and_columns(
                worksheet,
                ACCOUNT_HEADER_KEYS,
            )
            if found is None:
                # 검토필요 같은 보조 시트는 자동으로 건너뛴다.
                continue

            header_row, header_map = found
            account_column = find_column(header_map, ACCOUNT_HEADER_KEYS)
            trade_date_column = find_column(header_map, TRADE_DATE_HEADER_KEYS)
            source_pf_column = find_column(header_map, PF_HEADER_KEYS)
            contract_name_column = find_column(
                header_map,
                CONTRACT_NAME_HEADER_KEYS,
            )

            if account_column is None:
                continue

            if contract_name_column is None:
                raise ValueError(
                    f"{source_file.name}의 '{worksheet.title}' 시트에서 "
                    "계약자명/이름 열을 찾지 못했습니다."
                )

            if trade_date_column is None:
                raise ValueError(
                    f"{source_file.name}의 '{worksheet.title}' 시트에서 "
                    "거래일자/매매일자 열을 찾지 못했습니다."
                )

            last_column = last_nonempty_header_column(worksheet, header_row)
            if last_column == 0:
                continue

            source_headers = [
                worksheet.cell(header_row, column_number).value
                for column_number in range(1, last_column + 1)
            ]

            # PF명 열이 없는 파일도 동일한 출력 구조로 맞춘다.
            if source_pf_column is not None:
                current_headers = list(source_headers)
                pf_output_index = source_pf_column - 1
                insert_pf_column = False
            else:
                current_headers = list(source_headers)
                pf_output_index = contract_name_column
                current_headers.insert(pf_output_index, "PF명")
                insert_pf_column = True

            if headers is None:
                headers = current_headers
            elif [normalize_header(x) for x in headers] != [
                normalize_header(x) for x in current_headers
            ]:
                raise ValueError(
                    f"{source_file.name}의 시트별 헤더 구성이 서로 다릅니다. "
                    f"문제 시트: {worksheet.title}"
                )

            for row_number in range(header_row + 1, worksheet.max_row + 1):
                raw_account = worksheet.cell(row_number, account_column).value
                raw_name = worksheet.cell(row_number, contract_name_column).value

                account_key = normalize_account(raw_account)
                name_key = normalize_name(raw_name)
                contract_key = (account_key, name_key)

                if (
                    not account_key
                    or not name_key
                    or contract_key not in report_contracts
                ):
                    continue

                # 거래파일에 이 계좌가 존재했음을 표시한다.
                # 작성기간 밖의 거래만 있어도 계좌 자체는 매칭으로 본다.
                matched_accounts.add(contract_key)

                raw_trade_date = worksheet.cell(
                    row_number,
                    trade_date_column,
                ).value
                trade_date = parse_date_value(raw_trade_date)

                if trade_date is None:
                    raise ValueError(
                        f"{source_file.name}의 '{worksheet.title}' 시트 "
                        f"{row_number}행 거래일자를 날짜로 읽지 못했습니다. "
                        f"이름={clean_text(raw_name)}, 계좌번호={clean_text(raw_account)}, "
                        f"거래일자 값={clean_text(raw_trade_date)}"
                    )

                # 핵심: 고객별 시작일~마감일 안에 있는 거래만 통과시킨다.
                product_name = product_for_trade_date(
                    contract_key,
                    trade_date,
                    report_contracts,
                )

                if product_name is None:
                    continue

                row_values = [
                    worksheet.cell(row_number, column_number).value
                    for column_number in range(1, last_column + 1)
                ]

                # 계좌번호가 숫자로 저장돼도 결과에서는 텍스트로 저장한다.
                row_values[account_column - 1] = clean_text(raw_account)

                if not product_name:
                    raise ValueError(
                        "ContractList.xlsx에서 "
                        f"상품명을 찾지 못했습니다: 이름={clean_text(raw_name)}, "
                        f"계좌번호={clean_text(raw_account)}. "
                        "A열 상품명이 병합되었거나 비어 있는지 확인해 주세요."
                    )

                if insert_pf_column:
                    row_values.insert(pf_output_index, product_name)
                else:
                    row_values[pf_output_index] = product_name

                filtered_rows.append(row_values)
    finally:
        workbook.close()

    if headers is None:
        raise ValueError(
            f"{source_file.name}에서 '계좌번호' 헤더가 있는 거래내역 시트를 찾지 못했습니다."
        )

    return headers, filtered_rows, matched_accounts


# ============================================================
# 결과 엑셀 작성
# ============================================================

def apply_sheet_format(worksheet, headers: Sequence, row_count: int) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin_gray = Side(style="thin", color="D9E2F3")
    border = Border(bottom=thin_gray)

    worksheet.freeze_panes = "A2"

    if headers:
        worksheet.auto_filter.ref = (
            f"A1:{get_column_letter(len(headers))}{max(row_count + 1, 1)}"
        )

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    header_keys = [normalize_header(value) for value in headers]

    for column_number, header_key in enumerate(header_keys, start=1):
        column_letter = get_column_letter(column_number)

        if header_key in ACCOUNT_HEADER_KEYS:
            worksheet.column_dimensions[column_letter].width = 20
            for row_number in range(2, worksheet.max_row + 1):
                worksheet.cell(row_number, column_number).number_format = "@"

        elif "일자" in header_key or "날짜" in header_key:
            worksheet.column_dimensions[column_letter].width = 13
            for row_number in range(2, worksheet.max_row + 1):
                worksheet.cell(row_number, column_number).number_format = "yyyy-mm-dd"

        elif any(
            word in header_key
            for word in ("수량", "단가", "금액", "수수료", "세금")
        ):
            worksheet.column_dimensions[column_letter].width = 16
            for row_number in range(2, worksheet.max_row + 1):
                worksheet.cell(row_number, column_number).number_format = "#,##0.00"

        else:
            max_length = len(clean_text(headers[column_number - 1]))
            for row_number in range(2, worksheet.max_row + 1):
                value = worksheet.cell(row_number, column_number).value
                max_length = max(max_length, len(clean_text(value)))
            worksheet.column_dimensions[column_letter].width = min(
                max(max_length + 2, 10),
                30,
            )

    for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center")


def build_output(
    output_path: Path,
    source_results: Sequence[Tuple[str, List, List[List]]],
) -> None:
    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    for sheet_name, headers, rows in source_results:
        worksheet = workbook.create_sheet(sheet_name)
        worksheet.append(list(headers))

        for row_values in rows:
            worksheet.append(row_values)

        apply_sheet_format(
            worksheet,
            headers,
            len(rows),
        )

    temporary_path = output_path.with_name(
        f".{output_path.stem}_작성중.xlsx"
    )

    if temporary_path.exists():
        temporary_path.unlink()

    try:
        workbook.save(temporary_path)
        workbook.close()
        temporary_path.replace(output_path)
    except PermissionError as exc:
        workbook.close()
        if temporary_path.exists():
            temporary_path.unlink()
        raise PermissionError(
            f"'{output_path.name}' 파일이 Excel에서 열려 있습니다. "
            "파일을 완전히 닫은 뒤 다시 실행해 주세요."
        ) from exc
    except Exception:
        workbook.close()
        if temporary_path.exists():
            temporary_path.unlink()
        raise


# ============================================================
# 메인 실행
# ============================================================

def main() -> None:
    contract_file = contract_output()

    if not contract_file.exists():
        raise FileNotFoundError(
            "먼저 contract_report_list.py로 "
            "Output/계약/ContractList.xlsx를 생성하세요."
        )

    report_contracts = load_report_contracts(contract_file)
    report_window_count = sum(len(windows) for windows in report_contracts.values())

    print(f"계약파일: {contract_file.name}")
    print(
        f"보고서 작성 대상: {len(report_contracts):,}개 계좌/이름 "
        f"({report_window_count:,}개 작성기간)"
    )

    source_results: List[Tuple[str, List, List[List]]] = []
    all_matched_accounts: Set[Tuple[str, str]] = set()

    for sheet_name, broker in SOURCE_CONFIG:
        source_files = normalized_trade_files(broker)
        headers = None
        rows: List[List] = []
        matched_accounts: Set[Tuple[str, str]] = set()

        for source_file in source_files:
            current_headers, current_rows, current_accounts = read_filtered_rows(
                source_file,
                report_contracts,
            )

            if headers is not None and current_headers != headers:
                raise ValueError(f"{broker} 정리 파일의 열 구성이 서로 다릅니다.")

            headers = current_headers
            rows.extend(current_rows)
            matched_accounts.update(current_accounts)

        source_results.append((sheet_name, headers, rows))
        all_matched_accounts.update(matched_accounts)

        print(
            f"[{sheet_name}] {len(source_files)}개 파일 -> "
            f"거래 {len(rows):,}건 / 매칭계좌 {len(matched_accounts):,}개"
        )

    output_path = final_dir() / OUTPUT_FILE_NAME
    build_output(output_path, source_results)

    unmatched_count = len(set(report_contracts) - all_matched_accounts)

    print("-" * 60)
    print(f"완료: {output_path}")
    print(
        f"5개 증권사 파일에서 매칭되지 않은 보고서 대상 계좌: "
        f"{unmatched_count:,}개"
    )


if __name__ == "__main__":
    main()
